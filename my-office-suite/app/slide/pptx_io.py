"""PPTX import/export (python-pptx) including transitions (p:transition) and entrance animations (p:timing)."""
from __future__ import annotations

import io
import logging
import os
import tempfile

from lxml import etree

from app.slide.model import (EMU_PER_PT, THEMES, Animation, Paragraph, Presentation, Run, Shape, Slide, Theme,
                             Transition)

log = logging.getLogger(__name__)

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
P = "{%s}" % NS["p"]
A = "{%s}" % NS["a"]

AUTOSHAPES = {"rect": "RECTANGLE", "roundrect": "ROUNDED_RECTANGLE", "ellipse": "OVAL", "triangle": "ISOSCELES_TRIANGLE",
              "diamond": "DIAMOND", "pentagon": "REGULAR_PENTAGON", "hexagon": "HEXAGON", "arrow_right": "RIGHT_ARROW",
              "arrow_left": "LEFT_ARROW", "star5": "STAR_5_POINT", "chevron": "CHEVRON"}
PRESET_IDS = {"appear": 1, "fly_in": 2, "fade": 10, "wipe": 22, "float_in": 42, "zoom": 53}
PRESET_EFFECTS = {v: k for k, v in PRESET_IDS.items()}
FLY_SUBTYPE = {"top": 1, "right": 2, "bottom": 4, "left": 8}
WIPE_FILTER = {"bottom": "wipe(up)", "top": "wipe(down)", "left": "wipe(right)", "right": "wipe(left)"}


class PresentationError(Exception):
    pass


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_color.lstrip("#")[:6].upper())


# ============================================================================= save
def save_pptx(path: str, pres: Presentation) -> None:
    from pptx import Presentation as PptxPresentation
    from pptx.util import Emu

    prs = PptxPresentation()
    prs.slide_width = Emu(pres.width)
    prs.slide_height = Emu(pres.height)
    layout = prs.slide_layouts[6]
    if pres.title:
        prs.core_properties.title = pres.title
    for s in pres.slides:
        sl = prs.slides.add_slide(layout)
        bg = s.background or (pres.theme.background if pres.theme.background.lower() != "#ffffff" else None)
        if bg:
            fill = sl.background.fill
            fill.solid()
            fill.fore_color.rgb = _rgb(bg)
        if s.background_image:
            pic = sl.shapes.add_picture(io.BytesIO(s.background_image), 0, 0, pres.width, pres.height)
            pic.name = "Background"
        id_map = {}
        for sh in s.shapes:
            try:
                el = _add_shape(sl, sh, pres)
            except Exception as exc:
                log.warning("Shape export failed (%s): %s", sh.kind, exc)
                continue
            if el is not None:
                id_map[sh.id] = el
        if s.notes:
            sl.notes_slide.notes_text_frame.text = s.notes
        if s.hidden:
            sl._element.set("show", "0")
        _write_transition(sl, s.transition)
        _write_timing(sl, s, id_map)
    d = os.path.dirname(os.path.abspath(path)) or "."
    fd, tmp = tempfile.mkstemp(dir=d, suffix=".pptx.tmp")
    os.close(fd)
    try:
        prs.save(tmp)
        from app.services.fileutil import finish_atomic

        finish_atomic(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def _fill_text(tf, sh: Shape, pres: Presentation):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(91440)
    tf.margin_top = tf.margin_bottom = Emu(45720)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(sh.valign, MSO_ANCHOR.TOP)
    paras = sh.paragraphs or [Paragraph([Run("")])]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
                       "justify": PP_ALIGN.JUSTIFY}.get(para.align, PP_ALIGN.LEFT)
        p.level = max(0, min(8, para.level))
        pPr = p._p.get_or_add_pPr()
        if para.rtl:
            pPr.set("rtl", "1")
        if para.bullet:
            pPr.set("marL", str(342900 + 457200 * para.level))
            pPr.set("indent", "-342900")
            bu = etree.SubElement(pPr, A + "buChar")
            bu.set("char", "\u2022")
        for run in para.runs:
            r = p.add_run()
            r.text = run.text
            f = r.font
            if run.size:
                f.size = Pt(run.size)
            if run.bold is not None:
                f.bold = run.bold
            if run.italic is not None:
                f.italic = run.italic
            if run.underline is not None:
                f.underline = run.underline
            if run.color:
                f.color.rgb = _rgb(run.color)
            if run.font:
                f.name = run.font
                rPr = r._r.get_or_add_rPr()
                for tag in ("cs", "ea"):
                    el = rPr.find(A + tag)
                    if el is None:
                        el = etree.SubElement(rPr, A + tag)
                    el.set("typeface", run.font)
            if run.gradient or (run.outline and run.outline_width > 0):
                _text_effects(r._r.get_or_add_rPr(), run)


def _text_effects(rPr, run: Run) -> None:
    """WordArt: text outline (a:ln) and gradient fill (a:gradFill), in the order the schema wants (ln, fill)."""
    if run.gradient:
        for old in rPr.findall(A + "solidFill"):
            rPr.remove(old)
        c1, _, c2 = run.gradient.partition(",")
        grad = etree.Element(A + "gradFill")
        lst = etree.SubElement(grad, A + "gsLst")
        for pos, col in (("0", c1), ("100000", c2 or c1)):
            gs = etree.SubElement(lst, A + "gs")
            gs.set("pos", pos)
            etree.SubElement(gs, A + "srgbClr").set("val", col.lstrip("#").upper())
        lin = etree.SubElement(grad, A + "lin")
        lin.set("ang", "5400000")
        lin.set("scaled", "0")
        rPr.insert(0, grad)
    if run.outline and run.outline_width > 0:
        ln = etree.Element(A + "ln")
        ln.set("w", str(int(run.outline_width * 12700)))
        fill = etree.SubElement(ln, A + "solidFill")
        etree.SubElement(fill, A + "srgbClr").set("val", run.outline.lstrip("#").upper())
        rPr.insert(0, ln)


def _read_effects(r, run: Run) -> None:
    rPr = r._r.find(A + "rPr")
    if rPr is None:
        return
    ln = rPr.find(A + "ln")
    if ln is not None:
        clr = ln.find(f"{A}solidFill/{A}srgbClr")
        if clr is not None and ln.find(A + "noFill") is None:
            run.outline = "#" + clr.get("val", "000000").lower()
            run.outline_width = int(ln.get("w", "12700")) / 12700
    grad = rPr.find(A + "gradFill")
    if grad is not None:
        cols = ["#" + c.get("val", "000000").lower() for c in grad.iter(A + "srgbClr")]
        if cols:
            run.gradient = f"{cols[0]},{cols[-1]}"
            run.color = cols[0]


def _add_shape(sl, sh: Shape, pres: Presentation):
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Emu, Pt

    x, y, w, h = Emu(sh.x), Emu(sh.y), Emu(max(1, abs(sh.w))), Emu(max(1, abs(sh.h)))
    if sh.kind == "text":
        shape = sl.shapes.add_textbox(x, y, w, h)
        _fill_text(shape.text_frame, sh, pres)
        if sh.fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(sh.fill)
        if sh.line:
            shape.line.color.rgb = _rgb(sh.line)
            shape.line.width = Pt(sh.line_width)
    elif sh.kind in AUTOSHAPES:
        shape = sl.shapes.add_shape(getattr(MSO_SHAPE, AUTOSHAPES[sh.kind]), x, y, w, h)
        if sh.fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(sh.fill)
        else:
            shape.fill.background()
        if sh.line:
            shape.line.color.rgb = _rgb(sh.line)
            shape.line.width = Pt(sh.line_width)
        else:
            shape.line.fill.background()
        if sh.paragraphs and sh.text.strip():
            _fill_text(shape.text_frame, sh, pres)
        if not sh.shadow:
            sppr = shape._element.spPr
            etree.SubElement(sppr, A + "effectLst")
    elif sh.kind in ("line", "arrow_line"):
        shape = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(sh.x), Emu(sh.y), Emu(sh.x + sh.w), Emu(sh.y + sh.h))
        shape.line.color.rgb = _rgb(sh.line or "#404040")
        shape.line.width = Pt(sh.line_width or 1.5)
        if sh.kind == "arrow_line":
            ln = shape.line._get_or_add_ln()
            tail = etree.SubElement(ln, A + "tailEnd")
            tail.set("type", "triangle")
    elif sh.kind == "image" and sh.image:
        shape = sl.shapes.add_picture(io.BytesIO(_ensure_supported_image(sh.image, sh.image_ext)), x, y, w, h)
    elif sh.kind == "table" and sh.table:
        rows, cols = len(sh.table), max(len(r) for r in sh.table)
        gf = sl.shapes.add_table(rows, cols, x, y, w, h)
        tbl = gf.table
        tbl.first_row = bool(sh.table_style.get("header", True))
        tbl.horz_banding = bool(sh.table_style.get("banded", True))
        for i, row in enumerate(sh.table):
            for j in range(cols):
                cell = tbl.cell(i, j)
                cell.text = str(row[j]) if j < len(row) else ""
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(14)
        shape = gf
    elif sh.kind == "chart" and sh.chart:
        shape = _add_chart(sl, sh, x, y, w, h)
    else:
        return None
    if sh.rotation:
        shape.rotation = sh.rotation
    if sh.name:
        shape.name = sh.name
    return shape


def _ensure_supported_image(data: bytes, ext: str) -> bytes:
    if ext.lower() in ("png", "jpg", "jpeg", "gif", "bmp", "tiff", "tif"):
        return data
    from PIL import Image

    buf = io.BytesIO()
    Image.open(io.BytesIO(data)).save(buf, "PNG")
    return buf.getvalue()


def _add_chart(sl, sh: Shape, x, y, w, h):
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    ch = sh.chart
    kind = ch.get("kind", "column")
    if kind == "scatter":
        data = XyChartData()
        xs = [float(v) for v in ch.get("categories", [])]
        for name, vals in ch.get("series", []):
            s = data.add_series(name)
            for xv, yv in zip(xs, vals):
                if yv is not None:
                    s.add_data_point(xv, float(yv))
        ctype = XL_CHART_TYPE.XY_SCATTER
    else:
        data = CategoryChartData()
        data.categories = [str(c) for c in ch.get("categories", [])]
        for name, vals in ch.get("series", []):
            data.add_series(str(name), [None if v is None else float(v) for v in vals])
        ctype = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA, "pie": XL_CHART_TYPE.PIE,
                 "doughnut": XL_CHART_TYPE.DOUGHNUT}.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    gf = sl.shapes.add_chart(ctype, x, y, w, h, data)
    chart = gf.chart
    if ch.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = ch["title"]
    chart.has_legend = bool(ch.get("legend", True))
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    return gf


def _write_transition(sl, tr: Transition):
    if tr is None or tr.kind == "none":
        return
    root = sl._element
    spd = "fast" if tr.duration < 0.5 else ("med" if tr.duration < 1.0 else "slow")
    t = etree.Element(P + "transition")
    t.set("spd", spd)
    d = {"left": "l", "right": "r", "top": "u", "bottom": "d"}.get(tr.direction, "l")
    if tr.kind == "fade":
        etree.SubElement(t, P + "fade")
    elif tr.kind in ("push", "wipe", "cover"):
        etree.SubElement(t, P + tr.kind).set("dir", d)
    elif tr.kind == "zoom":
        etree.SubElement(t, P + "zoom")
    elif tr.kind == "split":
        e = etree.SubElement(t, P + "split")
        e.set("orient", "horz")
        e.set("dir", "out")
    anchor = root.find(P + "clrMapOvr")
    if anchor is None:
        anchor = root.find(P + "cSld")
    anchor.addnext(t)


def _ctn(parent, ids, **attrs):
    ids[0] += 1
    el = etree.SubElement(parent, P + "cTn")
    el.set("id", str(ids[0]))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _cond(parent, tag, delay):
    lst = etree.SubElement(parent, P + tag)
    c = etree.SubElement(lst, P + "cond")
    c.set("delay", str(delay))
    return lst


def _bhvr(parent, ids, spid, dur, attr_names=None, additive=None):
    cb = etree.SubElement(parent, P + "cBhvr")
    if additive:
        cb.set("additive", additive)
    ctn = _ctn(cb, ids, dur=dur, fill="hold")
    if dur == 1:
        _cond(ctn, "stCondLst", 0)
    tgt = etree.SubElement(cb, P + "tgtEl")
    etree.SubElement(tgt, P + "spTgt").set("spid", str(spid))
    if attr_names:
        lst = etree.SubElement(cb, P + "attrNameLst")
        for n in attr_names:
            etree.SubElement(lst, P + "attrName").text = n
    return cb


def _anim_prop(parent, ids, spid, dur, attr, frm, to):
    anim = etree.SubElement(parent, P + "anim")
    anim.set("calcmode", "lin")
    anim.set("valueType", "num")
    _bhvr(anim, ids, spid, dur, [attr], additive="base")
    tav = etree.SubElement(anim, P + "tavLst")
    for tm, val in (("0", frm), ("100000", to)):
        t = etree.SubElement(tav, P + "tav")
        t.set("tm", tm)
        v = etree.SubElement(t, P + "val")
        etree.SubElement(v, P + "strVal").set("val", val)


def _write_timing(sl, s: Slide, id_map: dict):
    anims = [sh for sh in s.animated_shapes() if sh.id in id_map]
    if not anims:
        return
    ids = [0]
    timing = etree.Element(P + "timing")
    tn = etree.SubElement(timing, P + "tnLst")
    par = etree.SubElement(tn, P + "par")
    root = _ctn(par, ids, dur="indefinite", restart="never", nodeType="tmRoot")
    rchild = etree.SubElement(root, P + "childTnLst")
    seq = etree.SubElement(rchild, P + "seq")
    seq.set("concurrent", "1")
    seq.set("nextAc", "seek")
    main = _ctn(seq, ids, dur="indefinite", nodeType="mainSeq")
    mchild = etree.SubElement(main, P + "childTnLst")
    click_children = None
    inner_children = None
    inner_end = 0
    bld = []
    for i, sh in enumerate(anims):
        an = sh.animation
        spid = id_map[sh.id].shape_id
        dur = max(1, int(an.duration * 1000))
        delay_ms = int(an.delay * 1000)
        trigger = an.trigger if i else "click"
        if trigger == "click" or click_children is None:
            cpar = etree.SubElement(mchild, P + "par")
            cctn = _ctn(cpar, ids, fill="hold")
            _cond(cctn, "stCondLst", "indefinite")
            click_children = etree.SubElement(cctn, P + "childTnLst")
            inner_children = None
            inner_end = 0
            node_type = "clickEffect"
        else:
            node_type = "withEffect" if trigger == "with" else "afterEffect"
        if inner_children is None or trigger == "after":
            ipar = etree.SubElement(click_children, P + "par")
            ictn = _ctn(ipar, ids, fill="hold")
            _cond(ictn, "stCondLst", inner_end if trigger == "after" else 0)
            inner_children = etree.SubElement(ictn, P + "childTnLst")
            start_offset = 0
        else:
            start_offset = 0
        epar = etree.SubElement(inner_children, P + "par")
        preset = PRESET_IDS.get(an.effect, 10)
        subtype = FLY_SUBTYPE.get(an.direction, 4) if an.effect in ("fly_in", "wipe") else (16 if an.effect == "zoom" else 0)
        has_text = id_map[sh.id]._element.find(f"{P}txBody") is not None
        ectn = _ctn(epar, ids, presetID=preset, presetClass="entr", presetSubtype=subtype, fill="hold",
                    nodeType=node_type, **({"grpId": "0"} if has_text else {}))
        _cond(ectn, "stCondLst", delay_ms + start_offset)
        ech = etree.SubElement(ectn, P + "childTnLst")
        st = etree.SubElement(ech, P + "set")
        _bhvr(st, ids, spid, 1, ["style.visibility"])
        to = etree.SubElement(st, P + "to")
        etree.SubElement(to, P + "strVal").set("val", "visible")
        if an.effect in ("fade", "zoom", "float_in"):
            ae = etree.SubElement(ech, P + "animEffect")
            ae.set("transition", "in")
            ae.set("filter", "fade")
            _bhvr(ae, ids, spid, dur)
        if an.effect == "wipe":
            ae = etree.SubElement(ech, P + "animEffect")
            ae.set("transition", "in")
            ae.set("filter", WIPE_FILTER.get(an.direction, "wipe(up)"))
            _bhvr(ae, ids, spid, dur)
        if an.effect == "fly_in":
            d = an.direction
            fx = {"left": "0-#ppt_w/2", "right": "1+#ppt_w/2"}.get(d, "#ppt_x")
            fy = {"top": "0-#ppt_h/2", "bottom": "1+#ppt_h/2"}.get(d, "#ppt_y")
            _anim_prop(ech, ids, spid, dur, "ppt_x", fx, "#ppt_x")
            _anim_prop(ech, ids, spid, dur, "ppt_y", fy, "#ppt_y")
        if an.effect == "zoom":
            _anim_prop(ech, ids, spid, dur, "ppt_w", "0", "#ppt_w")
            _anim_prop(ech, ids, spid, dur, "ppt_h", "0", "#ppt_h")
        if an.effect == "float_in":
            _anim_prop(ech, ids, spid, dur, "ppt_y", "#ppt_y+.1", "#ppt_y")
        inner_end = max(inner_end, delay_ms + dur) if trigger != "after" else inner_end + delay_ms + dur
        if has_text and spid not in bld:
            bld.append(spid)
    prev = etree.SubElement(seq, P + "prevCondLst")
    c = etree.SubElement(prev, P + "cond")
    c.set("evt", "onPrev")
    c.set("delay", "0")
    etree.SubElement(etree.SubElement(c, P + "tgtEl"), P + "sldTgt")
    nxt = etree.SubElement(seq, P + "nextCondLst")
    c = etree.SubElement(nxt, P + "cond")
    c.set("evt", "onNext")
    c.set("delay", "0")
    etree.SubElement(etree.SubElement(c, P + "tgtEl"), P + "sldTgt")
    if bld:
        bl = etree.SubElement(timing, P + "bldLst")
        for spid in bld:
            b = etree.SubElement(bl, P + "bldP")
            b.set("spid", str(spid))
            b.set("grpId", "0")
    root_el = sl._element
    anchor = root_el.find(P + "transition")
    if anchor is None:
        anchor = root_el.find(P + "clrMapOvr")
    if anchor is None:
        anchor = root_el.find(P + "cSld")
    anchor.addnext(timing)


# ============================================================================= load
def _color_of(fmt_color, theme: Theme) -> str | None:
    try:
        from pptx.enum.dml import MSO_COLOR_TYPE

        if fmt_color is None or fmt_color.type is None:
            return None
        if fmt_color.type == MSO_COLOR_TYPE.RGB:
            return "#" + str(fmt_color.rgb).lower()
        if fmt_color.type == MSO_COLOR_TYPE.SCHEME:
            name = str(fmt_color.theme_color).split(".")[-1].split(" ")[0].upper()
            return {"ACCENT_1": "#4472c4", "ACCENT_2": "#ed7d31", "ACCENT_3": "#a5a5a5", "ACCENT_4": "#ffc000",
                    "ACCENT_5": "#5b9bd5", "ACCENT_6": "#70ad47", "DARK_1": "#000000", "LIGHT_1": "#ffffff",
                    "DARK_2": "#44546a", "LIGHT_2": "#e7e6e6", "TEXT_1": "#000000", "BACKGROUND_1": "#ffffff",
                    "TEXT_2": "#44546a", "BACKGROUND_2": "#e7e6e6", "HYPERLINK": "#0563c1"}.get(name, theme.accent)
    except Exception:
        return None
    return None


def _read_text(tf, placeholder: str | None, theme: Theme) -> list:
    from pptx.enum.text import PP_ALIGN

    paras = []
    default_size = {"title": 40.0, "subtitle": 24.0, "body": 24.0}.get(placeholder or "", 18.0)
    for p in tf.paragraphs:
        pPr = p._p.find(A + "pPr")
        bullet = placeholder == "body"
        if pPr is not None:
            if pPr.find(A + "buNone") is not None:
                bullet = False
            if pPr.find(A + "buChar") is not None or pPr.find(A + "buAutoNum") is not None:
                bullet = True
        align = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right", PP_ALIGN.JUSTIFY: "justify"}.get(p.alignment, "left")
        if placeholder in ("title", "subtitle") and p.alignment is None and placeholder == "subtitle":
            align = "center"
        runs = []
        for r in p.runs:
            f = r.font
            size = f.size.pt if f.size is not None else (default_size if placeholder != "body" or p.level == 0
                                                          else max(12.0, default_size - 4 * p.level))
            # read WordArt effects from the XML first: python-pptx turns a gradient into a solid fill as soon as
            # font.color is accessed
            fx = Run("")
            _read_effects(r, fx)
            color = fx.color
            if fx.gradient is None:
                try:
                    color = _color_of(f.color, theme) if f.color and f.color.type is not None else None
                except Exception:
                    color = None
            run = Run(r.text, f.bold, f.italic, bool(f.underline) if f.underline is not None else None, size,
                      color, f.name, fx.outline, fx.outline_width, fx.gradient)
            runs.append(run)
        if not runs:
            runs = [Run("", size=default_size)]
        rtl = pPr is not None and pPr.get("rtl") == "1"
        paras.append(Paragraph(runs, align, p.level, bullet, rtl))
    return paras


def load_pptx(path: str) -> Presentation:
    from pptx import Presentation as PptxPresentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.exc import PackageNotFoundError

    try:
        prs = PptxPresentation(path)
    except PackageNotFoundError as exc:
        raise PresentationError(f"Not a PowerPoint file: {exc}") from exc
    except Exception as exc:
        raise PresentationError(f"The presentation is damaged or not a PPTX file ({exc.__class__.__name__}: {exc})") from exc
    pres = Presentation([], int(prs.slide_width or 12192000), int(prs.slide_height or 6858000), Theme(),
                        prs.core_properties.title or "")
    theme = pres.theme
    for sl in prs.slides:
        s = Slide(layout="blank")
        try:
            if sl.background.fill.type is not None:
                s.background = _color_of(sl.background.fill.fore_color, theme)
        except Exception:
            pass
        spid_map = {}
        for shape in sl.shapes:
            for sh in _convert_shape(shape, theme, PP_PLACEHOLDER, MSO_SHAPE_TYPE):
                s.shapes.append(sh)
                spid_map.setdefault(shape.shape_id, sh)
        if sl.has_notes_slide:
            s.notes = sl.notes_slide.notes_text_frame.text if sl.notes_slide.notes_text_frame is not None else ""
        s.hidden = sl._element.get("show") == "0"
        s.transition = _read_transition(sl._element)
        _read_timing(sl._element, spid_map)
        pres.slides.append(s)
    return pres


def _convert_shape(shape, theme, PP_PLACEHOLDER, MSO_SHAPE_TYPE, offset=None) -> list:
    out = []
    try:
        x, y, w, h = int(shape.left or 0), int(shape.top or 0), int(shape.width or 0), int(shape.height or 0)
    except Exception:
        return out
    if offset is not None:
        gx, gy, sx, sy, ox, oy = offset
        x, y, w, h = int(gx + (x - ox) * sx), int(gy + (y - oy) * sy), int(w * sx), int(h * sy)
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        grp = shape._element
        xfrm = grp.find(f"{P}grpSpPr/{A}xfrm")
        if xfrm is not None:
            ch_off, ch_ext = xfrm.find(A + "chOff"), xfrm.find(A + "chExt")
            ox, oy = int(ch_off.get("x")), int(ch_off.get("y"))
            cx, cy = max(1, int(ch_ext.get("cx"))), max(1, int(ch_ext.get("cy")))
            off = (x, y, w / cx, h / cy, ox, oy)
        else:
            off = (0, 0, 1, 1, 0, 0)
        for child in shape.shapes:
            out.extend(_convert_shape(child, theme, PP_PLACEHOLDER, MSO_SHAPE_TYPE, off))
        return out
    sh = Shape("text", x, y, w, h, rotation=float(getattr(shape, "rotation", 0) or 0), name=shape.name or "")
    placeholder = None
    if shape.is_placeholder:
        pt = shape.placeholder_format.type
        placeholder = {PP_PLACEHOLDER.TITLE: "title", PP_PLACEHOLDER.CENTER_TITLE: "title",
                       PP_PLACEHOLDER.SUBTITLE: "subtitle", PP_PLACEHOLDER.BODY: "body",
                       PP_PLACEHOLDER.OBJECT: "body"}.get(pt)
        sh.placeholder = placeholder
    if st == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, "image")):
        try:
            sh.kind = "image"
            sh.image = shape.image.blob
            sh.image_ext = shape.image.ext or "png"
            out.append(sh)
        except Exception:
            pass
        return out
    if getattr(shape, "has_table", False) and shape.has_table:
        sh.kind = "table"
        sh.table = [[cell.text for cell in row.cells] for row in shape.table.rows]
        out.append(sh)
        return out
    if getattr(shape, "has_chart", False) and shape.has_chart:
        sh.kind = "chart"
        sh.chart = _read_chart(shape.chart)
        out.append(sh)
        return out
    if st in (MSO_SHAPE_TYPE.LINE,) or shape._element.tag.endswith("}cxnSp"):
        sh.kind = "line"
        try:
            bx, by, ex, ey = shape.begin_x, shape.begin_y, shape.end_x, shape.end_y
            sh.x, sh.y, sh.w, sh.h = bx, by, ex - bx, ey - by
        except Exception:
            pass
        ln = shape._element.find(f".//{A}ln")
        if ln is not None and ln.find(A + "tailEnd") is not None and ln.find(A + "tailEnd").get("type") not in (None, "none"):
            sh.kind = "arrow_line"
        try:
            sh.line = _color_of(shape.line.color, theme) or "#404040"
            sh.line_width = shape.line.width.pt if shape.line.width else 1.5
        except Exception:
            sh.line = "#404040"
        out.append(sh)
        return out
    if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
        try:
            ast = str(shape.auto_shape_type).split(".")[-1].split(" ")[0]
        except Exception:
            ast = "RECTANGLE"
        rev = {v: k for k, v in AUTOSHAPES.items()}
        sh.kind = rev.get(ast, "rect")
        try:
            from pptx.enum.dml import MSO_FILL

            if shape.fill.type == MSO_FILL.SOLID:
                sh.fill = _color_of(shape.fill.fore_color, theme)
            elif shape.fill.type is None:
                sh.fill = "#4472c4"
        except Exception:
            sh.fill = "#4472c4"
        try:
            sh.line = _color_of(shape.line.color, theme) if shape.line.fill.type is not None else None
            if shape.line.width:
                sh.line_width = shape.line.width.pt
        except Exception:
            sh.line = None
    elif st not in (MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER, None) and not shape.has_text_frame:
        return out
    else:
        try:
            from pptx.enum.dml import MSO_FILL

            if shape.fill.type == MSO_FILL.SOLID:
                sh.fill = _color_of(shape.fill.fore_color, theme)
        except Exception:
            pass
    if shape.has_text_frame:
        sh.paragraphs = _read_text(shape.text_frame, placeholder, theme)
        try:
            from pptx.enum.text import MSO_ANCHOR

            va = shape.text_frame.vertical_anchor
            sh.valign = {MSO_ANCHOR.MIDDLE: "middle", MSO_ANCHOR.BOTTOM: "bottom"}.get(va, "middle" if sh.kind != "text" else "top")
        except Exception:
            pass
    out.append(sh)
    return out


def _read_chart(chart) -> dict:
    from pptx.enum.chart import XL_CHART_TYPE

    ct = chart.chart_type
    kind = "column"
    if ct in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED):
        kind = "bar"
    elif ct in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
        kind = "line"
    elif ct in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED):
        kind = "pie"
    elif ct == XL_CHART_TYPE.DOUGHNUT:
        kind = "doughnut"
    elif ct in (XL_CHART_TYPE.AREA, XL_CHART_TYPE.AREA_STACKED):
        kind = "area"
    elif ct in (XL_CHART_TYPE.XY_SCATTER, XL_CHART_TYPE.XY_SCATTER_LINES):
        kind = "scatter"
    cats, series = [], []
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        for s in plot.series:
            series.append([s.name, [None if v is None else float(v) for v in s.values]])
    except Exception:
        pass
    title = ""
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text
    except Exception:
        pass
    return {"kind": kind, "categories": cats, "series": series, "title": title, "legend": chart.has_legend}


def _read_transition(el) -> Transition:
    t = el.find(P + "transition")
    if t is None:
        return Transition()
    spd = {"fast": 0.4, "med": 0.7, "slow": 1.0}.get(t.get("spd", "med"), 0.7)
    for child in t:
        name = etree.QName(child).localname
        if name in ("fade", "push", "wipe", "cover", "zoom", "split"):
            d = {"l": "left", "r": "right", "u": "top", "d": "bottom"}.get(child.get("dir", "l"), "left")
            return Transition(name, spd, d)
    return Transition("fade", spd)


def _read_timing(el, spid_map: dict):
    timing = el.find(P + "timing")
    if timing is None:
        return
    order = 0
    for ctn in timing.iter(P + "cTn"):
        if ctn.get("presetClass") != "entr":
            continue
        tgt = ctn.find(f".//{P}spTgt")
        if tgt is None:
            continue
        sh = spid_map.get(int(tgt.get("spid", "0")))
        if sh is None:
            continue
        preset = int(ctn.get("presetID", "10"))
        effect = PRESET_EFFECTS.get(preset, "fade")
        node = ctn.get("nodeType", "clickEffect")
        trigger = {"clickEffect": "click", "withEffect": "with", "afterEffect": "after"}.get(node, "click")
        dur = 0.5
        for inner in ctn.iter(P + "cTn"):
            if inner is not ctn and inner.get("dur") not in (None, "1", "indefinite"):
                try:
                    dur = int(inner.get("dur")) / 1000
                    break
                except ValueError:
                    pass
        delay = 0.0
        cond = ctn.find(f"{P}stCondLst/{P}cond")
        if cond is not None and cond.get("delay", "0").isdigit():
            delay = int(cond.get("delay")) / 1000
        sub = int(ctn.get("presetSubtype", "4") or 4)
        direction = {1: "top", 2: "right", 4: "bottom", 8: "left"}.get(sub, "bottom")
        sh.animation = Animation(effect, trigger, dur, delay, direction, order)
        order += 1


def new_presentation(theme_name: str = "Office") -> Presentation:
    from app.slide.model import make_slide

    pres = Presentation(theme=Theme(**THEMES.get(theme_name, Theme()).__dict__))
    pres.slides.append(make_slide(pres, "title", "", ""))
    return pres


__all__ = ["save_pptx", "load_pptx", "new_presentation", "PresentationError", "EMU_PER_PT"]
