"""Resume Master: fill in your details once, pick one of the designs, get a finished CV as an editable DOCX
(Word, MyOffice Writer) or as a PDF."""
from __future__ import annotations

import io
import json
import os

from PySide6.QtCore import Qt
from PySide6.QtGui import QPixmap
from PySide6.QtWidgets import (QComboBox, QDialog, QDialogButtonBox, QFormLayout, QHBoxLayout, QLabel, QLineEdit,
                               QPlainTextEdit, QPushButton, QTabWidget, QVBoxLayout, QWidget)

from app.services.i18n import _t
from app.services.paths import user_subdir

STYLES = {  # name -> (accent colour, heading font, body font, layout)
    "Modern Blue": ("1A73E8", "Calibri Light", "Calibri", "sidebar"),
    "Classic": ("202124", "Georgia", "Georgia", "classic"),
    "Elegant Green": ("188038", "Cambria", "Calibri", "classic"),
    "Creative Orange": ("E8710A", "Segoe UI Semibold", "Segoe UI", "sidebar"),
    "Minimal Grey": ("5F6368", "Calibri Light", "Calibri", "minimal"),
    "Royal Purple": ("6A1B9A", "Calibri Light", "Calibri", "sidebar"),
}

FIELDS = [("name", "Full name"), ("title", "Job title"), ("email", "E-mail"), ("phone", "Phone"),
          ("location", "City, country"), ("links", "Website / LinkedIn")]
SECTIONS = [("summary", "Profile", "A short summary of who you are and what you offer."),
            ("experience", "Experience",
             "One job per block, separated by an empty line:\nJob title | Company | 2022 - now\n- what you achieved\n- ..."),
            ("education", "Education", "Degree | School | Year\n(empty line between entries)"),
            ("skills", "Skills", "One per line or comma-separated"),
            ("languages", "Languages", "e.g. Urdu (native), English (fluent)"),
            ("extra", "Certificates, projects, awards", "One per line")]
SAMPLE = {
    "name": "Ayesha Khan", "title": "Accountant", "email": "ayesha.khan@example.com", "phone": "+92 300 1234567",
    "location": "Lahore, Pakistan", "links": "linkedin.com/in/ayeshakhan",
    "summary": "Detail-oriented accountant with 6 years of experience in financial reporting, payroll and tax "
               "compliance for manufacturing companies.",
    "experience": "Senior Accountant | Punjab Textiles Ltd. | 2021 - now\n- Prepared monthly IFRS reports for "
                  "4 business units\n- Cut the month-end close from 9 to 5 days\n\nAccountant | Gulberg Traders | "
                  "2018 - 2021\n- Managed payroll for 220 employees\n- Handled sales tax returns with FBR",
    "education": "B.Com (Hons) | University of the Punjab | 2018",
    "skills": "Financial reporting, IFRS, Payroll, Sales tax, Excel (pivot tables, XLOOKUP), QuickBooks",
    "languages": "Urdu (native), English (fluent), Punjabi (native)",
    "extra": "ACCA - 9 papers passed\nAdvanced Excel certificate, 2023",
}


def _blocks(text: str) -> list[list[str]]:
    out, cur = [], []
    for line in text.splitlines():
        if not line.strip():
            if cur:
                out.append(cur)
                cur = []
        else:
            cur.append(line.rstrip())
    if cur:
        out.append(cur)
    return out


def _items(text: str) -> list[str]:
    parts = [p.strip() for line in text.splitlines() for p in (line.split(",") if "," in line and len(line) < 200
                                                             else [line])]
    return [p.lstrip("-\u2022 ").strip() for p in parts if p.strip()]


def build_docx(data: dict, style: str, path: str, photo: bytes | None = None) -> str:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    accent, hfont, bfont, layout = STYLES[style]
    rgb = RGBColor.from_string(accent)
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    normal = d.styles["Normal"]
    normal.font.name = bfont
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    def shade(cell, color):
        tcpr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), color)
        tcpr.append(sh)

    def heading(container, text, color=rgb, size=12.5, rule=True):
        p = container.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text.upper() if layout != "classic" else text)
        r.bold = True
        r.font.size = Pt(size)
        r.font.name = hfont
        r.font.color.rgb = color
        if rule:
            ppr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", accent)):
                bottom.set(qn(f"w:{k}"), v)
            bdr.append(bottom)
            ppr.append(bdr)
        return p

    def experience(container, text):
        for block in _blocks(text):
            head = [s.strip() for s in block[0].split("|")]
            p = container.add_paragraph()
            r = p.add_run(head[0])
            r.bold = True
            if len(head) > 1:
                p.add_run("  \u2022  " + head[1])
            if len(head) > 2:
                r3 = p.add_run("    " + head[2])
                r3.italic = True
                r3.font.color.rgb = RGBColor(0x5F, 0x63, 0x68)
            for line in block[1:]:
                bp = container.add_paragraph(line.lstrip("-\u2022 ").strip(), style="List Bullet")
                bp.paragraph_format.space_after = Pt(1)

    def contact_lines():
        return [data.get(k, "") for k in ("email", "phone", "location", "links") if data.get(k, "").strip()]

    if layout == "sidebar":
        t = d.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        left, right = t.rows[0].cells
        left.width, right.width = Cm(6.0), Cm(11.8)
        shade(left, accent)
        white = RGBColor(0xFF, 0xFF, 0xFF)
        lp = left.paragraphs[0]
        if photo:
            lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lp.add_run().add_picture(io.BytesIO(photo), width=Cm(3.6))
            lp = left.add_paragraph()
        r = lp.add_run(data.get("name", ""))
        r.bold = True
        r.font.size = Pt(18)
        r.font.name = hfont
        r.font.color.rgb = white
        if data.get("title"):
            tp = left.add_paragraph()
            tr = tp.add_run(data["title"])
            tr.font.color.rgb = white
            tr.font.size = Pt(11)
        heading(left, _t("Contact"), white, 11, rule=False)
        for line in contact_lines():
            cp = left.add_paragraph()
            cr = cp.add_run(line)
            cr.font.color.rgb = white
            cr.font.size = Pt(9.5)
        for key, label in (("skills", _t("Skills")), ("languages", _t("Languages"))):
            if data.get(key, "").strip():
                heading(left, label, white, 11, rule=False)
                for it in _items(data[key]):
                    sp = left.add_paragraph()
                    sr = sp.add_run("\u2022 " + it)
                    sr.font.color.rgb = white
                    sr.font.size = Pt(9.5)
        target = right
        right.paragraphs[0].text = ""
        order = (("summary", _t("Profile")), ("experience", _t("Experience")), ("education", _t("Education")),
                 ("extra", _t("Certificates, projects, awards")))
    else:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if layout == "classic" else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(data.get("name", ""))
        r.bold = True
        r.font.size = Pt(24)
        r.font.name = hfont
        r.font.color.rgb = rgb
        if data.get("title"):
            tp = d.add_paragraph()
            tp.alignment = p.alignment
            tr = tp.add_run(data["title"])
            tr.font.size = Pt(12)
        cp = d.add_paragraph("  |  ".join(contact_lines()))
        cp.alignment = p.alignment
        for run in cp.runs:
            run.font.size = Pt(9.5)
        target = d
        order = (("summary", _t("Profile")), ("experience", _t("Experience")), ("education", _t("Education")),
                 ("skills", _t("Skills")), ("languages", _t("Languages")),
                 ("extra", _t("Certificates, projects, awards")))
    for key, label in order:
        text = data.get(key, "").strip()
        if not text:
            continue
        heading(target, label)
        if key in ("experience", "education"):
            experience(target, text)
        elif key in ("skills", "languages", "extra"):
            if layout == "minimal" and key == "skills":
                target.add_paragraph(", ".join(_items(text)))
            else:
                for it in _items(text):
                    target.add_paragraph(it, style="List Bullet")
        else:
            target.add_paragraph(text)
    d.core_properties.title = f"{data.get('name', '')} - CV"
    d.save(path)
    return path


def docx_to_pdf(docx_path: str, pdf_path: str) -> str:
    from app.pdf.ops import to_pdf

    return to_pdf(docx_path, pdf_path)


class ResumeDialog(QDialog):
    def __init__(self, main):
        super().__init__(main)
        self.main = main
        self.setWindowTitle(_t("Resume Master"))
        self.resize(820, 680)
        self.photo: bytes | None = None
        self.store = user_subdir("resume") / "resume.json"
        v = QVBoxLayout(self)
        top = QHBoxLayout()
        top.addWidget(QLabel(_t("Design:")))
        self.style = QComboBox()
        self.style.addItems(list(STYLES))
        top.addWidget(self.style)
        sample = QPushButton(_t("Fill with Example"))
        sample.clicked.connect(lambda: self._fill(SAMPLE))
        photo = QPushButton(_t("Photo\u2026"))
        photo.clicked.connect(self._photo)
        self.photo_label = QLabel()
        self.photo_label.setFixedSize(48, 60)
        top.addStretch(1)
        top.addWidget(self.photo_label)
        top.addWidget(photo)
        top.addWidget(sample)
        v.addLayout(top)
        tabs = QTabWidget()
        pw = QWidget()
        form = QFormLayout(pw)
        self.edits: dict = {}
        for key, label in FIELDS:
            e = QLineEdit()
            form.addRow(_t(label) + ":", e)
            self.edits[key] = e
        tabs.addTab(pw, _t("Personal details"))
        for key, label, hint in SECTIONS:
            e = QPlainTextEdit()
            e.setPlaceholderText(_t(hint))
            self.edits[key] = e
            tabs.addTab(e, _t(label))
        v.addWidget(tabs, 1)
        bb = QDialogButtonBox()
        docx_btn = bb.addButton(_t("Create Word Document"), QDialogButtonBox.AcceptRole)
        pdf_btn = bb.addButton(_t("Create PDF"), QDialogButtonBox.ActionRole)
        bb.addButton(QDialogButtonBox.Close)
        docx_btn.clicked.connect(lambda: self.create("docx"))
        pdf_btn.clicked.connect(lambda: self.create("pdf"))
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        try:
            saved = json.loads(self.store.read_text(encoding="utf-8"))
            self._fill(saved.get("data", {}))
            self.style.setCurrentText(saved.get("style", "Modern Blue"))
        except (OSError, ValueError):
            pass

    def data(self) -> dict:
        return {k: (e.text() if isinstance(e, QLineEdit) else e.toPlainText()) for k, e in self.edits.items()}

    def _fill(self, data: dict):
        for k, e in self.edits.items():
            val = data.get(k, "")
            e.setText(val) if isinstance(e, QLineEdit) else e.setPlainText(val)

    def _photo(self):
        from app.shell.dialogs import get_open_file

        path = get_open_file(self, _t("Photo"), "Images (*.png *.jpg *.jpeg)")
        if path:
            from PIL import Image, ImageOps

            img = ImageOps.fit(Image.open(path).convert("RGB"), (420, 525), Image.LANCZOS, centering=(0.5, 0.35))
            buf = io.BytesIO()
            img.save(buf, "JPEG", quality=88)
            self.photo = buf.getvalue()
            pm = QPixmap()
            pm.loadFromData(self.photo)
            self.photo_label.setPixmap(pm.scaled(self.photo_label.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation))

    def create(self, kind: str):
        from app.shell.dialogs import get_save_file, handle_exception

        data = self.data()
        if not data.get("name", "").strip():
            self.edits["name"].setFocus()
            return
        self.store.write_text(json.dumps({"data": data, "style": self.style.currentText()}, ensure_ascii=False),
                              encoding="utf-8")
        base = data["name"].strip().replace(" ", "_") + "_CV"
        path = get_save_file(self, _t("Save CV"), "Word (*.docx)" if kind == "docx" else "PDF (*.pdf)",
                             f"{base}.{kind}", kind)
        if not path:
            return
        try:
            if kind == "docx":
                build_docx(data, self.style.currentText(), path, self.photo)
            else:
                tmp = os.path.splitext(path)[0] + ".tmp.docx"
                build_docx(data, self.style.currentText(), tmp, self.photo)
                docx_to_pdf(tmp, path)
                os.unlink(tmp)
        except Exception as exc:
            handle_exception(self, _t("Resume Master"), exc)
            return
        if self.main is not None and hasattr(self.main, "open_path"):
            self.main.open_path(path)
        self.accept()


def open_dialog(main) -> None:
    ResumeDialog(main).exec()
