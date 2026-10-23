"""File Repair (WPS File Repair equivalent) for DOCX/XLSX/PPTX/PDF/CSV/TXT.

Strategy for OOXML packages (all are ZIP files of XML parts):
1. salvage ZIP entries even when the central directory is damaged or the file is truncated
   (scan for local file headers and inflate each entry, tolerating a broken last entry);
2. repair malformed XML with lxml's recovering parser;
3. re-create missing package parts ([Content_Types].xml, relationships) from a clean template;
4. validate by loading with python-docx / openpyxl / python-pptx;
5. if the package still cannot be loaded, rescue the content (text, cell values, slide text)
   into a brand-new document.
PDFs are rebuilt with MuPDF's repair, then pypdf (non-strict), then Ghostscript.
"""
from __future__ import annotations

import io
import logging
import os
import re
import struct
import zipfile
import zlib
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)
KIND_BY_EXT = {".docx": "docx", ".docm": "docx", ".dotx": "docx", ".wpt": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx",
               ".ett": "xlsx", ".pptx": "pptx", ".pptm": "pptx", ".dpt": "pptx", ".pdf": "pdf", ".csv": "text",
               ".txt": "text", ".md": "text", ".tsv": "text"}
MAIN_PART = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml", "pptx": "ppt/presentation.xml"}


@dataclass
class RepairReport:
    source: str
    output: str = ""
    kind: str = ""
    ok: bool = False
    actions: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    level: str = ""          # intact, repaired, salvaged

    def summary(self) -> str:
        head = {"intact": "No damage was found; a clean copy was written.",
                "repaired": "The file was repaired.", "salvaged": "The file was badly damaged; its content was rescued "
                "into a new document (some formatting may be lost)."}.get(self.level, "Repair failed.")
        lines = [head]
        if self.output:
            lines.append(f"Output: {self.output}")
        lines += [f"\u2022 {a}" for a in self.actions]
        lines += [f"\u2716 {e}" for e in self.errors]
        return "\n".join(lines)


# ============================================================================= zip salvage
def salvage_zip_entries(data: bytes) -> dict[str, bytes]:
    """Recover entries from a (possibly corrupt or truncated) ZIP by scanning local headers."""
    entries: dict[str, bytes] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                try:
                    entries[info.filename] = z.read(info)
                except Exception:
                    pass
            if entries:
                return entries
    except Exception:
        pass
    pos = 0
    sig = b"PK\x03\x04"
    while True:
        i = data.find(sig, pos)
        if i < 0 or i + 30 > len(data):
            break
        try:
            (_sig, _ver, flags, method, _t, _d, _crc, csize, _usize, nlen, xlen) = struct.unpack("<IHHHHHIIIHH", data[i:i + 30])
        except struct.error:
            break
        name = data[i + 30:i + 30 + nlen].decode("utf-8", "replace")
        start = i + 30 + nlen + xlen
        if flags & 0x08 or csize == 0:  # sizes in data descriptor: find the next header
            nxt = data.find(sig, start)
            nxt_c = data.find(b"PK\x01\x02", start)
            ends = [x for x in (nxt, nxt_c) if x > 0]
            end = min(ends) if ends else len(data)
            blob = data[start:end]
        else:
            blob = data[start:start + csize]
            end = start + csize
        try:
            if method == 8:
                d = zlib.decompressobj(-15)
                content = d.decompress(blob)
                try:
                    content += d.flush()
                except zlib.error:
                    pass
            elif method == 0:
                content = blob
            else:
                content = None
        except zlib.error:
            content = None
        if content is not None and name and not name.endswith("/"):
            entries[name] = content
        pos = max(end, i + 4)
    return entries


def repair_xml(data: bytes) -> tuple[bytes, bool]:
    from lxml import etree

    try:
        etree.fromstring(data)
        return data, False
    except etree.XMLSyntaxError:
        pass
    parser = etree.XMLParser(recover=True, huge_tree=True)
    root = etree.fromstring(data, parser)
    if root is None:
        raise ValueError("XML part is unrecoverable")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), True


def _template_package(kind: str) -> dict[str, bytes]:
    buf = io.BytesIO()
    if kind == "docx":
        import docx

        docx.Document().save(buf)
    elif kind == "xlsx":
        import openpyxl

        openpyxl.Workbook().save(buf)
    else:
        import pptx

        pptx.Presentation().save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        return {n: z.read(n) for n in z.namelist()}


def _validate(kind: str, path: str):
    if kind == "docx":
        import docx

        return docx.Document(path)
    if kind == "xlsx":
        import openpyxl

        return openpyxl.load_workbook(path)
    import pptx

    return pptx.Presentation(path)


def _write_zip(entries: dict[str, bytes], path: str):
    order = sorted(entries, key=lambda n: (n != "[Content_Types].xml", not n.startswith("_rels/"), n))
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n in order:
            z.writestr(n, entries[n])
    os.replace(tmp, path)


def repair_ooxml(path: str, out: str, kind: str, rep: RepairReport):
    with open(path, "rb") as fh:
        data = fh.read()
    intact = False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            bad = z.testzip()
            intact = bad is None
    except Exception:
        intact = False
    entries = salvage_zip_entries(data)
    if not intact:
        rep.actions.append(f"Recovered {len(entries)} package part(s) from the damaged ZIP container")
    if not entries:
        raise ValueError("No recoverable content found (the file is not an OOXML package)")
    fixed_xml = 0
    for name in list(entries):
        if name.endswith((".xml", ".rels")):
            try:
                entries[name], changed = repair_xml(entries[name])
                fixed_xml += changed
            except Exception as exc:
                rep.errors.append(f"{name}: {exc}")
                del entries[name]
    if fixed_xml:
        rep.actions.append(f"Repaired {fixed_xml} malformed XML part(s)")
    template = _template_package(kind)
    added = 0
    for req in ("[Content_Types].xml", "_rels/.rels"):
        if req not in entries:
            entries[req] = template[req]
            added += 1
    if MAIN_PART[kind] not in entries:
        rep.errors.append(f"The main part {MAIN_PART[kind]} is missing")
    if added:
        rep.actions.append(f"Re-created {added} missing package part(s)")
    # make sure every part has a content type registered
    from lxml import etree

    try:
        ct = etree.fromstring(entries["[Content_Types].xml"])
        ns = ct.nsmap.get(None, "http://schemas.openxmlformats.org/package/2006/content-types")
        defaults = {d.get("Extension") for d in ct.findall(f"{{{ns}}}Default")}
        tct = etree.fromstring(template["[Content_Types].xml"])
        for d in tct.findall(f"{{{ns}}}Default"):
            if d.get("Extension") not in defaults:
                ct.append(d)
        overrides = {o.get("PartName") for o in ct.findall(f"{{{ns}}}Override")}
        for o in tct.findall(f"{{{ns}}}Override"):
            if o.get("PartName") not in overrides and o.get("PartName", "/")[1:] in entries:
                ct.append(o)
        for o in list(ct.findall(f"{{{ns}}}Override")):
            if o.get("PartName", "/")[1:] not in entries:
                ct.remove(o)
        entries["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
    except Exception as exc:
        rep.errors.append(f"[Content_Types].xml: {exc}")
    _write_zip(entries, out)
    try:
        _validate(kind, out)
        rep.level = "intact" if intact and not fixed_xml and not added else "repaired"
        rep.ok = True
        return
    except Exception as exc:
        rep.actions.append(f"Package still not loadable ({exc.__class__.__name__}); rescuing content")
    salvage_content(kind, entries, out, rep)


def _texts(xml: bytes, tag_local: str) -> list[str]:
    from lxml import etree

    root = etree.fromstring(xml, etree.XMLParser(recover=True, huge_tree=True))
    if root is None:
        return []
    return [t.text or "" for t in root.iter() if isinstance(t.tag, str) and etree.QName(t).localname == tag_local]


def salvage_content(kind: str, entries: dict, out: str, rep: RepairReport):
    from lxml import etree

    if kind == "docx":
        import docx

        d = docx.Document()
        xml = entries.get("word/document.xml", b"")
        root = etree.fromstring(xml, etree.XMLParser(recover=True, huge_tree=True)) if xml else None
        count = 0
        if root is not None:
            for p in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                text = "".join(t.text or "" for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
                d.add_paragraph(text)
                count += 1
        for name, blob in entries.items():
            if name.startswith("word/media/") and name.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp")):
                try:
                    d.add_picture(io.BytesIO(blob))
                except Exception:
                    pass
        d.save(out)
        rep.actions.append(f"Rescued {count} paragraph(s) of text and the embedded pictures")
    elif kind == "xlsx":
        import openpyxl

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        shared = _texts(entries["xl/sharedStrings.xml"], "t") if "xl/sharedStrings.xml" in entries else []
        sheets = sorted(n for n in entries if re.match(r"xl/worksheets/sheet\d+\.xml$", n))
        cells = 0
        for sn in sheets:
            ws = wb.create_sheet(Path(sn).stem.capitalize())
            root = etree.fromstring(entries[sn], etree.XMLParser(recover=True, huge_tree=True))
            if root is None:
                continue
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            for c in root.iter(ns + "c"):
                ref = c.get("r")
                v = c.find(ns + "v")
                f = c.find(ns + "f")
                is_ = c.find(ns + "is")
                if not ref:
                    continue
                if f is not None and f.text:
                    ws[ref] = "=" + f.text
                elif c.get("t") == "s" and v is not None and v.text and v.text.isdigit() and int(v.text) < len(shared):
                    ws[ref] = shared[int(v.text)]
                elif is_ is not None:
                    ws[ref] = "".join(t.text or "" for t in is_.iter(ns + "t"))
                elif v is not None and v.text is not None:
                    try:
                        ws[ref] = float(v.text)
                    except ValueError:
                        ws[ref] = v.text
                else:
                    continue
                cells += 1
        if not wb.worksheets:
            wb.create_sheet("Sheet1")
        wb.save(out)
        rep.actions.append(f"Rescued {cells} cell(s) from {len(sheets)} sheet(s)")
    else:
        import pptx
        from pptx.util import Pt

        prs = pptx.Presentation()
        slides = sorted((n for n in entries if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.search(r"(\d+)", Path(n).stem).group(1)))
        for sn in slides:
            texts = [t for t in _texts(entries[sn], "t") if t.strip()]
            s = prs.slides.add_slide(prs.slide_layouts[1] if texts else prs.slide_layouts[6])
            if texts:
                s.shapes.title.text = texts[0]
                body = s.placeholders[1].text_frame
                body.text = texts[1] if len(texts) > 1 else ""
                for t in texts[2:]:
                    p = body.add_paragraph()
                    p.text = t
                    p.font.size = Pt(18)
        prs.save(out)
        rep.actions.append(f"Rescued the text of {len(slides)} slide(s)")
    rep.level = "salvaged"
    rep.ok = True


# ============================================================================= pdf / text
def repair_pdf(path: str, out: str, rep: RepairReport):
    import pymupdf

    try:
        doc = pymupdf.open(path)
        repaired = doc.is_repaired if hasattr(doc, "is_repaired") else False
        if doc.page_count == 0:
            raise ValueError("no pages")
        doc.save(out, garbage=4, deflate=True, clean=True)
        doc.close()
        rep.actions.append("Rebuilt the cross-reference table and cleaned the object structure (MuPDF)")
        rep.level = "repaired" if repaired else "intact"
        rep.ok = True
        return
    except Exception as exc:
        rep.actions.append(f"MuPDF could not open the file ({exc}); trying pypdf")
    try:
        from pypdf import PdfReader, PdfWriter

        r = PdfReader(path, strict=False)
        w = PdfWriter()
        ok_pages = 0
        for p in r.pages:
            try:
                w.add_page(p)
                ok_pages += 1
            except Exception:
                continue
        if ok_pages:
            with open(out, "wb") as fh:
                w.write(fh)
            rep.actions.append(f"Recovered {ok_pages} page(s) with pypdf (non-strict mode)")
            rep.level = "repaired"
            rep.ok = True
            return
    except Exception as exc:
        rep.actions.append(f"pypdf failed: {exc}")
    from app.pdf.ops import find_ghostscript

    gs = find_ghostscript()
    if gs:
        import subprocess

        r = subprocess.run([gs, "-o", out, "-sDEVICE=pdfwrite", "-dPDFSTOPONERROR=false", path], capture_output=True,
                           timeout=600)
        if r.returncode == 0 and os.path.exists(out) and os.path.getsize(out) > 0:
            rep.actions.append("Re-distilled the document with Ghostscript")
            rep.level = "repaired"
            rep.ok = True
            return
    raise ValueError("The PDF could not be repaired")


def repair_text(path: str, out: str, rep: RepairReport):
    raw = open(path, "rb").read()
    raw = raw.replace(b"\x00", b"")
    detected = "utf-8 (with replacements)"
    text = None
    for enc in ("utf-8-sig", "cp1256", "cp1252"):
        try:
            text = raw.decode(enc)
            detected = enc
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", "replace")
    with open(out, "w", encoding="utf-8", newline="") as fh:
        fh.write(text)
    rep.actions.append(f"Removed invalid bytes and re-encoded the text as UTF-8 (source encoding: {detected})")
    rep.level = "repaired"
    rep.ok = True


def repair_file(path: str, out: str | None = None) -> RepairReport:
    ext = Path(path).suffix.lower()
    kind = KIND_BY_EXT.get(ext)
    rep = RepairReport(path, kind=kind or "")
    if kind is None:
        rep.errors.append(f"Unsupported file type: {ext}")
        return rep
    if not out:
        p = Path(path)
        out = str(p.with_name(f"{p.stem} (repaired){p.suffix if kind != 'docx' or ext == '.docx' else '.docx'}"))
    rep.output = out
    try:
        if kind in MAIN_PART:
            repair_ooxml(path, out, kind, rep)
        elif kind == "pdf":
            repair_pdf(path, out, rep)
        else:
            repair_text(path, out, rep)
    except Exception as exc:
        rep.ok = False
        rep.errors.append(str(exc))
    return rep
