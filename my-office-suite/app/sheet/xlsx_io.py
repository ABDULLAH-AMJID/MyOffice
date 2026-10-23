"""Spreadsheet file I/O: XLSX (openpyxl), CSV/TSV, XLS (xlrd / LibreOffice), ODS (LibreOffice)."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import logging
import os
import tempfile
from pathlib import Path

from app.sheet.formula.functions import FUNCTIONS
from app.sheet.formula.parser import function_names, parse
from app.sheet.formula.values import XLError, datetime_to_serial
from app.sheet.model import (Book, Border, Cell, ChartSpec, CondFormat, SheetImage, Style, Validation, parse_input)

log = logging.getLogger(__name__)

THEME_COLORS = ["#ffffff", "#000000", "#e7e6e6", "#44546a", "#4472c4", "#ed7d31", "#a5a5a5", "#ffc000", "#5b9bd5",
                "#70ad47", "#0563c1", "#954f72"]
LARGE_FILE_BYTES = 40 * 1024 * 1024


class SpreadsheetError(Exception):
    pass


# ============================================================================= colour helpers
def _tint(hex_color: str, tint: float) -> str:
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    if tint < 0:
        r, g, b = (int(x * (1 + tint)) for x in (r, g, b))
    else:
        r, g, b = (int(x + (255 - x) * tint) for x in (r, g, b))
    return f"#{max(0, min(255, r)):02x}{max(0, min(255, g)):02x}{max(0, min(255, b)):02x}"


def color_from_openpyxl(color) -> str | None:
    if color is None:
        return None
    try:
        if color.type == "rgb" and isinstance(color.rgb, str):
            rgb = color.rgb
            if len(rgb) == 8:
                if rgb[:2] == "00" and rgb[2:] == "000000" and False:
                    return None
                return "#" + rgb[2:].lower()
            if len(rgb) == 6:
                return "#" + rgb.lower()
        if color.type == "theme" and color.theme is not None:
            base = THEME_COLORS[color.theme] if color.theme < len(THEME_COLORS) else "#000000"
            return _tint(base, color.tint or 0.0) if color.tint else base
        if color.type == "indexed" and color.indexed is not None:
            from openpyxl.styles.colors import COLOR_INDEX

            if color.indexed < len(COLOR_INDEX):
                rgb = COLOR_INDEX[color.indexed]
                if color.indexed in (64, 65):  # system foreground/background
                    return None
                return "#" + rgb[2:].lower()
    except Exception:
        return None
    return None


def to_argb(hex_color: str | None) -> str | None:
    if not hex_color:
        return None
    return "FF" + hex_color.lstrip("#").upper()[:6]


# ============================================================================= load
def _style_from_cell(c) -> Style:
    f = c.font
    fill = c.fill
    al = c.alignment
    bd = c.border
    fill_color = None
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        fill_color = color_from_openpyxl(fill.fgColor) or color_from_openpyxl(fill.start_color)
    def side(s):
        if s is None or not s.style:
            return None
        return Border({"hair": "thin", "mediumDashed": "dashed", "dashDot": "dashed", "slantDashDot": "dashed"}.get(
            s.style, s.style), color_from_openpyxl(s.color) or "#000000")
    color = color_from_openpyxl(f.color) if f is not None else None
    if color in ("#000000",):
        color = None
    return Style(
        font=f.name if f is not None and f.name not in (None, "Calibri") else None,
        size=float(f.sz) if f is not None and f.sz and float(f.sz) != 11.0 else None,
        bold=bool(f.b) if f is not None else False, italic=bool(f.i) if f is not None else False,
        underline=bool(f.u) if f is not None else False, strike=bool(f.strike) if f is not None else False,
        color=color, fill=fill_color if fill_color not in ("#ffffff",) or True else None,
        halign={"general": None, "centerContinuous": "center", "distributed": "justify", "fill": None}.get(
            al.horizontal, al.horizontal) if al is not None else None,
        valign={"center": "center", "top": "top", "bottom": None, "justify": None, "distributed": None}.get(
            al.vertical, None) if al is not None and al.vertical else None,
        wrap=bool(al.wrap_text) if al is not None else False,
        numfmt=c.number_format or "General", indent=int(al.indent or 0) if al is not None else 0,
        rotation=int(al.text_rotation or 0) if al is not None else 0,
        border_top=side(bd.top) if bd else None, border_bottom=side(bd.bottom) if bd else None,
        border_left=side(bd.left) if bd else None, border_right=side(bd.right) if bd else None,
        locked=bool(c.protection.locked) if c.protection is not None else True,
    )


def _cell_value(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return datetime_to_serial(v)
    if isinstance(v, _dt.time):
        return datetime_to_serial(v)
    if isinstance(v, _dt.timedelta):
        return v.total_seconds() / 86400.0
    if isinstance(v, str):
        if v in ("#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"):
            return XLError(v)
        return v
    return None if v is None else str(v)


def _formula_text(v) -> str | None:
    if isinstance(v, str) and v.startswith("=") and len(v) > 1:
        return v
    try:  # openpyxl ArrayFormula / DataTableFormula
        from openpyxl.worksheet.formula import ArrayFormula

        if isinstance(v, ArrayFormula):
            return v.text if v.text.startswith("=") else "=" + v.text
    except ImportError:
        pass
    return None


def load_xlsx(path: str, book: Book, progress=None, fast: bool = False, max_rows: int | None = None) -> dict:
    """Load an XLSX/XLSM/ETT file into ``book``.  Returns info {"truncated": bool, "unsupported": set}."""
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    info = {"truncated": False, "unsupported": set(), "fast": fast}
    try:
        wb = openpyxl.load_workbook(path, data_only=fast, read_only=fast, keep_links=False, rich_text=False)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise SpreadsheetError(f"Not a valid Excel workbook: {exc}") from exc
    except Exception as exc:  # zipfile.BadZipFile, lxml errors ...
        raise SpreadsheetError(f"The workbook is damaged or not an XLSX file ({exc.__class__.__name__}: {exc})") from exc
    try:
        import zipfile

        with zipfile.ZipFile(path) as z:
            if any(n.lower().endswith("vbaproject.bin") for n in z.namelist()):
                with open(path, "rb") as fh:
                    book.vba_source = fh.read()
                info["vba"] = True
    except (OSError, zipfile.BadZipFile):
        pass
    style_cache: dict = {}
    try:
        if not fast:
            for name, dn in list(wb.defined_names.items()):
                if dn.attr_text and not dn.attr_text.startswith("#"):
                    book.names[name.upper()] = "=" + dn.attr_text.lstrip("=")
        total_sheets = len(wb.worksheets)
        for si, ws in enumerate(wb.worksheets):
            sh = book.add_sheet(ws.title)
            dims = ws.max_row or 0
            limit = max_rows or (10 ** 9)
            for r_idx, row in enumerate(ws.iter_rows(), start=1):
                if r_idx > limit:
                    info["truncated"] = True
                    break
                if progress and r_idx % 2000 == 0:
                    progress(int(100 * (si + min(1.0, r_idx / max(1, dims))) / total_sheets),
                             f"{ws.title}: row {r_idx:,}")
                for c in row:
                    v = c.value
                    if v is None and fast:
                        continue
                    col = c.column if not fast else c.column
                    if col is None:
                        continue
                    sid = 0
                    if not fast:
                        key = c._style if hasattr(c, "_style") else None
                        k = tuple(key) if key is not None else None
                        sid = style_cache.get(k)
                        if sid is None:
                            sid = book.styles.intern(_style_from_cell(c))
                            style_cache[k] = sid
                    if v is None and sid == 0:
                        continue
                    formula = None if fast else _formula_text(v)
                    cell = Cell(None if formula else _cell_value(v), None, sid)
                    if formula:
                        cell.formula = formula
                    sh.cells[(c.row, col)] = cell
            if fast:
                continue
            _load_sheet_features(ws, sh)
        # defined names local to sheets (openpyxl 3.1)
        if not fast:
            for ws, sh in zip(wb.worksheets, book.sheets):
                for name, dn in getattr(ws, "defined_names", {}).items():
                    if dn.attr_text:
                        book.names.setdefault(name.upper(), "=" + dn.attr_text.lstrip("="))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    if not fast:
        book.engine.rebuild()
        book.engine.recalc_all()
        unsupported = _apply_cached_values(path, book)
        info["unsupported"] = unsupported
    for sh in book.sheets:
        sh.invalidate_bounds()
    return info


def _load_sheet_features(ws, sh):
    from openpyxl.utils import column_index_from_string

    fmt = ws.sheet_format
    if fmt is not None and fmt.defaultColWidth:
        sh.default_col_width = int(float(fmt.defaultColWidth) * 7 + 5)
    if fmt is not None and fmt.defaultRowHeight:
        sh.default_row_height = max(18, int(float(fmt.defaultRowHeight) * 96 / 72 + 4))
    for letter, dim in ws.column_dimensions.items():
        try:
            lo = column_index_from_string(letter) if dim.min is None else dim.min
            hi = lo if dim.max is None else dim.max
        except ValueError:
            continue
        for ci in range(lo, min(hi, lo + 200) + 1):
            if dim.width:
                sh.col_widths[ci] = max(8, int(float(dim.width) * 7 + 5))
            if dim.hidden:
                sh.hidden_cols.add(ci)
    for r, dim in ws.row_dimensions.items():
        if dim.height:
            sh.row_heights[r] = max(8, int(float(dim.height) * 96 / 72))
        if dim.hidden:
            sh.hidden_rows.add(r)
    for m in ws.merged_cells.ranges:
        sh.merges.append((m.min_row, m.min_col, m.max_row, m.max_col))
    for t in getattr(ws, "tables", {}).values():
        try:
            sh.tables.append({"name": t.displayName or t.name, "ref": t.ref,
                              "style": t.tableStyleInfo.name if t.tableStyleInfo else "TableStyleMedium2",
                              "header": (t.headerRowCount is None or t.headerRowCount > 0)})
        except Exception as exc:
            log.debug("table skipped: %s", exc)
    if ws.freeze_panes:
        from openpyxl.utils.cell import coordinate_from_string

        col, row = coordinate_from_string(ws.freeze_panes)
        sh.freeze = (row - 1, column_index_from_string(col) - 1)
    try:
        if ws.sheet_properties.tabColor is not None:
            sh.tab_color = color_from_openpyxl(ws.sheet_properties.tabColor)
    except Exception:
        pass
    try:
        sh.show_grid = bool(ws.sheet_view.showGridLines) if ws.sheet_view.showGridLines is not None else True
    except Exception:
        pass
    sh.protected = bool(ws.protection.sheet)
    for row in ws.iter_rows():
        for c in row:
            if c.comment is not None:
                sh.comments[(c.row, c.column)] = c.comment.text
            if c.hyperlink is not None and (c.hyperlink.target or c.hyperlink.location):
                sh.links[(c.row, c.column)] = c.hyperlink.target or ("#" + c.hyperlink.location)
    if ws.auto_filter is not None and ws.auto_filter.ref:
        from openpyxl.utils.cell import range_boundaries

        c1, r1, c2, r2 = range_boundaries(ws.auto_filter.ref)
        sh.autofilter = (r1, c1, r2, c2)
    try:
        for dv in ws.data_validations.dataValidation:
            for rng in dv.sqref.ranges:
                kind = {"list": "list", "whole": "whole", "decimal": "decimal", "date": "date",
                        "textLength": "textlength"}.get(dv.type)
                if not kind:
                    continue
                v1 = (dv.formula1 or "").strip('"')
                sh.validations.append(Validation((rng.min_row, rng.min_col, rng.max_row, rng.max_col), kind,
                                                 dv.operator or "between", v1, (dv.formula2 or ""),
                                                 bool(dv.allow_blank), dv.error or "", dv.prompt or ""))
    except Exception as exc:
        log.debug("validation import failed: %s", exc)
    try:
        for cf_range, rules in ws.conditional_formatting._cf_rules.items():
            for rng in cf_range.sqref.ranges:
                box = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
                for rule in rules:
                    cf = _cf_from_rule(rule, box)
                    if cf is not None:
                        sh.cond_formats.append(cf)
    except Exception as exc:
        log.debug("conditional format import failed: %s", exc)
    try:
        for ch in getattr(ws, "_charts", []):
            spec = _chart_from_openpyxl(ch, sh.name)
            if spec is not None:
                sh.charts.append(spec)
        for img in getattr(ws, "_images", []):
            data = img._data()
            fr = img.anchor._from if hasattr(img.anchor, "_from") else None
            anchor = (fr.row + 1, fr.col + 1) if fr is not None else (1, 1)
            sh.images.append(SheetImage(data, (img.format or "png").lower(), anchor,
                                        (int(img.width or 240), int(img.height or 160))))
    except Exception as exc:
        log.debug("chart/image import failed: %s", exc)


def _cf_from_rule(rule, box) -> CondFormat | None:
    dxf = rule.dxf
    fill = color_from_openpyxl(dxf.fill.fgColor) or color_from_openpyxl(dxf.fill.bgColor) if dxf and dxf.fill else None
    color = color_from_openpyxl(dxf.font.color) if dxf and dxf.font and dxf.font.color else None
    bold = bool(dxf.font.b) if dxf and dxf.font else False
    t = rule.type
    if t == "cellIs":
        op = {"greaterThan": ">", "lessThan": "<", "greaterThanOrEqual": ">=", "lessThanOrEqual": "<=",
              "equal": "=", "notEqual": "<>", "between": "between", "notBetween": "between"}.get(rule.operator, ">")
        f = rule.formula or []
        return CondFormat(box, "cell", op, (f[0] if f else "").strip('"'), (f[1] if len(f) > 1 else "").strip('"'),
                          fill, color, bold)
    if t == "containsText":
        return CondFormat(box, "text", "contains", rule.text or "", "", fill, color, bold)
    if t == "top10":
        return CondFormat(box, "bottom" if rule.bottom else "top", "", str(rule.rank or 10), "", fill, color, bold)
    if t in ("duplicateValues", "uniqueValues"):
        return CondFormat(box, "duplicate" if t == "duplicateValues" else "unique", "", "", "", fill, color, bold)
    if t == "colorScale" and rule.colorScale is not None:
        cols = [color_from_openpyxl(c) or "#ffffff" for c in rule.colorScale.color]
        return CondFormat(box, "scale3" if len(cols) == 3 else "scale2", colors=tuple(cols))
    if t == "dataBar" and rule.dataBar is not None:
        return CondFormat(box, "databar", fill=color_from_openpyxl(rule.dataBar.color) or "#638ec6")
    if t == "expression":
        f = rule.formula or []
        return CondFormat(box, "formula", "", f[0] if f else "", "", fill, color, bold)
    return None


def _chart_from_openpyxl(ch, sheet_name: str) -> ChartSpec | None:
    from openpyxl.utils.cell import range_boundaries

    name = ch.__class__.__name__
    kind = {"BarChart": "bar" if getattr(ch, "type", "col") == "bar" else "column", "BarChart3D": "column",
            "LineChart": "line", "LineChart3D": "line", "PieChart": "pie", "PieChart3D": "pie",
            "DoughnutChart": "doughnut", "AreaChart": "area", "ScatterChart": "scatter"}.get(name)
    if kind is None:
        return None
    boxes = []
    sheet = sheet_name
    for s in ch.series:
        for ref in (getattr(s.val, "numRef", None) if s.val is not None else None,
                    getattr(s.cat, "numRef", None) if s.cat is not None else None,
                    getattr(s.cat, "strRef", None) if s.cat is not None else None,
                    getattr(getattr(s, "xVal", None), "numRef", None), getattr(getattr(s, "yVal", None), "numRef", None),
                    getattr(getattr(s, "tx", None), "strRef", None)):
            if ref is not None and ref.f:
                f = ref.f
                if "!" in f:
                    sheet, _, f = f.rpartition("!")
                    sheet = sheet.strip("'")
                try:
                    boxes.append(range_boundaries(f.replace("$", "")))
                except ValueError:
                    pass
    if not boxes:
        return None
    c1 = min(b[0] for b in boxes)
    r1 = min(b[1] for b in boxes)
    c2 = max(b[2] for b in boxes)
    r2 = max(b[3] for b in boxes)
    from app.sheet.formula.parser import range_name

    title = ""
    try:
        if ch.title is not None and ch.title.tx is not None and ch.title.tx.rich is not None:
            title = "".join(r.t for p in ch.title.tx.rich.p for r in (p.r or []))
    except Exception:
        title = ""
    fr = ch.anchor._from if hasattr(ch.anchor, "_from") else None
    anchor = (fr.row + 1, fr.col + 1) if fr is not None else (1, c2 + 2)
    w = int((ch.width or 15) * 37.8)
    h = int((ch.height or 7.5) * 37.8)
    series_in_rows = (c2 - c1) > (r2 - r1) and kind not in ("scatter",)
    return ChartSpec(kind, range_name(r1, c1, r2, c2, sheet), title, series_in_rows, True, "", "", True, anchor,
                     (w, h))


def _apply_cached_values(path: str, book: Book) -> set:
    """Formulas using functions MyOffice does not implement keep Excel's cached results."""
    targets = []
    unsupported: set = set()
    for sh in book.sheets:
        for (r, c), cell in sh.cells.items():
            if cell.formula:
                try:
                    missing = {n for n in function_names(parse(cell.formula)) if n not in FUNCTIONS}
                except Exception:
                    missing = set()
                if missing:
                    unsupported |= missing
                    targets.append((sh, r, c))
    if not targets:
        return unsupported
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return unsupported
    try:
        by_sheet: dict = {}
        for sh, r, c in targets:
            by_sheet.setdefault(sh.name, []).append((sh, r, c))
        for name, cells in by_sheet.items():
            ws = wb[name]
            wanted = {(r, c) for _s, r, c in cells}
            maxr = max(r for r, _ in wanted)
            for row in ws.iter_rows(min_row=1, max_row=maxr):
                for c in row:
                    if getattr(c, "row", None) and (c.row, c.column) in wanted:
                        val = _cell_value(c.value)
                        cell = cells[0][0].cells[(c.row, c.column)]
                        cell.value = val

                        def const(ctx, _v=val):
                            return _v
                        cell.fn.fn = const
    finally:
        wb.close()
    return unsupported


# ============================================================================= save
_BORDER_STYLES = {"thin": "thin", "medium": "medium", "thick": "thick", "dashed": "dashed", "dotted": "dotted",
                  "double": "double"}


def _openpyxl_style(st: Style, cache: dict, sid: int):
    if sid in cache:
        return cache[sid]
    from openpyxl.styles import Alignment, Border as XBorder, Font, PatternFill, Protection, Side

    font = Font(name=st.font or "Calibri", sz=st.size or 11, b=st.bold, i=st.italic,
                u="single" if st.underline else None, strike=st.strike, color=to_argb(st.color))
    fill = PatternFill("solid", fgColor=to_argb(st.fill)) if st.fill else PatternFill()
    al = Alignment(horizontal=st.halign, vertical={"center": "center", "top": "top", "bottom": "bottom"}.get(st.valign),
                   wrap_text=st.wrap or None, indent=st.indent or 0, text_rotation=st.rotation or 0)

    def side(b: Border | None):
        return Side(style=_BORDER_STYLES.get(b.style, "thin"), color=to_argb(b.color)) if b else Side()
    border = XBorder(left=side(st.border_left), right=side(st.border_right), top=side(st.border_top),
                     bottom=side(st.border_bottom))
    res = (font, fill, al, border, Protection(locked=st.locked), st.numfmt or "General")
    cache[sid] = res
    return res


def _xl_value(v):
    if isinstance(v, XLError):
        return v.code
    return v


def save_xlsx(path: str, book: Book, progress=None) -> None:
    import io
    import zipfile

    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cache: dict = {}
    default_style = book.styles.get(0)
    total = max(1, sum(len(s.cells) for s in book.sheets))
    done = 0
    for sh in book.sheets:
        ws = wb.create_sheet(sh.name)
        for (r, c), cell in sorted(sh.cells.items()):
            done += 1
            if progress and done % 5000 == 0:
                progress(int(90 * done / total), f"{sh.name}: {done:,} cells")
            xc = ws.cell(row=r, column=c)
            if cell.formula:
                xc.value = cell.formula
            elif cell.value is not None:
                xc.value = _xl_value(cell.value)
            if cell.style:
                st = book.styles.get(cell.style)
                if st != default_style:
                    font, fill, al, border, prot, numfmt = _openpyxl_style(st, cache, cell.style)
                    xc.font, xc.fill, xc.alignment, xc.border, xc.protection = font, fill, al, border, prot
                    xc.number_format = numfmt
        for c, w in sh.col_widths.items():
            ws.column_dimensions[get_column_letter(c)].width = max(1.0, (w - 5) / 7)
        for c in sh.hidden_cols:
            ws.column_dimensions[get_column_letter(c)].hidden = True
        for r, h in sh.row_heights.items():
            ws.row_dimensions[r].height = h * 72 / 96
        for r in sh.hidden_rows | sh.filtered_rows:
            ws.row_dimensions[r].hidden = True
        if sh.default_col_width != 90:
            ws.sheet_format.defaultColWidth = (sh.default_col_width - 5) / 7
        for (r1, c1, r2, c2) in sh.merges:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        if sh.freeze != (0, 0):
            ws.freeze_panes = ws.cell(row=sh.freeze[0] + 1, column=sh.freeze[1] + 1)
        for (r, c), text in sh.comments.items():
            ws.cell(row=r, column=c).comment = Comment(text, "MyOffice")
        for (r, c), url in sh.links.items():
            xc = ws.cell(row=r, column=c)
            if url.startswith("#"):
                xc.hyperlink = openpyxl.worksheet.hyperlink.Hyperlink(ref=xc.coordinate, location=url[1:])
            else:
                xc.hyperlink = url
        if sh.tab_color:
            ws.sheet_properties.tabColor = sh.tab_color.lstrip("#").upper()
        ws.sheet_view.showGridLines = sh.show_grid
        if sh.protected:
            ws.protection.sheet = True
            if sh.protect_password:
                ws.protection.password = sh.protect_password
        if sh.autofilter:
            from app.sheet.formula.parser import range_name

            ws.auto_filter.ref = range_name(*sh.autofilter)
        _save_validations(ws, sh)
        _save_cond_formats(ws, sh)
        for spec in sh.charts:
            try:
                ch = build_openpyxl_chart(spec, book, wb)
                if ch is not None:
                    ws.add_chart(ch, ws.cell(row=spec.anchor[0], column=spec.anchor[1]).coordinate)
            except Exception as exc:
                log.warning("Chart export failed: %s", exc)
        for img in sh.images:
            try:
                from openpyxl.drawing.image import Image as XImage

                xi = XImage(io.BytesIO(img.data))
                xi.width, xi.height = img.size
                ws.add_image(xi, ws.cell(row=img.anchor[0], column=img.anchor[1]).coordinate)
            except Exception as exc:
                log.warning("Image export failed: %s", exc)
    for name, text in book.names.items():
        try:
            dn = DefinedName(name, attr_text=text.lstrip("="))
            wb.defined_names[name] = dn
        except Exception as exc:
            log.warning("Defined name %s skipped: %s", name, exc)
    for sh in book.sheets:
        ws = wb[sh.name] if sh.name in wb.sheetnames else None
        if ws is None:
            continue
        for t in getattr(sh, "tables", []):
            try:
                tab = Table(displayName=t["name"], ref=t["ref"], headerRowCount=1 if t.get("header", True) else 0)
                tab.tableStyleInfo = TableStyleInfo(name=t.get("style") or "TableStyleMedium2", showRowStripes=True,
                                                    showColumnStripes=False, showFirstColumn=False,
                                                    showLastColumn=False)
                ws.add_table(tab)
                if ws.auto_filter and ws.auto_filter.ref:  # a table brings its own filter; two would clash
                    ws.auto_filter.ref = None
            except Exception as exc:
                log.warning("Table %s not saved: %s", t.get("name"), exc)
    if book.vba_source and path.lower().endswith((".xlsm", ".xltm")):
        wb.vba_archive = zipfile.ZipFile(io.BytesIO(book.vba_source))
    if not wb.worksheets:
        wb.create_sheet("Sheet1")
    if progress:
        progress(95, "Writing file")
    d = os.path.dirname(os.path.abspath(path)) or "."
    fd, tmp = tempfile.mkstemp(dir=d, suffix=".xlsx.tmp")
    os.close(fd)
    try:
        wb.save(tmp)
        from app.services.fileutil import finish_atomic

        finish_atomic(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def _save_validations(ws, sh):
    from openpyxl.worksheet.datavalidation import DataValidation

    from app.sheet.formula.parser import range_name

    for v in sh.validations:
        kind = {"list": "list", "whole": "whole", "decimal": "decimal", "date": "date", "textlength": "textLength"}[v.kind]
        f1 = f'"{v.value1}"' if v.kind == "list" and not v.value1.startswith("=") else v.value1.lstrip("=")
        dv = DataValidation(type=kind, operator=None if v.kind == "list" else v.op, formula1=f1,
                            formula2=v.value2 or None, allow_blank=v.allow_blank, showErrorMessage=True,
                            error=v.error_message or None, prompt=v.prompt or None, showInputMessage=bool(v.prompt))
        dv.add(range_name(*v.range))
        ws.add_data_validation(dv)


def _save_cond_formats(ws, sh):
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle

    from app.sheet.formula.parser import range_name

    for cf in sh.cond_formats:
        ref = range_name(*cf.range)
        fill = PatternFill(start_color=to_argb(cf.fill), end_color=to_argb(cf.fill), fill_type="solid") if cf.fill else None
        font = Font(color=to_argb(cf.color), b=cf.bold or None) if (cf.color or cf.bold) else None
        try:
            if cf.kind == "cell":
                op = {">": "greaterThan", "<": "lessThan", ">=": "greaterThanOrEqual", "<=": "lessThanOrEqual",
                      "=": "equal", "<>": "notEqual", "between": "between"}[cf.op]
                vals = [cf.value1] + ([cf.value2] if cf.op == "between" else [])
                vals = [v if _is_number(v) else f'"{v}"' for v in vals]
                ws.conditional_formatting.add(ref, CellIsRule(operator=op, formula=vals, fill=fill, font=font))
            elif cf.kind == "text":
                first = range_name(cf.range[0], cf.range[1], cf.range[0], cf.range[1])
                ws.conditional_formatting.add(ref, FormulaRule(
                    formula=[f'NOT(ISERROR(SEARCH("{cf.value1}",{first})))'], fill=fill, font=font))
            elif cf.kind in ("top", "bottom"):
                rule = Rule(type="top10", rank=int(float(cf.value1 or 10)), bottom=cf.kind == "bottom",
                            dxf=DifferentialStyle(fill=fill, font=font))
                ws.conditional_formatting.add(ref, rule)
            elif cf.kind in ("duplicate", "unique"):
                rule = Rule(type="duplicateValues" if cf.kind == "duplicate" else "uniqueValues",
                            dxf=DifferentialStyle(fill=fill, font=font))
                ws.conditional_formatting.add(ref, rule)
            elif cf.kind == "scale2":
                ws.conditional_formatting.add(ref, ColorScaleRule(start_type="min", start_color=to_argb(cf.colors[0]),
                                                                  end_type="max", end_color=to_argb(cf.colors[-1])))
            elif cf.kind == "scale3":
                ws.conditional_formatting.add(ref, ColorScaleRule(
                    start_type="min", start_color=to_argb(cf.colors[0]), mid_type="percentile", mid_value=50,
                    mid_color=to_argb(cf.colors[1]), end_type="max", end_color=to_argb(cf.colors[2])))
            elif cf.kind == "databar":
                ws.conditional_formatting.add(ref, DataBarRule(start_type="min", end_type="max",
                                                               color=to_argb(cf.fill or "#638ec6")))
            elif cf.kind == "formula":
                ws.conditional_formatting.add(ref, FormulaRule(formula=[cf.value1.lstrip("=")], fill=fill, font=font))
        except Exception as exc:
            log.warning("Conditional format export failed: %s", exc)


def _is_number(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


def build_openpyxl_chart(spec: ChartSpec, book: Book, wb):
    """Create a native Excel chart for a ChartSpec."""
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series

    from app.sheet.formula.parser import parse_ref_text

    ref = parse_ref_text(spec.data)
    sh = book.sheet_by_name(ref.sheet) if ref.sheet else None
    if sh is None:
        return None
    ws = wb[sh.name]
    r1, c1, r2, c2 = ref.r1, ref.c1, ref.r2, ref.c2
    kind = spec.kind
    if kind in ("column", "bar"):
        ch = BarChart()
        ch.type = "col" if kind == "column" else "bar"
    elif kind == "line":
        ch = LineChart()
    elif kind == "area":
        ch = AreaChart()
    elif kind == "pie":
        ch = PieChart()
    elif kind == "doughnut":
        ch = DoughnutChart()
    else:
        ch = ScatterChart()
        ch.style = 13
    hdr = 1 if spec.has_header else 0
    if kind == "scatter":
        xref = Reference(ws, min_col=c1, min_row=r1 + hdr, max_row=r2)
        for c in range(c1 + 1, c2 + 1):
            yref = Reference(ws, min_col=c, min_row=r1, max_row=r2)
            s = Series(yref, xref, title_from_data=bool(hdr))
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
            ch.series.append(s)
    elif spec.series_in_rows:
        data = Reference(ws, min_col=c1 + 1 if hdr else c1, min_row=r1 + hdr, max_col=c2, max_row=r2)
        cats = Reference(ws, min_col=c1 + 1, min_row=r1, max_col=c2, max_row=r1) if hdr else None
        ch.add_data(data, from_rows=True, titles_from_data=bool(hdr))
        if cats is not None:
            ch.set_categories(cats)
    else:
        data = Reference(ws, min_col=c1 + 1 if (c2 > c1) else c1, min_row=r1, max_col=c2, max_row=r2)
        ch.add_data(data, titles_from_data=bool(hdr))
        if c2 > c1:
            ch.set_categories(Reference(ws, min_col=c1, min_row=r1 + hdr, max_row=r2))
    if spec.title:
        ch.title = spec.title
    if kind not in ("pie", "doughnut"):
        try:
            if spec.x_title:
                ch.x_axis.title = spec.x_title
            if spec.y_title:
                ch.y_axis.title = spec.y_title
            ch.x_axis.delete = False
            ch.y_axis.delete = False
        except Exception:
            pass
    if not spec.legend:
        ch.legend = None
    ch.width = spec.size[0] / 37.8
    ch.height = spec.size[1] / 37.8
    return ch


# ============================================================================= CSV / XLS / ODS
def sniff_encoding(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    try:
        raw.decode("cp1256")  # Urdu/Arabic Windows code page
        if any(0xC1 <= b <= 0xFE for b in raw[:4000]):
            return "cp1256"
    except UnicodeDecodeError:
        pass
    return "cp1252"


def load_csv(path: str, book: Book, delimiter: str | None = None, encoding: str | None = None) -> dict:
    with open(path, "rb") as fh:
        raw = fh.read()
    enc = encoding or sniff_encoding(raw)
    text = raw.decode(enc, errors="replace")
    if delimiter is None:
        if path.lower().endswith(".tsv"):
            delimiter = "\t"
        else:
            try:
                delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
    sh = book.add_sheet(Path(path).stem[:31] or "Sheet1")
    styles: dict = {}
    for r, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        for c, val in enumerate(row, start=1):
            if val == "":
                continue
            value, formula, fmt = parse_input(val)
            sid = 0
            if fmt:
                sid = styles.get(fmt)
                if sid is None:
                    sid = styles[fmt] = book.styles.intern(Style(numfmt=fmt))
            cell = Cell(value, None, sid)
            if formula:
                cell.formula = formula
            sh.cells[(r, c)] = cell
    sh.invalidate_bounds()
    book.engine.rebuild()
    book.engine.recalc_all()
    return {"encoding": enc, "delimiter": delimiter}


def save_csv(path: str, book: Book, sheet=None, delimiter: str = ",", encoding: str = "utf-8-sig") -> None:
    sh = sheet or book.sheets[0]
    r1, c1, r2, c2 = sh.used_range()
    tmp = path + ".tmp"
    with open(tmp, "w", newline="", encoding=encoding) as fh:
        w = csv.writer(fh, delimiter=delimiter)
        for r in range(1, r2 + 1):
            w.writerow([sh.display(r, c)[0] for c in range(1, c2 + 1)])
    os.replace(tmp, path)


def load_xls(path: str, book: Book) -> dict:
    """Legacy .xls: LibreOffice (if installed) also keeps the formulas; otherwise xlrd reads values and
    formatting."""
    from app.services.office_convert import available, convert

    if available():
        out = convert(path, "xlsx")
        return load_xlsx(out, book)
    from app.sheet.features import load_xls_formatted

    return load_xls_formatted(path, book)


def load_any(path: str, book: Book, progress=None, fast: bool = False, max_rows: int | None = None) -> dict:
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm", ".ett") :
        import zipfile

        if not zipfile.is_zipfile(path):
            if ext == ".ett":
                return load_xls(path, book)
            raise SpreadsheetError("The file is not a valid XLSX workbook (it is not a ZIP package).")
        return load_xlsx(path, book, progress, fast, max_rows)
    if ext in (".csv", ".tsv", ".txt"):
        return load_csv(path, book)
    if ext in (".xls", ".et"):
        return load_xls(path, book)
    if ext in (".ods", ".fods"):
        from app.services.office_convert import convert

        return load_xlsx(convert(path, "xlsx"), book, progress, fast, max_rows)
    raise SpreadsheetError(f"Unsupported spreadsheet format: {ext}")
