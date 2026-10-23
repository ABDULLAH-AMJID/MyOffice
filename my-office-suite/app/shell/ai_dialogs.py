"""MyOffice AI - the dialogs shared by Writer, Spreadsheets, Presentation and PDF.

All network work runs in worker threads; answers stream into the dialogs where it helps (writing, chat).
"""
from __future__ import annotations

import html
import logging
import os

from PySide6.QtCore import QThread, Qt, QUrl, Signal
from PySide6.QtGui import QDesktopServices, QFont, QImage, QPainter, QTextDocument
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QCheckBox, QComboBox, QDialog, QDialogButtonBox,
                               QDoubleSpinBox, QFormLayout, QHBoxLayout, QHeaderView, QLabel, QLineEdit,
                               QListWidget, QListWidgetItem, QMessageBox, QPlainTextEdit, QPushButton, QSpinBox,
                               QTableWidget, QTableWidgetItem, QTextBrowser, QVBoxLayout, QWidget)

from app.services import ai
from app.services.i18n import _t

log = logging.getLogger(__name__)


# ============================================================================ worker
class AIWorker(QThread):
    """Runs ``fn(on_delta, is_cancelled) -> str`` in a thread; streams ``delta`` and ends with ``done``/``failed``."""
    delta = Signal(str)
    done = Signal(str)
    failed = Signal(str)

    def __init__(self, fn, parent=None):
        super().__init__(parent)
        self._fn = fn
        self._cancel = False

    def cancel(self):
        self._cancel = True

    def run(self):
        try:
            result = self._fn(self.delta.emit, lambda: self._cancel)
        except ai.AICancelled:
            self.failed.emit("")
        except Exception as exc:  # shown in the dialog
            log.info("AI request failed: %s", exc)
            self.failed.emit(str(exc))
        else:
            self.done.emit(result or "")


def ensure_ai(parent) -> bool:
    """True when an AI provider is set up; otherwise offers AI Settings first."""
    if ai.configured():
        return True
    box = QMessageBox(parent)
    box.setWindowTitle(_t("MyOffice AI"))
    box.setIcon(QMessageBox.Information)
    box.setText(_t("MyOffice AI needs an AI provider."))
    box.setInformativeText(_t("Use a free local AI server on this PC (Ollama or LM Studio - works offline, nothing "
                              "leaves your computer) or an online provider such as OpenAI, Google Gemini or "
                              "Anthropic with your own API key."))
    setup = box.addButton(_t("Set Up AI\u2026"), QMessageBox.AcceptRole)
    box.addButton(QMessageBox.Cancel)
    box.exec()
    if box.clickedButton() is not setup:
        return False
    return AISettingsDialog(parent).exec() == QDialog.Accepted and ai.configured()


# ============================================================================ settings
class AISettingsDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle(_t("AI Settings"))
        self.setMinimumWidth(560)
        s = ai.settings()
        v = QVBoxLayout(self)
        intro = QLabel(_t("Choose where MyOffice AI runs. A local server (Ollama, LM Studio) works without internet "
                          "and keeps your documents on this PC. Online providers need your own API key; the text "
                          "you send is processed by that provider."))
        intro.setWordWrap(True)
        intro.setStyleSheet("color:#5f6368")
        v.addWidget(intro)
        form = QFormLayout()
        self.provider = QComboBox()
        for key, (label, *_rest) in ai.PRESETS.items():
            self.provider.addItem(_t(label), key)
        self.provider.setCurrentIndex(max(0, self.provider.findData(s.provider)))
        self.provider.currentIndexChanged.connect(self._preset)
        form.addRow(_t("Provider:"), self.provider)
        self.url = QLineEdit(s.base_url)
        form.addRow(_t("Server address:"), self.url)
        self.key = QLineEdit(s.api_key)
        self.key.setEchoMode(QLineEdit.Password)
        self.key.setPlaceholderText(_t("not needed for local servers"))
        form.addRow(_t("API key:"), self.key)
        mrow = QHBoxLayout()
        self.model = QComboBox()
        self.model.setEditable(True)
        self.model.setCurrentText(s.model)
        self.model.setMinimumWidth(260)
        fetch = QPushButton(_t("Load Models"))
        fetch.clicked.connect(self._fetch)
        mrow.addWidget(self.model, 1)
        mrow.addWidget(fetch)
        form.addRow(_t("Model:"), mrow)
        self.temp = QDoubleSpinBox()
        self.temp.setRange(0.0, 1.5)
        self.temp.setSingleStep(0.1)
        self.temp.setValue(s.temperature)
        form.addRow(_t("Creativity (temperature):"), self.temp)
        self.timeout = QSpinBox()
        self.timeout.setRange(10, 900)
        self.timeout.setSuffix(" s")
        self.timeout.setValue(int(s.timeout))
        form.addRow(_t("Time limit:"), self.timeout)
        v.addLayout(form)
        self.status = QLabel("")
        self.status.setWordWrap(True)
        v.addWidget(self.status)
        row = QHBoxLayout()
        test = QPushButton(_t("Test Connection"))
        test.clicked.connect(self._test)
        help_btn = QPushButton(_t("How to Set Up Offline AI"))
        help_btn.clicked.connect(lambda: QDesktopServices.openUrl(QUrl("https://ollama.com/download")))
        row.addWidget(test)
        row.addWidget(help_btn)
        row.addStretch(1)
        v.addLayout(row)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.accepted.connect(self._save)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        self._worker = None

    def current(self) -> ai.AISettings:
        return ai.AISettings(provider=self.provider.currentData(), base_url=self.url.text().strip().rstrip("/"),
                             model=self.model.currentText().strip(), api_key=self.key.text().strip(),
                             timeout=float(self.timeout.value()), temperature=float(self.temp.value()))

    def _preset(self):
        label, url, model, _needs = ai.PRESETS[self.provider.currentData()]
        self.url.setText(url)
        self.model.setCurrentText(model)

    def _run(self, fn, ok):
        if self._worker is not None and self._worker.isRunning():
            return
        self.status.setText(_t("Contacting the AI provider\u2026"))
        self._worker = AIWorker(lambda _d, _c: fn(), self)
        self._worker.done.connect(ok)
        self._worker.failed.connect(lambda msg: self.status.setText("<span style='color:#c5221f'>%s</span>"
                                                                   % html.escape(msg)))
        self._worker.start()

    def _fetch(self):
        s = self.current()

        def ok(text):
            names = [n for n in text.split("\n") if n]
            cur = self.model.currentText()
            self.model.clear()
            self.model.addItems(names)
            self.model.setCurrentText(cur if cur in names or not names else names[0])
            self.status.setText(_t("{0} models found.").format(len(names)))
        self._run(lambda: "\n".join(ai.list_models(s)), ok)

    def _test(self):
        s = self.current()
        self._run(lambda: ai.test_connection(s),
                  lambda text: self.status.setText("<span style='color:#188038'>%s</span> %s" % (
                      _t("Connected. The model answered:"), html.escape(text[:80]))))

    def _save(self):
        s = self.current()
        if not s.base_url or not s.model:
            self.status.setText(_t("Enter the server address and the model."))
            return
        ai.save_settings(s)
        self.accept()


# ============================================================================ writing assistant
class AIWriterDialog(QDialog):
    """Write, continue, rewrite, shorten, expand, change the tone, summarize - with live output."""
    ACTIONS = [("improve", "Improve writing"), ("continue", "Continue writing"), ("shorten", "Make shorter"),
               ("expand", "Make longer"), ("formal", "More formal"), ("casual", "More casual"),
               ("summarize", "Summarize"), ("bullets", "Bullet points"), ("write", "Write new text"),
               ("custom", "Custom instruction")]

    def __init__(self, parent, text: str, insert_cb, title: str = ""):
        super().__init__(parent)
        self.setWindowTitle(title or _t("AI Writer"))
        self.resize(720, 560)
        self.source = text
        self.insert_cb = insert_cb
        v = QVBoxLayout(self)
        top = QHBoxLayout()
        self.action = QComboBox()
        for key, label in self.ACTIONS:
            self.action.addItem(_t(label), key)
        if not text.strip():
            self.action.setCurrentIndex(self.action.findData("write"))
        self.lang = QComboBox()
        self.lang.addItem(_t("Same language"), "")
        for name in ai.LANGUAGES:
            self.lang.addItem(_t(name), name)
        top.addWidget(QLabel(_t("Task:")))
        top.addWidget(self.action, 1)
        top.addWidget(QLabel(_t("Language:")))
        top.addWidget(self.lang)
        v.addLayout(top)
        self.instruction = QLineEdit()
        self.instruction.setPlaceholderText(_t("Topic or instruction, e.g. \"a polite reminder about the unpaid "
                                               "invoice\""))
        v.addWidget(self.instruction)
        if text.strip():
            src = QPlainTextEdit(text)
            src.setReadOnly(True)
            src.setMaximumHeight(120)
            v.addWidget(QLabel(_t("Selected text:")))
            v.addWidget(src)
        v.addWidget(QLabel(_t("Result:")))
        self.out = QPlainTextEdit()
        v.addWidget(self.out, 1)
        self.status = QLabel("")
        self.status.setStyleSheet("color:#5f6368")
        v.addWidget(self.status)
        row = QHBoxLayout()
        self.go = QPushButton(_t("Generate"))
        self.go.setObjectName("Primary")
        self.go.clicked.connect(self.generate)
        self.stop = QPushButton(_t("Stop"))
        self.stop.setEnabled(False)
        self.stop.clicked.connect(self._stop)
        self.replace = QPushButton(_t("Replace Selection") if text.strip() else _t("Insert"))
        self.replace.clicked.connect(lambda: self._use("replace"))
        self.below = QPushButton(_t("Insert Below"))
        self.below.clicked.connect(lambda: self._use("below"))
        self.below.setVisible(bool(text.strip()))
        copy = QPushButton(_t("Copy"))
        copy.clicked.connect(lambda: QApplication.clipboard().setText(self.out.toPlainText()))
        close = QPushButton(_t("Close"))
        close.clicked.connect(self.reject)
        for b in (self.go, self.stop):
            row.addWidget(b)
        row.addStretch(1)
        for b in (copy, self.below, self.replace, close):
            row.addWidget(b)
        v.addLayout(row)
        self._worker = None

    def generate(self):
        action = self.action.currentData()
        instr = self.instruction.text().strip()
        if action in ("write", "custom") and not instr:
            self.status.setText(_t("Enter a topic or an instruction."))
            return
        msgs = ai.writer_messages(action, "" if action == "write" else self.source, instr, self.lang.currentData())
        self.out.clear()
        self.status.setText(_t("Writing\u2026"))
        self.go.setEnabled(False)
        self.stop.setEnabled(True)
        self._worker = AIWorker(lambda on_delta, cancelled: ai.chat(msgs, on_delta=on_delta, is_cancelled=cancelled),
                                self)
        self._worker.delta.connect(self._append)
        self._worker.done.connect(self._done)
        self._worker.failed.connect(self._failed)
        self._worker.start()

    def _append(self, text):
        c = self.out.textCursor()
        c.movePosition(c.MoveOperation.End)
        c.insertText(text)
        self.out.setTextCursor(c)

    def _done(self, text):
        final = ai.clean_text(text)
        if final and final != self.out.toPlainText():
            self.out.setPlainText(final)
        self.status.setText(_t("Done. Edit the result if you like, then insert it."))
        self.go.setEnabled(True)
        self.stop.setEnabled(False)

    def _failed(self, msg):
        self.status.setText(msg or _t("Stopped."))
        self.go.setEnabled(True)
        self.stop.setEnabled(False)

    def _stop(self):
        if self._worker:
            self._worker.cancel()

    def _use(self, mode):
        text = self.out.toPlainText().strip()
        if text:
            self.insert_cb(text, mode)
            self.accept()

    def reject(self):
        if self._worker and self._worker.isRunning():
            self._worker.cancel()
            self._worker.wait(3000)
        super().reject()


# ============================================================================ grammar check
class GrammarDialog(QDialog):
    """Spelling (offline) and grammar (AI) suggestions for a text; accepted ones are applied together."""

    def __init__(self, parent, text: str, apply_cb, variant: str = "American English", spell_lang: str = "en_US"):
        super().__init__(parent)
        self.setWindowTitle(_t("Spelling & Grammar"))
        self.resize(680, 520)
        self.text = text
        self.apply_cb = apply_cb
        self.variant = variant
        self.spell_lang = spell_lang
        self.items: list[ai.Correction] = []
        v = QVBoxLayout(self)
        self.info = QLabel("")
        self.info.setWordWrap(True)
        v.addWidget(self.info)
        self.list = QListWidget()
        self.list.setSelectionMode(QAbstractItemView.NoSelection)
        v.addWidget(self.list, 1)
        row = QHBoxLayout()
        self.ai_btn = QPushButton(_t("Check Grammar with AI"))
        self.ai_btn.clicked.connect(self.run_ai)
        all_btn = QPushButton(_t("Select All"))
        all_btn.clicked.connect(lambda: self._check_all(True))
        none_btn = QPushButton(_t("Select None"))
        none_btn.clicked.connect(lambda: self._check_all(False))
        self.apply_btn = QPushButton(_t("Apply Selected"))
        self.apply_btn.setObjectName("Primary")
        self.apply_btn.clicked.connect(self._apply)
        close = QPushButton(_t("Close"))
        close.clicked.connect(self.reject)
        for b in (self.ai_btn, all_btn, none_btn):
            row.addWidget(b)
        row.addStretch(1)
        row.addWidget(self.apply_btn)
        row.addWidget(close)
        v.addLayout(row)
        self._worker = None
        self.run_spelling()

    def run_spelling(self):
        from app.services.spell import speller

        sp = speller(self.spell_lang)
        found = []
        for start, end, word in sp.errors(self.text):
            sugg = sp.suggest(word, 1)
            if sugg:
                found.append(ai.Correction(start, end, word, sugg[0]))
        self._show(found, _t("Spelling: {0} suggestions (offline).").format(len(found)))

    def run_ai(self):
        if not ensure_ai(self):
            return
        msgs = ai.grammar_messages(self.text, self.variant)
        self.info.setText(_t("Checking grammar with AI\u2026"))
        self.ai_btn.setEnabled(False)
        self._worker = AIWorker(lambda _d, cancelled: ai.chat(msgs, temperature=0, is_cancelled=cancelled), self)
        self._worker.done.connect(self._ai_done)
        self._worker.failed.connect(lambda m: (self.info.setText(m or _t("Stopped.")), self.ai_btn.setEnabled(True)))
        self._worker.start()

    def _ai_done(self, answer):
        self.ai_btn.setEnabled(True)
        found = ai.corrections(self.text, answer)
        self._show(found, _t("Spelling, grammar and punctuation: {0} suggestions.").format(len(found)))

    def _show(self, found, info):
        self.items = found
        self.list.clear()
        for c in found:
            ctx_a = self.text[max(0, c.start - 30):c.start].replace("\n", " ")
            ctx_b = self.text[c.end:c.end + 30].replace("\n", " ")
            empty = "\u2205"  # (no backslashes inside f-string expressions: Python 3.11)
            it = QListWidgetItem(f"\u2026{ctx_a}[{c.original or empty} \u2192 {c.suggestion or empty}]{ctx_b}\u2026")
            it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Checked)
            self.list.addItem(it)
        self.info.setText(info if found else _t("No problems found."))
        self.apply_btn.setEnabled(bool(found))

    def _check_all(self, on):
        for i in range(self.list.count()):
            self.list.item(i).setCheckState(Qt.Checked if on else Qt.Unchecked)

    def _apply(self):
        """``apply_cb`` receives the accepted corrections (offsets into the checked text), so the caller can apply
        them in place and keep the formatting of everything else."""
        chosen = [c for i, c in enumerate(self.items) if self.list.item(i).checkState() == Qt.Checked]
        if chosen:
            self.apply_cb(chosen)
        self.accept()


# ============================================================================ parallel translate
def export_translation(parent, rows: list[tuple[str, str]], target: str, base: str = "translation") -> None:
    from app.shell.dialogs import get_save_file, show_info

    path = get_save_file(parent, _t("Export Translation"),
                         "Word (*.docx);;Excel (*.xlsx);;PDF (*.pdf);;Text (*.txt);;PNG image (*.png)",
                         base + ".docx", "docx")
    if not path:
        return
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        import openpyxl
        from openpyxl.styles import Alignment, Font

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Translation"
        ws.append([_t("Original"), target])
        for a, b in rows:
            ws.append([a, b])
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.column_dimensions["A"].width = ws.column_dimensions["B"].width = 70
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(path)
    elif ext == ".txt":
        with open(path, "w", encoding="utf-8") as f:
            for a, b in rows:
                f.write(f"{a}\n{b}\n\n")
    elif ext in (".pdf", ".png"):
        doc = QTextDocument()
        doc.setDefaultFont(QFont("Segoe UI", 10))
        cells = "".join(f"<tr><td width='50%'>{html.escape(a)}</td><td width='50%'>{html.escape(b)}</td></tr>"
                        for a, b in rows)
        doc.setHtml(f"<table border='1' cellspacing='0' cellpadding='6' width='100%'><tr><th>{_t('Original')}</th>"
                    f"<th>{html.escape(target)}</th></tr>{cells}</table>")
        if ext == ".pdf":
            from PySide6.QtGui import QPageSize, QPdfWriter

            w = QPdfWriter(path)
            w.setPageSize(QPageSize(QPageSize.A4))
            w.setResolution(96)
            doc.print_(w)
        else:
            doc.setTextWidth(1100)
            img = QImage(int(doc.size().width()), int(doc.size().height()), QImage.Format_ARGB32)
            img.fill(Qt.white)
            p = QPainter(img)
            doc.drawContents(p)
            p.end()
            img.save(path)
    else:
        import docx

        d = docx.Document()
        t = d.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = _t("Original"), target
        for a, b in rows:
            r = t.add_row().cells
            r[0].text, r[1].text = a, b
        d.save(path)
    show_info(parent, _t("Export Translation"), _t("Saved:\n{0}").format(path))


class TranslateDialog(QDialog):
    """Parallel Translate: the original and the translation side by side, paragraph by paragraph."""

    def __init__(self, parent, paragraphs: list[str], insert_cb=None, title: str = ""):
        super().__init__(parent)
        self.setWindowTitle(title or _t("Parallel Translate"))
        self.resize(980, 640)
        self.paragraphs = [p for p in paragraphs]
        self.insert_cb = insert_cb
        self.result: list[str] = []
        v = QVBoxLayout(self)
        top = QHBoxLayout()
        top.addWidget(QLabel(_t("Translate into:")))
        self.target = QComboBox()
        for name in ai.LANGUAGES:
            self.target.addItem(_t(name), name)
        self.target.setCurrentIndex(max(0, self.target.findData("Urdu")))
        top.addWidget(self.target)
        self.go = QPushButton(_t("Translate"))
        self.go.setObjectName("Primary")
        self.go.clicked.connect(self.translate)
        top.addWidget(self.go)
        top.addStretch(1)
        v.addLayout(top)
        self.table = QTableWidget(len(self.paragraphs), 2)
        self.table.setHorizontalHeaderLabels([_t("Original"), _t("Translation")])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setWordWrap(True)
        self.table.verticalHeader().setVisible(False)
        for i, p in enumerate(self.paragraphs):
            it = QTableWidgetItem(p)
            it.setFlags(it.flags() & ~Qt.ItemIsEditable)
            self.table.setItem(i, 0, it)
            self.table.setItem(i, 1, QTableWidgetItem(""))
        self.table.resizeRowsToContents()
        v.addWidget(self.table, 1)
        row = QHBoxLayout()
        exp = QPushButton(_t("Export\u2026"))
        exp.clicked.connect(self._export)
        row.addWidget(exp)
        row.addStretch(1)
        if insert_cb:
            ins = QPushButton(_t("Insert Translation Below"))
            ins.clicked.connect(lambda: self._use("below"))
            rep = QPushButton(_t("Replace with Translation"))
            rep.clicked.connect(lambda: self._use("replace"))
            row.addWidget(ins)
            row.addWidget(rep)
        close = QPushButton(_t("Close"))
        close.clicked.connect(self.reject)
        row.addWidget(close)
        v.addLayout(row)

    def translate(self):
        if not ensure_ai(self):
            return
        from app.shell.tasks import TaskError, run_task
        from app.shell.dialogs import show_error

        target = self.target.currentData()
        try:
            self.result = run_task(self, _t("Translating\u2026"), lambda progress=None, is_cancelled=None:
                                   ai.translate_batches(self.paragraphs, target, progress), cancellable=False)
        except TaskError as exc:
            show_error(self, _t("Parallel Translate"), str(exc.original), exc.tb)
            return
        for i, t in enumerate(self.result):
            self.table.item(i, 1).setText(t)
        self.table.resizeRowsToContents()

    def rows(self) -> list[tuple[str, str]]:
        return [(self.table.item(i, 0).text(), self.table.item(i, 1).text()) for i in range(self.table.rowCount())]

    def _export(self):
        export_translation(self, self.rows(), self.target.currentData())

    def _use(self, mode):
        translated = [b for _a, b in self.rows()]
        if any(t.strip() for t in translated):
            self.insert_cb(translated, mode)
            self.accept()


# ============================================================================ Chat PDF
class ChatPdfPanel(QWidget):
    """Summary and question-answering for the open PDF, with clickable page citations.

    Only the passages that answer the question are sent to the AI (found with BM25 on this PC); with a local AI
    server the document never leaves the computer."""

    def __init__(self, parent, pages_provider, goto_page, ocr_cb=None):
        super().__init__(parent)
        self.pages_provider = pages_provider
        self.goto_page = goto_page
        self.ocr_cb = ocr_cb
        self.history: list[tuple[str, str]] = []
        self._index = None
        self._pages = None
        v = QVBoxLayout(self)
        v.setContentsMargins(6, 6, 6, 6)
        head = QLabel("<b>%s</b>" % _t("Chat with PDF"))
        v.addWidget(head)
        row = QHBoxLayout()
        self.summary_btn = QPushButton(_t("Summarize"))
        self.summary_btn.clicked.connect(self.summarize)
        self.lang = QComboBox()
        self.lang.addItem(_t("Auto language"), "")
        for name in ai.LANGUAGES:
            self.lang.addItem(_t(name), name)
        row.addWidget(self.summary_btn)
        row.addWidget(self.lang, 1)
        v.addLayout(row)
        self.view = QTextBrowser()
        self.view.setOpenLinks(False)
        self.view.anchorClicked.connect(self._link)
        v.addWidget(self.view, 1)
        self.question = QLineEdit()
        self.question.setPlaceholderText(_t("Ask a question about this PDF\u2026"))
        self.question.returnPressed.connect(self.ask)
        ask = QPushButton(_t("Ask"))
        ask.setObjectName("Primary")
        ask.clicked.connect(self.ask)
        qrow = QHBoxLayout()
        qrow.addWidget(self.question, 1)
        qrow.addWidget(ask)
        v.addLayout(qrow)
        self.hint = QLabel(_t("Answers cite pages - click a citation to jump there."))
        self.hint.setStyleSheet("color:#5f6368; font-size:11px")
        self.hint.setWordWrap(True)
        v.addWidget(self.hint)
        self._worker = None
        self._answer = ""

    def reset(self):
        self._index = None
        self._pages = None
        self.history.clear()
        self.view.clear()

    def _ensure_pages(self) -> bool:
        if self._pages is None:
            self._pages = self.pages_provider()
            self._index = ai.BM25(ai.make_chunks(self._pages))
        if not any((p or "").strip() for p in self._pages):
            self._say("<i>%s</i>" % html.escape(_t("This PDF has no text layer (it is a scan). Run OCR first, then "
                                                   "ask again.")))
            if self.ocr_cb:
                self.ocr_cb()
            self._pages = None
            return False
        return True

    def _say(self, html_text: str):
        self.view.append(html_text)

    def _busy(self, on: bool):
        self.summary_btn.setEnabled(not on)
        self.question.setEnabled(not on)

    def summarize(self):
        if not ensure_ai(self) or not self._ensure_pages():
            return
        pages, lang = self._pages, self.lang.currentData()
        self._say("<p><b>%s</b></p>" % _t("Summary"))
        self._busy(True)
        self._worker = AIWorker(lambda _d, cancelled: ai.summarize_pages(pages, lang, is_cancelled=cancelled), self)
        self._worker.done.connect(self._show_answer)
        self._worker.failed.connect(self._failed)
        self._worker.start()

    def ask(self):
        q = self.question.text().strip()
        if not q or not ensure_ai(self) or not self._ensure_pages():
            return
        hits = [c for _s, c in self._index.search(q, 6)]
        if not hits:  # nothing matches literally: give the model the start of the document
            hits = ai.make_chunks(self._pages)[:4]
        msgs = ai.pdf_question_messages(q, hits, self.history, self.lang.currentData())
        self._say("<p style='color:#1a73e8'><b>%s</b> %s</p>" % (_t("You:"), html.escape(q)))
        self.question.clear()
        self._busy(True)
        self._pending_q = q
        self._answer = ""
        self._worker = AIWorker(lambda on_delta, cancelled: ai.chat(msgs, on_delta=on_delta, is_cancelled=cancelled),
                                self)
        self._worker.done.connect(self._qa_done)
        self._worker.failed.connect(self._failed)
        self._worker.start()

    def _qa_done(self, text):
        self.history.append((self._pending_q, text))
        self._show_answer(text)

    def _show_answer(self, text):
        self._busy(False)
        body = html.escape(ai.clean_text(text)).replace("\n", "<br>")
        self._say("<p>%s</p>" % ai.link_page_refs(body))

    def _failed(self, msg):
        self._busy(False)
        self._say("<p style='color:#c5221f'>%s</p>" % html.escape(msg or _t("Stopped.")))

    def _link(self, url: QUrl):
        s = url.toString()
        if s.startswith("page:"):
            try:
                self.goto_page(int(s[5:]) - 1)
            except ValueError:
                pass


# ============================================================================ spreadsheet
class FormulaDialog(QDialog):
    def __init__(self, parent, context: str, cell: str, insert_cb):
        super().__init__(parent)
        self.setWindowTitle(_t("AI Formula"))
        self.resize(620, 360)
        self.context, self.cell, self.insert_cb = context, cell, insert_cb
        v = QVBoxLayout(self)
        v.addWidget(QLabel(_t("Describe what the formula in {0} should calculate:").format(cell)))
        self.request = QPlainTextEdit()
        self.request.setPlaceholderText(_t("e.g. total sales of the \"North\" region where the amount is over "
                                           "1000"))
        self.request.setMaximumHeight(90)
        v.addWidget(self.request)
        self.go = QPushButton(_t("Generate Formula"))
        self.go.setObjectName("Primary")
        self.go.clicked.connect(self.generate)
        v.addWidget(self.go, 0, Qt.AlignLeft)
        self.formula = QLineEdit()
        self.formula.setFont(QFont("Consolas", 11))
        v.addWidget(QLabel(_t("Formula:")))
        v.addWidget(self.formula)
        self.expl = QLabel("")
        self.expl.setWordWrap(True)
        self.expl.setStyleSheet("color:#5f6368")
        v.addWidget(self.expl, 1)
        bb = QDialogButtonBox()
        self.ins = bb.addButton(_t("Insert into {0}").format(cell), QDialogButtonBox.AcceptRole)
        bb.addButton(QDialogButtonBox.Close)
        bb.accepted.connect(self._insert)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        self._worker = None

    def generate(self):
        req = self.request.toPlainText().strip()
        if not req or not ensure_ai(self):
            return
        msgs = ai.formula_messages(req, self.context, self.cell)
        self.go.setEnabled(False)
        self.expl.setText(_t("Thinking\u2026"))
        self._worker = AIWorker(lambda _d, c: ai.chat(msgs, temperature=0, is_cancelled=c, json_mode=True), self)
        self._worker.done.connect(self._done)
        self._worker.failed.connect(lambda m: (self.expl.setText(m), self.go.setEnabled(True)))
        self._worker.start()

    def _done(self, answer):
        self.go.setEnabled(True)
        try:
            f, e = ai.parse_formula(answer)
        except ai.AIError as exc:
            self.expl.setText(str(exc))
            return
        self.formula.setText(f)
        self.expl.setText(e)

    def _insert(self):
        f = self.formula.text().strip()
        if f:
            self.insert_cb(f if f.startswith("=") else "=" + f)
            self.accept()


class AnalyzeDialog(QDialog):
    def __init__(self, parent, summary: str):
        super().__init__(parent)
        self.setWindowTitle(_t("AI Data Analysis"))
        self.resize(720, 560)
        self.summary = summary
        v = QVBoxLayout(self)
        v.addWidget(QLabel(_t("Statistics of the selected range (computed on this PC):")))
        st = QPlainTextEdit(summary)
        st.setReadOnly(True)
        st.setMaximumHeight(170)
        st.setFont(QFont("Consolas", 9))
        v.addWidget(st)
        q = QHBoxLayout()
        self.question = QLineEdit()
        self.question.setPlaceholderText(_t("Optional question, e.g. \"which product grows fastest?\""))
        self.go = QPushButton(_t("Analyze"))
        self.go.setObjectName("Primary")
        self.go.clicked.connect(self.run)
        q.addWidget(self.question, 1)
        q.addWidget(self.go)
        v.addLayout(q)
        self.out = QTextBrowser()
        v.addWidget(self.out, 1)
        bb = QDialogButtonBox(QDialogButtonBox.Close)
        bb.rejected.connect(self.reject)
        copy = bb.addButton(_t("Copy"), QDialogButtonBox.ActionRole)
        copy.clicked.connect(lambda: QApplication.clipboard().setText(self.out.toPlainText()))
        v.addWidget(bb)
        self._worker = None

    def run(self):
        if not ensure_ai(self):
            return
        msgs = ai.analysis_messages(self.summary, self.question.text().strip())
        self.go.setEnabled(False)
        self.out.setPlainText("")
        self._worker = AIWorker(lambda d, c: ai.chat(msgs, on_delta=d, is_cancelled=c), self)
        self._worker.delta.connect(lambda t: self.out.insertPlainText(t))
        self._worker.done.connect(lambda t: (self.out.setPlainText(ai.clean_text(t)), self.go.setEnabled(True)))
        self._worker.failed.connect(lambda m: (self.out.setPlainText(m), self.go.setEnabled(True)))
        self._worker.start()


# ============================================================================ presentation
class AISlidesDialog(QDialog):
    def __init__(self, parent, build_cb):
        super().__init__(parent)
        self.setWindowTitle(_t("AI Slides"))
        self.resize(620, 420)
        self.build_cb = build_cb
        v = QVBoxLayout(self)
        v.addWidget(QLabel(_t("What is the presentation about?")))
        self.topic = QPlainTextEdit()
        self.topic.setPlaceholderText(_t("e.g. \"Solar energy for small businesses in Pakistan: costs, savings, "
                                         "financing\""))
        self.topic.setMaximumHeight(100)
        v.addWidget(self.topic)
        form = QFormLayout()
        self.count = QSpinBox()
        self.count.setRange(2, 25)
        self.count.setValue(6)
        form.addRow(_t("Content slides:"), self.count)
        self.lang = QComboBox()
        self.lang.addItem(_t("Same as the topic"), "")
        for name in ai.LANGUAGES:
            self.lang.addItem(_t(name), name)
        form.addRow(_t("Language:"), self.lang)
        self.audience = QLineEdit()
        self.audience.setPlaceholderText(_t("optional, e.g. investors, students"))
        form.addRow(_t("Audience:"), self.audience)
        self.notes = QCheckBox(_t("Write speaker notes"))
        self.notes.setChecked(True)
        form.addRow("", self.notes)
        v.addLayout(form)
        self.status = QLabel("")
        self.status.setWordWrap(True)
        v.addWidget(self.status)
        bb = QDialogButtonBox()
        self.go = bb.addButton(_t("Create Slides"), QDialogButtonBox.AcceptRole)
        self.go.setObjectName("Primary")
        bb.addButton(QDialogButtonBox.Cancel)
        bb.accepted.connect(self.run)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)
        self._worker = None

    def run(self):
        topic = self.topic.toPlainText().strip()
        if not topic or not ensure_ai(self):
            return
        msgs = ai.slides_messages(topic, self.count.value(), self.lang.currentData(), self.audience.text().strip())
        self.go.setEnabled(False)
        self.status.setText(_t("Creating the outline\u2026"))
        self._worker = AIWorker(lambda _d, c: ai.chat(msgs, is_cancelled=c, json_mode=True), self)
        self._worker.done.connect(self._done)
        self._worker.failed.connect(lambda m: (self.status.setText(m), self.go.setEnabled(True)))
        self._worker.start()

    def _done(self, answer):
        self.go.setEnabled(True)
        try:
            title, subtitle, slides = ai.parse_slides(answer)
        except ai.AIError as exc:
            self.status.setText(str(exc))
            return
        if not self.notes.isChecked():
            for s in slides:
                s.notes = ""
        self.build_cb(title or self.topic.toPlainText().strip()[:80], subtitle, slides)
        self.accept()
