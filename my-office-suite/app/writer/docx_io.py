"""DOCX (OOXML WordprocessingML) import/export for the Writer.

Import walks ``word/document.xml`` directly (via python-docx's package loader) so it can
handle style inheritance, numbering, tables with merged cells, images, hyperlinks,
fields, bookmarks, comments and tracked changes (``w:ins`` / ``w:del``).
Export builds the WordprocessingML with python-docx + lxml, writing run properties,
real numbering, complex-script (Urdu/Hindi) fonts, bidi paragraphs, revisions,
comments, headers/footers with PAGE fields and page setup.
"""
from __future__ import annotations

import io
import json
import logging
import os
import re
import tempfile
import uuid
import zipfile
from datetime import datetime

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt
from docx.text.run import Run
from lxml import etree
from PySide6.QtCore import QBuffer, QByteArray, QIODevice, Qt
from PySide6.QtGui import (QBrush, QColor, QFont, QImage, QPainter, QTextBlockFormat, QTextCharFormat,
                           QTextCursor, QTextDocument, QTextFormat, QTextFrameFormat, QTextImageFormat, QTextLength,
                           QTextListFormat, QTextTable, QTextTableFormat)

from app.shell.document import DocumentError
from app.writer.model import (P_EQUATION, P_FOOTNOTE, P_COMMENTS, P_DEL, P_INS, P_STYLE, P_TOC, Comment,
                              HeaderFooter, PageSetup, color_hex, has_arabic, has_devanagari, has_rtl,
                              parse_stamp, style_for_heading)

log = logging.getLogger(__name__)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
V = "urn:schemas-microsoft-com:vml"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
CUSTOM_NS = "urn:myoffice:writer"
CUSTOMXML_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
CUSTOMXMLPROPS_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps"

EMU_PER_PX = 9525
TWIP_PER_PX = 15.0  # 1px = 0.75pt = 15 twips

HIGHLIGHT_COLORS = {
    "yellow": "FFFF00", "green": "00FF00", "cyan": "00FFFF", "magenta": "FF00FF", "blue": "0000FF",
    "red": "FF0000", "darkBlue": "000080", "darkCyan": "008080", "darkGreen": "008000",
    "darkMagenta": "800080", "darkRed": "800000", "darkYellow": "808000", "darkGray": "808080",
    "lightGray": "C0C0C0", "black": "000000", "white": "FFFFFF",
}


def w(tag: str) -> str:
    return "{%s}%s" % (W, tag)


def _on(el) -> bool | None:
    """Boolean OOXML toggle property."""
    if el is None:
        return None
    val = el.get(w("val"))
    return val is None or val.lower() not in ("0", "false", "off", "none")


def _int(v, default=0):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


# ============================================================================ run properties
class RunProps:
    __slots__ = ("fonts", "size", "size_cs", "bold", "bold_cs", "italic", "italic_cs", "underline", "strike",
                 "color", "highlight", "shading", "valign", "caps", "small_caps", "rtl", "vanish", "spacing", "style_link")

    def __init__(self):
        self.fonts = {}
        for s in self.__slots__[1:]:
            setattr(self, s, None)

    def copy(self) -> "RunProps":
        rp = RunProps()
        rp.fonts = dict(self.fonts)
        for s in self.__slots__[1:]:
            setattr(rp, s, getattr(self, s))
        return rp

    def apply_xml(self, rpr, theme_fonts: dict) -> None:
        if rpr is None:
            return
        f = rpr.find(w("rFonts"))
        if f is not None:
            for key, theme_key in (("ascii", "asciiTheme"), ("hAnsi", "hAnsiTheme"), ("cs", "cstheme"),
                                   ("eastAsia", "eastAsiaTheme")):
                val = f.get(w(key))
                th = f.get(w(theme_key))
                if th:
                    val = theme_fonts.get("major" if th.startswith("major") else "minor") or val
                if val:
                    self.fonts[key] = val
        for tag, attr in (("b", "bold"), ("bCs", "bold_cs"), ("i", "italic"), ("iCs", "italic_cs"),
                          ("caps", "caps"), ("smallCaps", "small_caps"), ("rtl", "rtl"), ("vanish", "vanish")):
            v = _on(rpr.find(w(tag)))
            if v is not None:
                setattr(self, attr, v)
        st = _on(rpr.find(w("strike")))
        dst = _on(rpr.find(w("dstrike")))
        if st is not None or dst is not None:
            self.strike = bool(st) or bool(dst)
        u = rpr.find(w("u"))
        if u is not None:
            self.underline = (u.get(w("val")) or "single") != "none"
        c = rpr.find(w("color"))
        if c is not None and c.get(w("val")) and c.get(w("val")).lower() != "auto":
            self.color = c.get(w("val"))
        sz = rpr.find(w("sz"))
        if sz is not None and sz.get(w("val")):
            self.size = _int(sz.get(w("val"))) / 2.0
        szc = rpr.find(w("szCs"))
        if szc is not None and szc.get(w("val")):
            self.size_cs = _int(szc.get(w("val"))) / 2.0
        hl = rpr.find(w("highlight"))
        if hl is not None and hl.get(w("val")) in HIGHLIGHT_COLORS:
            self.highlight = HIGHLIGHT_COLORS[hl.get(w("val"))]
        shd = rpr.find(w("shd"))
        if shd is not None and shd.get(w("fill")) and shd.get(w("fill")).lower() != "auto":
            self.shading = shd.get(w("fill"))
        va = rpr.find(w("vertAlign"))
        if va is not None:
            self.valign = va.get(w("val"))
        sp = rpr.find(w("spacing"))
        if sp is not None and sp.get(w("val")):
            self.spacing = _int(sp.get(w("val"))) / TWIP_PER_PX

    def char_format(self, text: str, base: QTextCharFormat | None = None) -> QTextCharFormat:
        cf = QTextCharFormat(base) if base is not None else QTextCharFormat()
        complex_script = bool(self.rtl) or has_arabic(text) or has_devanagari(text) or has_rtl(text)
        fam = None
        if complex_script:
            fam = self.fonts.get("cs") or self.fonts.get("ascii") or self.fonts.get("hAnsi")
        else:
            fam = self.fonts.get("ascii") or self.fonts.get("hAnsi") or self.fonts.get("eastAsia")
        if fam:
            cf.setFontFamilies([fam])
        size = (self.size_cs or self.size) if complex_script else self.size
        if size:
            cf.setFontPointSize(size)
        bold = (self.bold_cs if self.bold_cs is not None else self.bold) if complex_script else self.bold
        if bold is not None:
            cf.setFontWeight(QFont.Bold if bold else QFont.Normal)
        italic = (self.italic_cs if self.italic_cs is not None else self.italic) if complex_script else self.italic
        if italic is not None:
            cf.setFontItalic(bool(italic))
        if self.underline is not None:
            cf.setFontUnderline(bool(self.underline))
        if self.strike is not None:
            cf.setFontStrikeOut(bool(self.strike))
        if self.color:
            cf.setForeground(QBrush(QColor("#" + self.color)))
        if self.highlight or self.shading:
            cf.setBackground(QBrush(QColor("#" + (self.highlight or self.shading))))
        if self.valign == "superscript":
            cf.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
        elif self.valign == "subscript":
            cf.setVerticalAlignment(QTextCharFormat.AlignSubScript)
        if self.caps:
            cf.setFontCapitalization(QFont.AllUppercase)
        elif self.small_caps:
            cf.setFontCapitalization(QFont.SmallCaps)
        if self.spacing:
            cf.setFontLetterSpacingType(QFont.AbsoluteSpacing)
            cf.setFontLetterSpacing(self.spacing)
        return cf


class ParaProps:
    __slots__ = ("jc", "before", "after", "line", "line_rule", "ind_left", "ind_right", "first_line", "hanging",
                 "num_id", "ilvl", "outline", "bidi", "page_break_before", "keep_next", "shading", "border_bottom")

    def __init__(self):
        for s in self.__slots__:
            setattr(self, s, None)

    def copy(self) -> "ParaProps":
        pp = ParaProps()
        for s in self.__slots__:
            setattr(pp, s, getattr(self, s))
        return pp

    def apply_xml(self, ppr) -> None:
        if ppr is None:
            return
        jc = ppr.find(w("jc"))
        if jc is not None:
            self.jc = jc.get(w("val"))
        sp = ppr.find(w("spacing"))
        if sp is not None:
            if sp.get(w("before")) is not None:
                self.before = _int(sp.get(w("before")))
            if sp.get(w("after")) is not None:
                self.after = _int(sp.get(w("after")))
            if sp.get(w("line")) is not None:
                self.line = _int(sp.get(w("line")))
                self.line_rule = sp.get(w("lineRule")) or "auto"
        ind = ppr.find(w("ind"))
        if ind is not None:
            for attr, keys in (("ind_left", ("left", "start")), ("ind_right", ("right", "end")),
                               ("first_line", ("firstLine",)), ("hanging", ("hanging",))):
                for k in keys:
                    if ind.get(w(k)) is not None:
                        setattr(self, attr, _int(ind.get(w(k))))
                        break
        num = ppr.find(w("numPr"))
        if num is not None:
            nid = num.find(w("numId"))
            il = num.find(w("ilvl"))
            if nid is not None:
                self.num_id = nid.get(w("val"))
            if il is not None:
                self.ilvl = _int(il.get(w("val")))
        ol = ppr.find(w("outlineLvl"))
        if ol is not None:
            self.outline = _int(ol.get(w("val")), 9)
        b = _on(ppr.find(w("bidi")))
        if b is not None:
            self.bidi = b
        pb = _on(ppr.find(w("pageBreakBefore")))
        if pb is not None:
            self.page_break_before = pb
        shd = ppr.find(w("shd"))
        if shd is not None and shd.get(w("fill")) and shd.get(w("fill")).lower() != "auto":
            self.shading = shd.get(w("fill"))
        bdr = ppr.find(w("pBdr"))
        if bdr is not None and bdr.find(w("bottom")) is not None:
            self.border_bottom = True


# ============================================================================ importer
class _Ctx:
    def __init__(self, cursor: QTextCursor, reuse: bool = True):
        self.cur = cursor
        self.reuse = reuse


def _open_docx(path: str):
    """Open .docx/.docm/.dotx/.dotm (templates are patched to the document content-type)."""
    try:
        with open(path, "rb") as fh:
            data = fh.read()
    except OSError as exc:
        raise DocumentError(f"Cannot read file: {exc}") from exc
    if not data.startswith(b"PK"):
        raise DocumentError("This file is not a valid DOCX (OOXML) document. If it is a legacy .doc file, "
                            "install LibreOffice to convert it, or use Tools \u2192 File Repair.")
    try:
        zin = zipfile.ZipFile(io.BytesIO(data))
        names = zin.namelist()
        ct = zin.read("[Content_Types].xml").decode("utf-8", "replace")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise DocumentError("The document is damaged (invalid ZIP container). Try Tools \u2192 File Repair.") from exc
    patched = re.sub(r"wordprocessingml\.(template\.main|document\.macroEnabled\.main|template\.macroEnabledTemplate\.main)",
                     "wordprocessingml.document.main", ct)
    patched = patched.replace("application/vnd.ms-word.document.macroEnabled.main+xml",
                              "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
    patched = patched.replace("application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
                              "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
    if patched != ct:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for n in names:
                zout.writestr(n, patched if n == "[Content_Types].xml" else zin.read(n))
        data = out.getvalue()
    try:
        return docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError) as exc:
        raise DocumentError("The document is damaged (missing parts). Try Tools \u2192 File Repair.") from exc
    except etree.XMLSyntaxError as exc:
        raise DocumentError(f"The document XML is damaged ({exc}). Try Tools \u2192 File Repair.") from exc
    except ValueError as exc:
        raise DocumentError(f"Unsupported Word document: {exc}") from exc


class DocxImporter:
    def __init__(self, qdoc: QTextDocument, register_image):
        self.qdoc = qdoc
        self.register_image = register_image
        self.page = PageSetup()
        self.hf = HeaderFooter()
        self.comments: dict[str, Comment] = {}
        self.sources: list[dict] = []
        self.citation_style = "APA"
        self.watermark = ""
        self.title = ""
        self._lists: dict = {}
        self._style_cache: dict = {}
        self._field_stack: list[dict] = []
        self._active_comments: set[str] = set()
        self._pending_bookmarks: list[str] = []
        self._extra_paras: list = []
        self._notes: dict[tuple[str, str], str] = {}   # ("footnote"|"endnote", id) -> text
        self.protection: dict[str, str] = {}           # w:documentProtection attributes (enforced only)

    # ------------------------------------------------------------------ entry
    def load(self, path: str) -> None:
        d = _open_docx(path)
        self.d = d
        self._read_theme()
        self._read_styles()
        self._read_numbering()
        self._read_comments()
        self._read_custom()
        self._read_notes()
        self._read_protection()
        self.qdoc.clear()
        try:
            self.title = d.core_properties.title or ""
        except Exception:
            self.title = ""
        ctx = _Ctx(QTextCursor(self.qdoc), True)
        ctx.cur.beginEditBlock()
        body = d.element.body
        for child in body.iterchildren():
            self._block_level(child, ctx)
        ctx.cur.endEditBlock()
        self._read_section(body)
        if getattr(self, "_has_notes", False):
            from app.writer.notes import renumber

            renumber(self.qdoc)

    # ------------------------------------------------------------------ package parts
    def _read_notes(self):
        for rel in self.d.part.rels.values():
            kind = {RT.FOOTNOTES: "footnote", RT.ENDNOTES: "endnote"}.get(rel.reltype)
            if not kind:
                continue
            try:
                part = rel.target_part
                root = part.element if hasattr(part, "element") else etree.fromstring(part.blob)
            except Exception as exc:
                log.warning("%ss not read: %s", kind, exc)
                continue
            for note in root.findall(w(kind)):
                if note.get(w("type")) in ("separator", "continuationSeparator", "continuationNotice"):
                    continue
                paras = []
                for p_el in note.findall(".//" + w("p")):
                    paras.append("".join((t.text or "") if t.tag == w("t") else "\t" if t.tag == w("tab") else ""
                                         for t in p_el.iter(w("t"), w("tab"))).strip())
                self._notes[(kind, note.get(w("id")) or "")] = "\n".join(x for x in paras if x)

    def _read_protection(self):
        try:
            el = self.d.settings.element.find(w("documentProtection"))
        except Exception:
            el = None
        if el is None:
            return
        attrs = {etree.QName(k).localname: v for k, v in el.attrib.items()}
        if attrs.get("enforcement") in ("1", "true", "on") and attrs.get("edit") in (
                "readOnly", "comments", "trackedChanges", "forms"):
            self.protection = attrs

    def _read_theme(self):
        self.theme_fonts = {"minor": "Calibri", "major": "Calibri Light"}
        try:
            for rel in self.d.part.rels.values():
                if rel.reltype == RT.THEME:
                    root = etree.fromstring(rel.target_part.blob)
                    for kind in ("minor", "major"):
                        el = root.find(f".//{{{A}}}{kind}Font/{{{A}}}latin")
                        if el is not None and el.get("typeface"):
                            self.theme_fonts[kind] = el.get("typeface")
        except Exception as exc:
            log.debug("theme read failed: %s", exc)

    def _read_styles(self):
        self.styles = {}
        self.default_pstyle = None
        self.default_rpr = RunProps()
        self.default_ppr = ParaProps()
        try:
            root = self.d.styles.element
        except Exception:
            return
        dd = root.find(w("docDefaults"))
        if dd is not None:
            self.default_rpr.apply_xml(dd.find(f"{w('rPrDefault')}/{w('rPr')}"), self.theme_fonts)
            self.default_ppr.apply_xml(dd.find(f"{w('pPrDefault')}/{w('pPr')}"))
        for st in root.findall(w("style")):
            sid = st.get(w("styleId"))
            name_el = st.find(w("name"))
            based = st.find(w("basedOn"))
            self.styles[sid] = {
                "name": name_el.get(w("val")) if name_el is not None else sid,
                "type": st.get(w("type")),
                "based": based.get(w("val")) if based is not None else None,
                "ppr": st.find(w("pPr")), "rpr": st.find(w("rPr")),
            }
            if st.get(w("type")) == "paragraph" and st.get(w("default")) in ("1", "true"):
                self.default_pstyle = sid

    def _style_props(self, sid: str | None):
        """Resolved (ParaProps, RunProps, name) for a style id (including docDefaults)."""
        if sid in self._style_cache:
            return self._style_cache[sid]
        chain = []
        seen = set()
        cur = sid
        while cur and cur in self.styles and cur not in seen:
            seen.add(cur)
            chain.append(self.styles[cur])
            cur = self.styles[cur]["based"]
        pp = self.default_ppr.copy()
        rp = self.default_rpr.copy()
        for st in reversed(chain):
            pp.apply_xml(st["ppr"])
            rp.apply_xml(st["rpr"], self.theme_fonts)
        name = self.styles[sid]["name"] if sid in self.styles else "Normal"
        res = (pp, rp, name)
        self._style_cache[sid] = res
        return res

    def _char_style_props(self, sid: str | None) -> RunProps | None:
        if not sid or sid not in self.styles:
            return None
        key = ("char", sid)
        if key in self._style_cache:
            return self._style_cache[key]
        chain = []
        cur = sid
        seen = set()
        while cur and cur in self.styles and cur not in seen:
            seen.add(cur)
            chain.append(self.styles[cur])
            cur = self.styles[cur]["based"]
        rp = RunProps()
        for st in reversed(chain):
            rp.apply_xml(st["rpr"], self.theme_fonts)
        self._style_cache[key] = rp
        return rp

    def _read_numbering(self):
        self.num_fmt = {}  # (numId, ilvl) -> (numFmt, lvlText)
        try:
            numbering = self.d.part.numbering_part.element
        except Exception:
            return
        abstract = {}
        for an in numbering.findall(w("abstractNum")):
            aid = an.get(w("abstractNumId"))
            lv = {}
            for lvl in an.findall(w("lvl")):
                il = _int(lvl.get(w("ilvl")))
                fmt = lvl.find(w("numFmt"))
                txt = lvl.find(w("lvlText"))
                lv[il] = (fmt.get(w("val")) if fmt is not None else "decimal",
                          txt.get(w("val")) if txt is not None else "%1.")
            abstract[aid] = lv
        for num in numbering.findall(w("num")):
            nid = num.get(w("numId"))
            an = num.find(w("abstractNumId"))
            if an is None:
                continue
            for il, val in abstract.get(an.get(w("val")), {}).items():
                self.num_fmt[(nid, il)] = val

    def _read_comments(self):
        try:
            for rel in self.d.part.rels.values():
                if rel.reltype == RT.COMMENTS:
                    part = rel.target_part
                    root = part.element if hasattr(part, "element") else etree.fromstring(part.blob)
                    for c in root.findall(w("comment")):
                        cid = c.get(w("id"))
                        paras = ["".join(t.text or "" for t in p.iter(w("t"))) for p in c.findall(f".//{w('p')}")]
                        self.comments[f"c{cid}"] = Comment(id=f"c{cid}", author=c.get(w("author")) or "",
                                                           text="\n".join(paras).strip(),
                                                           date=c.get(w("date")) or "",
                                                           initials=c.get(w("initials")) or "")
        except Exception as exc:
            log.warning("Could not read comments: %s", exc)

    def _read_custom(self):
        try:
            for rel in self.d.part.rels.values():
                if rel.reltype == CUSTOMXML_RT:
                    blob = rel.target_part.blob
                    if CUSTOM_NS.encode() in blob:
                        root = etree.fromstring(blob)
                        data = json.loads(root.text or "{}")
                        self.sources = data.get("sources", [])
                        self.citation_style = data.get("citation_style", "APA")
                        self.watermark = data.get("watermark", "") or ""
                        hf = data.get("header_footer")
                        if hf:
                            self._custom_hf = hf
        except Exception as exc:
            log.warning("Custom data unreadable: %s", exc)

    def _read_section(self, body):
        sect = body.find(w("sectPr"))
        if sect is None:
            secs = body.findall(f".//{w('sectPr')}")
            sect = secs[-1] if secs else None
        if sect is not None:
            pg = sect.find(w("pgSz"))
            if pg is not None:
                wmm = _int(pg.get(w("w")), 11906) / 1440 * 25.4
                hmm = _int(pg.get(w("h")), 16838) / 1440 * 25.4
                orient = pg.get(w("orient")) or ("landscape" if wmm > hmm else "portrait")
                self.page.orientation = orient
                self.page.width_mm, self.page.height_mm = (min(wmm, hmm), max(wmm, hmm))
                self.page.size = PageSetup.detect_size(wmm, hmm)
            mar = sect.find(w("pgMar"))
            if mar is not None:
                tw = lambda k, d: abs(_int(mar.get(w(k)), d)) / 1440 * 25.4  # noqa: E731
                self.page.margins_mm = [tw("top", 1440), tw("right", 1440), tw("bottom", 1440), tw("left", 1440)]
                self.page.header_mm = tw("header", 720)
                self.page.footer_mm = tw("footer", 720)
        try:
            sec = self.d.sections[0]
            if not sec.header.is_linked_to_previous:
                if not self.watermark:
                    for tp in sec.header._element.iter("{%s}textpath" % V):
                        if tp.get("string"):
                            self.watermark = tp.get("string")
                            break
                self.hf.header, self.hf.header_align = self._hf_text(sec.header._element)
            if not sec.footer.is_linked_to_previous:
                self.hf.footer, self.hf.footer_align = self._hf_text(sec.footer._element)
        except Exception as exc:
            log.debug("header/footer read failed: %s", exc)
        custom = getattr(self, "_custom_hf", None)
        if custom:
            self.hf.header_align = custom.get("header_align", self.hf.header_align)
            self.hf.footer_align = custom.get("footer_align", self.hf.footer_align)

    def _hf_text(self, el) -> tuple[str, str]:
        parts: list[str] = []
        align = "center"
        for p in el.findall(f".//{w('p')}"):
            jc = p.find(f"{w('pPr')}/{w('jc')}")
            txt = []
            in_instr = False
            skip_result = False
            for node in p.iter():
                if node.tag == w("fldSimple"):
                    token = _hf_token(node.get(w("instr")) or "")
                    if token:
                        txt.append(token)
                        node.set("_skip", "1")
                elif node.tag == w("fldChar"):
                    t = node.get(w("fldCharType"))
                    if t == "begin":
                        in_instr = True
                    elif t == "separate":
                        in_instr = False
                    elif t == "end":
                        skip_result = False
                        in_instr = False
                elif node.tag == w("instrText") and in_instr:
                    token = _hf_token(node.text or "")
                    if token:
                        txt.append(token)
                        skip_result = True
                elif node.tag == w("t") and not in_instr and not skip_result:
                    anc = node.getparent()
                    inside_simple = False
                    while anc is not None and anc is not p:
                        if anc.tag == w("fldSimple") and anc.get("_skip"):
                            inside_simple = True
                            break
                        anc = anc.getparent()
                    if not inside_simple:
                        txt.append(node.text or "")
                elif node.tag == w("tab"):
                    txt.append("  ")
            line = "".join(txt).strip()
            if line:
                parts.append(line)
                if jc is not None:
                    align = {"left": "left", "start": "left", "right": "right", "end": "right"}.get(jc.get(w("val")), "center")
        return " ".join(parts), align

    # ------------------------------------------------------------------ block level
    def _block_level(self, el, ctx: _Ctx):
        tag = el.tag
        if tag == w("p"):
            self._paragraph(el, ctx)
            while self._extra_paras:
                extra = self._extra_paras
                self._extra_paras = []
                for p in extra:
                    self._paragraph(p, ctx)
        elif tag == w("tbl"):
            self._table(el, ctx)
        elif tag in (w("sdt"),):
            content = el.find(w("sdtContent"))
            if content is not None:
                for ch in content.iterchildren():
                    self._block_level(ch, ctx)
        elif tag in (w("customXml"), w("smartTag")):
            for ch in el.iterchildren():
                self._block_level(ch, ctx)

    def _paragraph(self, p_el, ctx: _Ctx):
        ppr = p_el.find(w("pPr"))
        sid = None
        if ppr is not None and ppr.find(w("pStyle")) is not None:
            sid = ppr.find(w("pStyle")).get(w("val"))
        sid = sid or self.default_pstyle
        spp, srp, sname = self._style_props(sid)
        pp = spp.copy()
        pp.apply_xml(ppr)
        rp_base = srp.copy()
        bf = QTextBlockFormat()
        bf.setProperty(P_STYLE, sname)
        heading = 0
        m = re.match(r"(?i)heading\s*(\d)", sname or "")
        if m:
            heading = int(m.group(1))
        elif pp.outline is not None and pp.outline < 9:
            heading = pp.outline + 1
        if heading:
            bf.setHeadingLevel(min(heading, 6))
        jc = pp.jc
        if jc in ("center",):
            bf.setAlignment(Qt.AlignHCenter)
        elif jc in ("right", "end"):
            bf.setAlignment(Qt.AlignRight)
        elif jc in ("both", "distribute", "lowKashida", "mediumKashida", "highKashida", "thaiDistribute"):
            bf.setAlignment(Qt.AlignJustify)
        elif jc in ("left", "start"):
            bf.setAlignment(Qt.AlignLeft)
        if pp.bidi:
            bf.setLayoutDirection(Qt.RightToLeft)
        if pp.before:
            bf.setTopMargin(pp.before / TWIP_PER_PX)
        if pp.after:
            bf.setBottomMargin(pp.after / TWIP_PER_PX)
        if pp.line:
            if pp.line_rule == "exact":
                bf.setLineHeight(pp.line / TWIP_PER_PX, QTextBlockFormat.FixedHeight.value)
            elif pp.line_rule == "atLeast":
                bf.setLineHeight(pp.line / TWIP_PER_PX, QTextBlockFormat.MinimumHeight.value)
            else:
                bf.setLineHeight(max(50.0, pp.line / 240.0 * 100.0), QTextBlockFormat.ProportionalHeight.value)
        in_list = pp.num_id not in (None, "0") and self._has_marker(pp.num_id, pp.ilvl or 0)
        if pp.ind_left and not in_list:
            bf.setLeftMargin(pp.ind_left / TWIP_PER_PX)
        if pp.ind_right:
            bf.setRightMargin(pp.ind_right / TWIP_PER_PX)
        if pp.hanging and not in_list:
            bf.setTextIndent(-pp.hanging / TWIP_PER_PX)
        elif pp.first_line:
            bf.setTextIndent(pp.first_line / TWIP_PER_PX)
        if pp.page_break_before:
            bf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
        if pp.shading:
            bf.setBackground(QColor("#" + pp.shading))
        if pp.border_bottom:
            bf.setProperty(QTextFormat.BlockTrailingHorizontalRulerWidth, QTextLength(QTextLength.PercentageLength, 100))
        base_cf = rp_base.char_format("")
        if self._field_stack and self._field_stack[-1].get("toc"):
            bf.setProperty(P_TOC, True)
        cur = ctx.cur
        if ctx.reuse:
            cur.setBlockFormat(bf)
            cur.setBlockCharFormat(base_cf)
            ctx.reuse = False
        else:
            cur.insertBlock(bf, base_cf)
        if in_list:
            self._attach_list(cur, pp.num_id, pp.ilvl or 0)
        self._inline_children(p_el, ctx, rp_base, {}, bf)

    def _has_marker(self, num_id: str, ilvl: int) -> bool:
        """False for a numbering level that shows nothing: numFmt "none" or an empty lvlText.  LibreOffice links
        every heading style to such an outline level; it must not turn the headings into list items."""
        fmt = self.num_fmt.get((num_id, ilvl))
        return fmt is None or (fmt[0] != "none" and fmt[1] != "")

    def _attach_list(self, cur: QTextCursor, num_id: str, ilvl: int):
        key = (num_id, ilvl)
        lst = self._lists.get(key)
        if lst is not None:
            lst.add(cur.block())
            return
        fmt_name, lvl_text = self.num_fmt.get(key, ("bullet" if ilvl % 2 == 0 else "decimal", ""))
        lf = QTextListFormat()
        style = {
            "decimal": QTextListFormat.ListDecimal, "decimalZero": QTextListFormat.ListDecimal,
            "lowerLetter": QTextListFormat.ListLowerAlpha, "upperLetter": QTextListFormat.ListUpperAlpha,
            "lowerRoman": QTextListFormat.ListLowerRoman, "upperRoman": QTextListFormat.ListUpperRoman,
        }.get(fmt_name)
        if style is None:
            style = [QTextListFormat.ListDisc, QTextListFormat.ListCircle, QTextListFormat.ListSquare][min(ilvl, 2)]
        lf.setStyle(style)
        lf.setIndent(ilvl + 1)
        if lvl_text and style not in (QTextListFormat.ListDisc, QTextListFormat.ListCircle, QTextListFormat.ListSquare):
            suffix = lvl_text.split("%" + str(ilvl + 1))[-1] if ("%" + str(ilvl + 1)) in lvl_text else "."
            try:
                lf.setNumberSuffix(suffix)
            except Exception:
                pass
        self._lists[key] = cur.createList(lf)

    # ------------------------------------------------------------------ inline content
    def _inline_children(self, parent, ctx: _Ctx, rp_base: RunProps, extra: dict, bf: QTextBlockFormat):
        for node in parent.iterchildren():
            tag = node.tag
            if tag == w("r"):
                self._run(node, ctx, rp_base, extra, bf)
            elif tag == w("hyperlink"):
                href = None
                rid = node.get(qn("r:id"))
                if rid:
                    try:
                        href = self.d.part.rels[rid].target_ref
                    except KeyError:
                        href = None
                if node.get(w("anchor")):
                    href = "#" + node.get(w("anchor"))
                e2 = dict(extra)
                if href:
                    e2["href"] = href
                self._inline_children(node, ctx, rp_base, e2, bf)
            elif tag in (w("ins"), w("moveTo")):
                e2 = dict(extra)
                e2["ins"] = f"{node.get(w('author')) or 'Unknown'}|{node.get(w('date')) or ''}"
                self._inline_children(node, ctx, rp_base, e2, bf)
            elif tag in (w("del"), w("moveFrom")):
                e2 = dict(extra)
                e2["del"] = f"{node.get(w('author')) or 'Unknown'}|{node.get(w('date')) or ''}"
                self._inline_children(node, ctx, rp_base, e2, bf)
            elif tag == w("fldSimple"):
                instr = (node.get(w("instr")) or "").strip()
                e2 = dict(extra)
                m = re.match(r'(?i)HYPERLINK\s+"([^"]+)"', instr)
                if m:
                    e2["href"] = m.group(1)
                self._inline_children(node, ctx, rp_base, e2, bf)
            elif tag in (w("smartTag"), w("customXml"), w("dir"), w("bdo")):
                self._inline_children(node, ctx, rp_base, extra, bf)
            elif tag == w("sdt"):
                content = node.find(w("sdtContent"))
                if content is not None:
                    self._inline_children(content, ctx, rp_base, extra, bf)
            elif tag == w("commentRangeStart"):
                self._active_comments.add("c" + (node.get(w("id")) or ""))
            elif tag == w("commentRangeEnd"):
                self._active_comments.discard("c" + (node.get(w("id")) or ""))
            elif tag == w("bookmarkStart"):
                name = node.get(w("name"))
                if name and name != "_GoBack":
                    self._pending_bookmarks.append(name)
            elif tag in ("{%s}oMath" % M_NS, "{%s}oMathPara" % M_NS):
                self._equation(node, ctx, rp_base, extra)
            elif tag == "{%s}AlternateContent" % MC:
                choice = node.find("{%s}Choice" % MC)
                target = choice if choice is not None else node.find("{%s}Fallback" % MC)
                if target is not None:
                    for r in target.iterchildren():
                        if r.tag == w("r"):
                            self._run(r, ctx, rp_base, extra, bf)

    def _run(self, r_el, ctx: _Ctx, rp_base: RunProps, extra: dict, bf: QTextBlockFormat):
        rpr = r_el.find(w("rPr"))
        rp = rp_base.copy()
        if rpr is not None:
            rs = rpr.find(w("rStyle"))
            if rs is not None:
                csp = self._char_style_props(rs.get(w("val")))
                if csp is not None:
                    for s in RunProps.__slots__[1:]:
                        v = getattr(csp, s)
                        if v is not None:
                            setattr(rp, s, v)
                    rp.fonts.update(csp.fonts)
            rp.apply_xml(rpr, self.theme_fonts)
        if rp.vanish:
            return
        cur = ctx.cur
        for node in r_el.iterchildren():
            tag = node.tag
            if tag == w("fldChar"):
                t = node.get(w("fldCharType"))
                if t == "begin":
                    self._field_stack.append({"state": "instr", "instr": "", "toc": False})
                elif t == "separate" and self._field_stack:
                    fld = self._field_stack[-1]
                    fld["state"] = "result"
                    instr = fld["instr"].strip()
                    fld["toc"] = instr.upper().startswith("TOC")
                    m = re.match(r'(?i)HYPERLINK\s+(?:\\l\s+)?"([^"]+)"', instr)
                    if m:
                        target = m.group(1)
                        fld["href"] = ("#" + target) if "\\l" in instr else target
                    if fld["toc"]:
                        b = cur.block()
                        f = b.blockFormat()
                        f.setProperty(P_TOC, True)
                        cur.setBlockFormat(f)
                elif t == "end" and self._field_stack:
                    self._field_stack.pop()
                continue
            if tag == w("instrText"):
                if self._field_stack and self._field_stack[-1]["state"] == "instr":
                    self._field_stack[-1]["instr"] += node.text or ""
                continue
            if self._field_stack and self._field_stack[-1]["state"] == "instr":
                continue  # inside field instructions
            text = None
            if tag in (w("t"), w("delText")):
                text = node.text or ""
            elif tag == w("tab"):
                text = "\t"
            elif tag in (w("br"), w("cr")):
                if node.get(w("type")) == "page":
                    if cur.block().length() > 1:
                        nbf = QTextBlockFormat(bf)
                        nbf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
                        nbf.clearProperty(P_TOC)
                        cur.insertBlock(nbf, cur.blockCharFormat())
                    else:
                        f = cur.blockFormat()
                        f.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
                        cur.setBlockFormat(f)
                    continue
                text = "\u2028"
            elif tag == w("noBreakHyphen"):
                text = "\u2011"
            elif tag == w("softHyphen"):
                text = "\u00ad"
            elif tag == w("sym"):
                try:
                    code = int(node.get(w("char")), 16)
                    if 0xF000 <= code <= 0xF0FF:
                        code -= 0xF000
                    text = chr(code)
                except (TypeError, ValueError):
                    text = ""
            elif tag == w("drawing") or tag == w("pict") or tag == "{%s}AlternateContent" % MC:
                self._drawing(node, ctx, rp, extra)
                continue
            elif tag in (w("footnoteReference"), w("endnoteReference")):
                kind = "footnote" if tag == w("footnoteReference") else "endnote"
                note = self._notes.get((kind, node.get(w("id")) or ""), "")
                from app.writer.notes import reference_format

                base = rp.char_format("1")
                self._decorate(base, extra)
                cur.insertText("1", reference_format(base, note))
                self._has_notes = True
                continue
            elif tag in (w("footnoteRef"), w("endnoteRef"), w("separator"), w("continuationSeparator")):
                continue
            if text:
                cf = rp.char_format(text)
                self._decorate(cf, extra)
                cur.insertText(text, cf)

    def _decorate(self, cf: QTextCharFormat, extra: dict):
        href = extra.get("href") or (self._field_stack[-1].get("href") if self._field_stack else None)
        if href:
            cf.setAnchor(True)
            cf.setAnchorHref(href)
            if not cf.foreground().style() or cf.foreground().color() == QColor("#000000"):
                cf.setForeground(QColor("#0563c1"))
            cf.setFontUnderline(True)
        if self._pending_bookmarks:
            cf.setAnchor(True)
            cf.setAnchorNames(list(self._pending_bookmarks))
            self._pending_bookmarks.clear()
        if extra.get("ins"):
            cf.setProperty(P_INS, extra["ins"])
        if extra.get("del"):
            cf.setProperty(P_DEL, extra["del"])
        if self._active_comments:
            cf.setProperty(P_COMMENTS, ",".join(sorted(self._active_comments)))

    def _equation(self, node, ctx: _Ctx, rp_base: RunProps, extra: dict):
        from app.writer import equations

        maths = [node] if node.tag == "{%s}oMath" % M_NS else node.findall("{%s}oMath" % M_NS)
        for om in maths:
            latex = equations.omml_to_latex(om)
            if not latex:
                continue
            try:
                png = equations.render_png(latex)
            except equations.EquationError:
                cf = rp_base.char_format(latex)
                ctx.cur.insertText(equations.plain_text(latex), cf)
                continue
            img = QImage.fromData(png)
            name = self.register_image(png, "png", img)
            imf = QTextImageFormat()
            imf.setName(name)
            imf.setWidth(img.width() / 2.0)   # rendered at 192 dpi, shown at 96
            imf.setHeight(img.height() / 2.0)
            imf.setVerticalAlignment(QTextCharFormat.AlignMiddle)
            imf.setProperty(P_EQUATION, latex)
            ctx.cur.insertImage(imf)

    def _drawing(self, node, ctx: _Ctx, rp: RunProps, extra: dict):
        # text boxes -> extract paragraphs after the current one
        for tb in node.iter(w("txbxContent")):
            for p in tb.findall(w("p")):
                self._extra_paras.append(p)
            return
        blip = None
        for b in node.iter("{%s}blip" % A):
            blip = b
            break
        rid = blip.get(qn("r:embed")) if blip is not None else None
        if rid is None:
            for im in node.iter("{%s}imagedata" % V):
                rid = im.get(qn("r:id"))
                break
        if not rid:
            return
        try:
            part = self.d.part.related_parts[rid]
            data = part.blob
            ext = os.path.splitext(str(part.partname))[1].lstrip(".").lower() or "png"
        except KeyError:
            return
        cx = cy = None
        ext_el = None
        for e in node.iter("{%s}extent" % WP):
            ext_el = e
            break
        if ext_el is not None:
            cx = _int(ext_el.get("cx")) / EMU_PER_PX
            cy = _int(ext_el.get("cy")) / EMU_PER_PX
        img = QImage.fromData(data)
        if img.isNull():
            img = _placeholder_image(int(cx or 200), int(cy or 120), ext.upper())
            buf = _png_bytes(img)
            data, ext = buf, "png"
        if ext == "jpeg":
            ext = "jpg"
        name = self.register_image(data, ext, img)
        fmt = QTextImageFormat()
        fmt.setName(name)
        fmt.setWidth(cx or img.width())
        fmt.setHeight(cy or img.height())
        self._decorate(fmt, extra)
        ctx.cur.insertImage(fmt)

    # ------------------------------------------------------------------ tables
    def _table(self, tbl, ctx: _Ctx):
        grid = [(_int(g.get(w("w")), 0) / TWIP_PER_PX) for g in tbl.findall(f"{w('tblGrid')}/{w('gridCol')}")]
        rows = tbl.findall(w("tr"))
        if not rows:
            return
        layout = []  # per row: list of (col, span, vmerge, tc)
        ncols = len(grid)
        for tr in rows:
            col = 0
            cells = []
            for tc in tr.findall(w("tc")):
                tcpr = tc.find(w("tcPr"))
                span = 1
                vm = None
                if tcpr is not None:
                    gs = tcpr.find(w("gridSpan"))
                    if gs is not None:
                        span = max(1, _int(gs.get(w("val")), 1))
                    vme = tcpr.find(w("vMerge"))
                    if vme is not None:
                        vm = vme.get(w("val")) or "continue"
                cells.append((col, span, vm, tc))
                col += span
            ncols = max(ncols, col)
            layout.append(cells)
        if ncols == 0:
            return
        tf = QTextTableFormat()
        tf.setBorder(1)
        tf.setBorderStyle(QTextFrameFormat.BorderStyle_Solid)
        tf.setBorderBrush(QBrush(QColor("#000000")))
        tf.setCellSpacing(0)
        tf.setCellPadding(4)
        try:
            tf.setBorderCollapse(True)
        except AttributeError:
            pass
        if grid and len(grid) == ncols and all(g > 0 for g in grid):
            tf.setColumnWidthConstraints([QTextLength(QTextLength.FixedLength, g) for g in grid])
        else:
            tf.setWidth(QTextLength(QTextLength.PercentageLength, 100))
        jc = tbl.find(f"{w('tblPr')}/{w('jc')}")
        if jc is not None and jc.get(w("val")) == "center":
            tf.setAlignment(Qt.AlignHCenter)
        bidi = tbl.find(f"{w('tblPr')}/{w('bidiVisual')}")
        if bidi is not None:
            tf.setLayoutDirection(Qt.RightToLeft)
        cur = ctx.cur
        table = cur.insertTable(len(rows), ncols, tf)
        # merges first (on an empty table)
        for r, cells in enumerate(layout):
            for (c, span, vm, _tc) in cells:
                if vm == "continue":
                    continue
                rowspan = 1
                if vm == "restart":
                    rr = r + 1
                    while rr < len(layout):
                        nxt = [x for x in layout[rr] if x[0] == c]
                        if nxt and nxt[0][2] == "continue":
                            rowspan += 1
                            rr += 1
                        else:
                            break
                if (span > 1 or rowspan > 1) and c < ncols:
                    table.mergeCells(r, c, rowspan, min(span, ncols - c))
        for r, cells in enumerate(layout):
            for (c, span, vm, tc) in cells:
                if vm == "continue" or c >= ncols:
                    continue
                cell = table.cellAt(r, c)
                tcpr = tc.find(w("tcPr"))
                if tcpr is not None:
                    shd = tcpr.find(w("shd"))
                    if shd is not None and shd.get(w("fill")) and shd.get(w("fill")).lower() != "auto":
                        cfmt = cell.format().toTableCellFormat()
                        cfmt.setBackground(QColor("#" + shd.get(w("fill"))))
                        cell.setFormat(cfmt)
                cctx = _Ctx(cell.firstCursorPosition(), True)
                for ch in tc.iterchildren():
                    if ch.tag in (w("p"), w("tbl"), w("sdt")):
                        self._block_level(ch, cctx)
        ctx.cur = table.lastCursorPosition()
        ctx.cur.movePosition(QTextCursor.NextBlock)
        ctx.reuse = True


def _png_bytes(img: QImage) -> bytes:
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QIODevice.WriteOnly)
    img.save(buf, "PNG")
    return bytes(ba.data())


def _placeholder_image(wd: int, ht: int, label: str) -> QImage:
    img = QImage(max(40, wd), max(30, ht), QImage.Format_RGB32)
    img.fill(QColor("#f1f3f4"))
    p = QPainter(img)
    p.setPen(QColor("#9aa0a6"))
    p.drawRect(0, 0, img.width() - 1, img.height() - 1)
    p.drawText(img.rect(), Qt.AlignCenter, f"[{label} image]")
    p.end()
    return img


_HF_FIELDS = {"{PAGE}": (" PAGE ", "1"), "{NUMPAGES}": (" NUMPAGES ", "1"),
              "{DATE}": (' DATE \\@ "yyyy-MM-dd" ', ""), "{TITLE}": (" TITLE ", "")}


def _hf_token(instr: str) -> str | None:
    key = instr.strip().upper().split(" ")[0] if instr.strip() else ""
    return {"PAGE": "{PAGE}", "NUMPAGES": "{NUMPAGES}", "SECTIONPAGES": "{NUMPAGES}", "DATE": "{DATE}",
            "TIME": "{DATE}", "TITLE": "{TITLE}"}.get(key)


# ============================================================================ watermark (VML text path)
_WATERMARK_XML = """<w:r xmlns:w="%(w)s" xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:w10="urn:schemas-microsoft-com:office:word">
<w:rPr><w:noProof/></w:rPr><w:pict>
<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">
<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/>
<v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/>
<v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/>
<v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>
<v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>
<v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles>
<o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>
<v:shape id="MyOfficeWaterMark" o:spid="_x0000_s2049" type="#_x0000_t136"
 style="position:absolute;margin-left:0;margin-top:0;width:468pt;height:117pt;rotation:315;z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
 o:allowincell="f" fillcolor="silver" stroked="f"><v:fill opacity=".5"/>
<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="%(text)s"/>
<w10:wrap anchorx="margin" anchory="margin"/></v:shape></w:pict></w:r>"""


def _watermark_run(text: str):
    from xml.sax.saxutils import quoteattr

    esc = quoteattr(text)[1:-1]
    return etree.fromstring(_WATERMARK_XML % {"w": W, "text": esc})


# ============================================================================ exporter
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"]
_RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def _el(tag: str, **attrs):
    e = OxmlElement(f"w:{tag}")
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _ordered(parent_tag: str, children: dict, order: list):
    parent = OxmlElement(f"w:{parent_tag}")
    for name in order:
        if name in children and children[name] is not None:
            val = children[name]
            for item in (val if isinstance(val, list) else [val]):
                parent.append(item)
    return parent


class DocxExporter:
    def __init__(self, qdoc: QTextDocument, page: PageSetup, hf: HeaderFooter, comments: dict,
                 image_bytes, title: str = "", author: str = "", sources=None, citation_style: str = "APA",
                 watermark: str = "", protection: dict | None = None):
        self.watermark = watermark or ""
        self.protection = protection or {}
        self._footnotes: list[str] = []
        self.qdoc = qdoc
        self.page = page
        self.hf = hf
        self.comments = comments
        self.image_bytes = image_bytes
        self.title = title
        self.author = author
        self.sources = sources or []
        self.citation_style = citation_style
        self._rev_id = 1000
        self._comment_runs: dict[str, list] = {}
        self._list_nums: dict[int, str] = {}
        self._toc_open = False

    # ------------------------------------------------------------------ public
    def save(self, path: str) -> None:
        d = docx.Document()
        self.d = d
        self._prepare_styles()
        body = d.element.body
        for child in list(body):
            if child.tag != w("sectPr"):
                body.remove(child)
        self._sect = body.find(w("sectPr"))
        self._export_frame_into(self.qdoc.rootFrame(), self._append_body)
        self._close_toc_if_needed(force=True)
        self._page_setup()
        self._headers_footers()
        self._add_comments()
        self._add_footnotes()
        self._add_protection()
        self._custom_data()
        cp = d.core_properties
        cp.title = self.title or ""
        cp.author = self.author or ""
        cp.last_modified_by = self.author or ""
        cp.modified = datetime.now()
        folder = os.path.dirname(os.path.abspath(path)) or "."
        fd, tmp = tempfile.mkstemp(suffix=".docx", dir=folder)
        os.close(fd)
        try:
            d.save(tmp)
            from app.services.fileutil import finish_atomic

            finish_atomic(tmp, path)
        except Exception:
            try:
                os.unlink(tmp)
            except OSError:
                pass
            raise

    # ------------------------------------------------------------------ styles
    def _prepare_styles(self):
        df = self.qdoc.defaultFont()
        normal = self.d.styles["Normal"]
        normal.font.name = df.family()
        normal.font.size = Pt(df.pointSizeF() if df.pointSizeF() > 0 else 11)
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            fonts.set(qn("w:cs"), df.family())
            fonts.set(qn("w:eastAsia"), df.family())
        pf = normal.paragraph_format
        pf.space_after = Pt(0)
        self._style_ids = {}
        for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5",
                     "Heading 6", "Quote", "List Paragraph", "List Bullet", "List Number", "Normal"):
            try:
                self._style_ids[name] = self.d.styles[name].style_id
            except KeyError:
                pass
        try:  # our Title style has no rule underneath; drop the template's bottom border
            tppr = self.d.styles["Title"].element.find(qn("w:pPr"))
            if tppr is not None:
                for bdr in tppr.findall(qn("w:pBdr")):
                    tppr.remove(bdr)
        except KeyError:
            pass
        try:
            code = self.d.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
            code.base_style = normal
            code.font.name = "Consolas"
            self._style_ids["Code"] = code.style_id
        except Exception:
            pass

    def _numbering_ids(self):
        if hasattr(self, "_abstract_ids"):
            return self._abstract_ids
        self._abstract_ids = {}
        try:
            numbering = self.d.part.numbering_part.element
            for name, key in (("List Bullet", "bullet"), ("List Number", "number")):
                st = self.d.styles[name].element
                numpr = st.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
                nid = numpr.find(qn("w:numId")).get(qn("w:val"))
                for num in numbering.findall(qn("w:num")):
                    if num.get(qn("w:numId")) == nid:
                        self._abstract_ids[key] = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            self._numbering_el = numbering
        except Exception as exc:
            log.warning("numbering setup failed: %s", exc)
        return self._abstract_ids

    def _num_id_for_list(self, lst) -> str | None:
        key = lst.objectIndex()
        if key in self._list_nums:
            return self._list_nums[key]
        ids = self._numbering_ids()
        style = lst.format().style()
        bullet = style in (QTextListFormat.ListDisc, QTextListFormat.ListCircle, QTextListFormat.ListSquare)
        abstract = ids.get("bullet" if bullet else "number")
        if abstract is None:
            return None
        numbering = self._numbering_el
        existing = [_int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        new_id = str(max(existing + [0]) + 1)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), new_id)
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), abstract)
        num.append(an)
        if not bullet:
            for lvl in range(9):
                ov = OxmlElement("w:lvlOverride")
                ov.set(qn("w:ilvl"), str(lvl))
                so = OxmlElement("w:startOverride")
                so.set(qn("w:val"), "1")
                ov.append(so)
                num.append(ov)
        numbering.append(num)
        self._list_nums[key] = new_id
        return new_id

    # ------------------------------------------------------------------ frames / blocks
    def _append_body(self, el):
        self._sect.addprevious(el)

    def _export_frame_into(self, frame, append):
        it = frame.begin()
        while not it.atEnd():
            child = it.currentFrame()
            block = it.currentBlock()
            if child is not None:
                if isinstance(child, QTextTable):
                    self._close_toc_if_needed(force=True)
                    append(self._table(child))
                else:
                    self._export_frame_into(child, append)
            elif block.isValid():
                append(self._paragraph(block))
            it += 1

    def _close_toc_if_needed(self, force=False, next_is_toc=False):
        if self._toc_open and (force or not next_is_toc):
            if getattr(self, "_last_toc_p", None) is not None:
                r = _el("r")
                r.append(_el("fldChar", fldCharType="end"))
                self._last_toc_p.append(r)
            self._toc_open = False

    def _paragraph(self, block):
        bf = block.blockFormat()
        is_toc = bool(bf.property(P_TOC)) if bf.hasProperty(P_TOC) else False
        if not is_toc:
            self._close_toc_if_needed(force=True)
        p = _el("p")
        ppr = {}
        style_name = str(bf.property(P_STYLE)) if bf.hasProperty(P_STYLE) else ""
        if bf.headingLevel() > 0:
            style_name = style_for_heading(bf.headingLevel()) if bf.headingLevel() <= 4 else f"Heading {bf.headingLevel()}"
        lst = block.textList()
        if lst is not None and style_name in ("", "Normal"):
            style_name = "List Paragraph"
        sid = self._style_ids.get(style_name)
        if sid and style_name != "Normal":
            ppr["pStyle"] = _el("pStyle", val=sid)
        if bf.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore:
            ppr["pageBreakBefore"] = _el("pageBreakBefore")
        if lst is not None:
            nid = self._num_id_for_list(lst)
            if nid:
                numpr = _el("numPr")
                numpr.append(_el("ilvl", val=max(0, lst.format().indent() - 1)))
                numpr.append(_el("numId", val=nid))
                ppr["numPr"] = numpr
        if bf.hasProperty(QTextFormat.BlockTrailingHorizontalRulerWidth):
            bdr = _el("pBdr")
            bdr.append(_el("bottom", val="single", sz="6", space="1", color="auto"))
            ppr["pBdr"] = bdr
        if bf.background().style() != Qt.NoBrush:
            ppr["shd"] = _el("shd", val="clear", color="auto", fill=color_hex(bf.background()) or "auto")
        rtl = bf.layoutDirection() == Qt.RightToLeft or (
            bf.layoutDirection() == Qt.LayoutDirectionAuto and block.textDirection() == Qt.RightToLeft)
        if rtl:
            ppr["bidi"] = _el("bidi")
        spacing = {}
        if bf.topMargin():
            spacing["before"] = int(round(bf.topMargin() * TWIP_PER_PX))
        if bf.bottomMargin():
            spacing["after"] = int(round(bf.bottomMargin() * TWIP_PER_PX))
        else:
            spacing["after"] = 0
        lht = bf.lineHeightType()
        if lht == QTextBlockFormat.ProportionalHeight.value and bf.lineHeight() not in (0, 100):
            spacing["line"] = int(round(bf.lineHeight() * 2.4))
            spacing["lineRule"] = "auto"
        elif lht == QTextBlockFormat.FixedHeight.value:
            spacing["line"] = int(round(bf.lineHeight() * TWIP_PER_PX))
            spacing["lineRule"] = "exact"
        elif lht == QTextBlockFormat.MinimumHeight.value:
            spacing["line"] = int(round(bf.lineHeight() * TWIP_PER_PX))
            spacing["lineRule"] = "atLeast"
        if spacing:
            ppr["spacing"] = _el("spacing", **spacing)
        ind = {}
        if bf.leftMargin() and lst is None:
            ind["left"] = int(round(bf.leftMargin() * TWIP_PER_PX))
        if bf.rightMargin():
            ind["right"] = int(round(bf.rightMargin() * TWIP_PER_PX))
        if bf.textIndent() > 0:
            ind["firstLine"] = int(round(bf.textIndent() * TWIP_PER_PX))
        elif bf.textIndent() < 0:
            ind["hanging"] = int(round(-bf.textIndent() * TWIP_PER_PX))
        if ind:
            ppr["ind"] = _el("ind", **ind)
        al = bf.alignment() & Qt.AlignHorizontal_Mask
        if al & Qt.AlignHCenter:
            ppr["jc"] = _el("jc", val="center")
        elif al & Qt.AlignJustify:
            ppr["jc"] = _el("jc", val="both")
        elif al & Qt.AlignRight:
            ppr["jc"] = _el("jc", val="right")
        if bf.headingLevel() > 0:
            ppr["outlineLvl"] = _el("outlineLvl", val=bf.headingLevel() - 1)
        if ppr:
            p.append(_ordered("pPr", ppr, _PPR_ORDER))
        if is_toc and not self._toc_open:
            for el in (_el("fldChar", fldCharType="begin"),):
                r = _el("r")
                r.append(el)
                p.append(r)
            r = _el("r")
            it = _el("instrText")
            it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            it.text = ' TOC \\o "1-3" \\h \\z \\u '
            r.append(it)
            p.append(r)
            r = _el("r")
            r.append(_el("fldChar", fldCharType="separate"))
            p.append(r)
            self._toc_open = True
        self._runs(block, p)
        if is_toc:
            self._last_toc_p = p
        return p

    def _runs(self, block, p):
        items = []
        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            it += 1
            if not frag.isValid():
                continue
            fmt = frag.charFormat()
            text = frag.text()
            if fmt.isImageFormat():
                imf = fmt.toImageFormat()
                for _ in text:
                    r = self._equation_xml(imf) if imf.hasProperty(P_EQUATION) else None
                    if r is None:
                        r = self._image_run(imf)
                    if r is not None:
                        items.append((r, fmt))
                continue
            if fmt.hasProperty(P_FOOTNOTE):
                items.append((self._footnote_ref(str(fmt.property(P_FOOTNOTE))), fmt))
                continue
            for r in self._text_runs(text, fmt):
                items.append((r, fmt))
        # group into hyperlink / revision wrappers
        i = 0
        while i < len(items):
            r, fmt = items[i]
            href = fmt.anchorHref() if fmt.isAnchor() else ""
            if href:
                j = i
                group = []
                while j < len(items) and (items[j][1].anchorHref() if items[j][1].isAnchor() else "") == href:
                    group.append(items[j])
                    j += 1
                hl = _el("hyperlink")
                if href.startswith("#"):
                    hl.set(qn("w:anchor"), href[1:])
                else:
                    rid = self.d.part.relate_to(href, RT.HYPERLINK, is_external=True)
                    hl.set(qn("r:id"), rid)
                self._emit_revision_groups(group, hl)
                p.append(hl)
                i = j
            else:
                j = i
                group = []
                while j < len(items) and not (items[j][1].isAnchor() and items[j][1].anchorHref()):
                    group.append(items[j])
                    j += 1
                self._emit_revision_groups(group, p)
                i = j
        # bookmarks (anchor names) -> wrap whole paragraph start
        names = []
        it = block.begin()
        while not it.atEnd():
            f = it.fragment()
            it += 1
            if f.isValid() and f.charFormat().isAnchor():
                names.extend(f.charFormat().anchorNames() or [])
        for n in names:
            bid = str(self._next_id())
            bs = _el("bookmarkStart", id=bid, name=n)
            be = _el("bookmarkEnd", id=bid)
            ppr = p.find(qn("w:pPr"))
            if ppr is not None:
                ppr.addnext(bs)
            else:
                p.insert(0, bs)
            p.append(be)

    def _next_id(self) -> int:
        self._rev_id += 1
        return self._rev_id

    def _emit_revision_groups(self, group, parent):
        i = 0
        while i < len(group):
            r, fmt = group[i]
            ins = str(fmt.property(P_INS)) if fmt.hasProperty(P_INS) else ""
            dele = str(fmt.property(P_DEL)) if fmt.hasProperty(P_DEL) else ""
            if not ins and not dele:
                parent.append(r)
                self._track_comment(r, fmt)
                i += 1
                continue
            key = ("ins", ins) if ins else ("del", dele)
            j = i
            wrapper = _el(key[0], id=self._next_id())
            author, date = parse_stamp(key[1])
            wrapper.set(qn("w:author"), author or "Unknown")
            if date:
                wrapper.set(qn("w:date"), date if date.endswith("Z") else (date[:19] + "Z"))
            while j < len(group):
                r2, f2 = group[j]
                ins2 = str(f2.property(P_INS)) if f2.hasProperty(P_INS) else ""
                del2 = str(f2.property(P_DEL)) if f2.hasProperty(P_DEL) else ""
                if (("ins", ins2) if ins2 else ("del", del2)) != key or not (ins2 or del2):
                    break
                if key[0] == "del":
                    for t in r2.findall(qn("w:t")):
                        t.tag = qn("w:delText")
                wrapper.append(r2)
                self._track_comment(r2, f2)
                j += 1
            parent.append(wrapper)
            i = j

    def _track_comment(self, r, fmt):
        ids = [x for x in str(fmt.property(P_COMMENTS) or "").split(",") if x] if fmt.hasProperty(P_COMMENTS) else []
        for cid in ids:
            self._comment_runs.setdefault(cid, []).append(r)

    def _rpr(self, fmt: QTextCharFormat, text: str):
        ch = {}
        fams = fmt.fontFamilies() or []
        fam = fams[0] if fams else None
        complex_script = has_rtl(text) or has_devanagari(text)
        if fam:
            ch["rFonts"] = _el("rFonts", ascii=fam, hAnsi=fam, cs=fam, eastAsia=fam)
        if fmt.hasProperty(QTextFormat.FontWeight):
            if fmt.fontWeight() >= 600:
                ch["b"] = _el("b")
                ch["bCs"] = _el("bCs")
            else:
                ch["b"] = _el("b", val="0")
                ch["bCs"] = _el("bCs", val="0")
        if fmt.hasProperty(QTextFormat.FontItalic):
            if fmt.fontItalic():
                ch["i"] = _el("i")
                ch["iCs"] = _el("iCs")
            else:
                ch["i"] = _el("i", val="0")
        cap = fmt.fontCapitalization()
        if cap == QFont.AllUppercase:
            ch["caps"] = _el("caps")
        elif cap == QFont.SmallCaps:
            ch["smallCaps"] = _el("smallCaps")
        if fmt.fontStrikeOut():
            ch["strike"] = _el("strike")
        col = color_hex(fmt.foreground()) if fmt.hasProperty(QTextFormat.ForegroundBrush) else None
        if col:
            ch["color"] = _el("color", val=col)
        if fmt.fontLetterSpacingType() == QFont.AbsoluteSpacing and fmt.fontLetterSpacing():
            ch["spacing"] = _el("spacing", val=int(round(fmt.fontLetterSpacing() * TWIP_PER_PX)))
        if fmt.fontPointSize() > 0:
            hp = int(round(fmt.fontPointSize() * 2))
            ch["sz"] = _el("sz", val=hp)
            ch["szCs"] = _el("szCs", val=hp)
        if fmt.fontUnderline() and not fmt.isAnchor():
            ch["u"] = _el("u", val="single")
        elif fmt.isAnchor() and fmt.anchorHref():
            ch["u"] = _el("u", val="single")
        bg = color_hex(fmt.background()) if fmt.hasProperty(QTextFormat.BackgroundBrush) else None
        if bg:
            ch["shd"] = _el("shd", val="clear", color="auto", fill=bg)
        va = fmt.verticalAlignment()
        if va == QTextCharFormat.AlignSuperScript:
            ch["vertAlign"] = _el("vertAlign", val="superscript")
        elif va == QTextCharFormat.AlignSubScript:
            ch["vertAlign"] = _el("vertAlign", val="subscript")
        if has_rtl(text):
            ch["rtl"] = _el("rtl")
        if complex_script:
            ch["cs"] = _el("cs")
            lang = "ur-PK" if has_arabic(text) else ("hi-IN" if has_devanagari(text) else "ar-SA")
            ch["lang"] = _el("lang", bidi=lang)
        return _ordered("rPr", ch, _RPR_ORDER) if ch else None

    def _text_runs(self, text: str, fmt: QTextCharFormat):
        if not text:
            return []
        rpr = self._rpr(fmt, text)
        runs = []
        r = _el("r")
        if rpr is not None:
            r.append(rpr)
        buf = []

        def flush():
            if buf:
                t = _el("t")
                t.text = "".join(buf)
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)
                buf.clear()

        for chx in text:
            if chx == "\t":
                flush()
                r.append(_el("tab"))
            elif chx in ("\u2028", "\n", "\r"):
                flush()
                r.append(_el("br"))
            elif chx == "\u2011":
                flush()
                r.append(_el("noBreakHyphen"))
            elif chx in ("\ufffc", "\u2029"):
                continue
            elif ord(chx) < 0x20:
                continue
            else:
                buf.append(chx)
        flush()
        runs.append(r)
        return runs

    def _image_run(self, imf: QTextImageFormat):
        data = self.image_bytes(imf.name())
        if not data:
            return None
        blob, ext = data
        if ext.lower() not in ("png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"):
            img = QImage.fromData(blob)
            if img.isNull():
                return None
            blob, ext = _png_bytes(img), "png"
        img = QImage.fromData(blob)
        wpx = imf.width() if imf.width() > 0 else img.width()
        hpx = imf.height() if imf.height() > 0 else img.height()
        try:
            inline = self.d.part.new_pic_inline(io.BytesIO(blob), Emu(int(wpx * EMU_PER_PX)), Emu(int(hpx * EMU_PER_PX)))
        except Exception as exc:
            log.warning("image export failed (%s): %s", imf.name(), exc)
            return None
        r = _el("r")
        drawing = _el("drawing")
        drawing.append(inline)
        r.append(drawing)
        return r

    # ------------------------------------------------------------------ equations / footnotes / protection
    def _equation_xml(self, imf: QTextImageFormat):
        from app.writer import equations

        try:
            return equations.latex_to_omml(str(imf.property(P_EQUATION)))
        except equations.EquationError as exc:
            log.warning("equation kept as picture: %s", exc)
            return None

    def _footnote_ref(self, text: str):
        self._footnotes.append(text)
        r = _el("r")
        rpr = _el("rPr")
        rpr.append(_el("rStyle", val="FootnoteReference"))
        r.append(rpr)
        r.append(_el("footnoteReference", id=str(len(self._footnotes))))
        return r

    def _note_styles(self):
        styles = self.d.styles
        names = {s.name for s in styles}
        if "footnote text" not in {n.lower() for n in names}:
            st = styles.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            st.font.size = Pt(9)
            st.element.set(qn("w:styleId"), "FootnoteText")
        if "footnote reference" not in {n.lower() for n in names}:
            st = styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
            st.font.superscript = True
            st.element.set(qn("w:styleId"), "FootnoteReference")

    def _add_footnotes(self):
        if not self._footnotes:
            return
        self._note_styles()
        W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        root = etree.Element("{%s}footnotes" % W, nsmap={"w": W, "r": "http://schemas.openxmlformats.org/"
                                                                      "officeDocument/2006/relationships"})
        for nid, kind in (("-1", "separator"), ("0", "continuationSeparator")):
            fn = etree.SubElement(root, w("footnote"))
            fn.set(w("type"), kind)
            fn.set(w("id"), nid)
            p_el = etree.SubElement(fn, w("p"))
            ppr = etree.SubElement(p_el, w("pPr"))
            sp = etree.SubElement(ppr, w("spacing"))
            sp.set(w("after"), "0")
            sp.set(w("line"), "240")
            sp.set(w("lineRule"), "auto")
            etree.SubElement(etree.SubElement(p_el, w("r")), w(kind))
        for i, text in enumerate(self._footnotes, 1):
            fn = etree.SubElement(root, w("footnote"))
            fn.set(w("id"), str(i))
            for n, line in enumerate((text or " ").split("\n")):
                p_el = etree.SubElement(fn, w("p"))
                etree.SubElement(etree.SubElement(p_el, w("pPr")), w("pStyle")).set(w("val"), "FootnoteText")
                if n == 0:
                    r = etree.SubElement(p_el, w("r"))
                    etree.SubElement(etree.SubElement(r, w("rPr")), w("rStyle")).set(w("val"), "FootnoteReference")
                    etree.SubElement(r, w("footnoteRef"))
                    r2 = etree.SubElement(p_el, w("r"))
                    t = etree.SubElement(r2, w("t"))
                    t.text = " " + line
                else:
                    t = etree.SubElement(etree.SubElement(p_el, w("r")), w("t"))
                    t.text = line
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        part = Part(PackURI("/word/footnotes.xml"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", blob,
                    self.d.part.package)
        self.d.part.relate_to(part, RT.FOOTNOTES)

    def _add_protection(self):
        if not self.protection:
            return
        settings = self.d.settings.element
        for old in settings.findall(qn("w:documentProtection")):
            settings.remove(old)
        el = OxmlElement("w:documentProtection")
        for k, v in self.protection.items():
            el.set(qn(f"w:{k}"), str(v))
        # schema order: documentProtection comes before these elements
        later = {"autoFormatOverride", "styleLockTheme", "styleLockQFSet", "defaultTabStop", "autoHyphenation",
                 "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps", "showEnvelope",
                 "summaryLength", "clickAndTypeStyle", "defaultTableStyle", "evenAndOddHeaders",
                 "bookFoldRevPrinting", "bookFoldPrinting", "bookFoldPrintingSheets", "drawingGridHorizontalSpacing",
                 "drawingGridVerticalSpacing", "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery",
                 "doNotUseMarginsForDrawingGridOrigin", "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin",
                 "doNotShadeFormData", "noPunctuationKerning", "characterSpacingControl", "printTwoOnOne",
                 "strictFirstAndLastChars", "noLineBreaksAfter", "noLineBreaksBefore", "savePreviewPicture",
                 "doNotValidateAgainstSchema", "saveInvalidXml", "ignoreMixedContent", "alwaysShowPlaceholderText",
                 "doNotDemarcateInvalidXml", "saveXmlDataOnly", "useXSLTWhenSaving", "saveThroughXslt",
                 "showXMLTags", "alwaysMergeEmptyNamespace", "updateFields", "hdrShapeDefaults", "footnotePr",
                 "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema", "themeFontLang",
                 "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                 "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults",
                 "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
        anchor = next((c for c in settings if etree.QName(c).localname in later), None)
        if anchor is not None:
            anchor.addprevious(el)
        else:
            settings.append(el)

    # ------------------------------------------------------------------ tables
    def _table(self, table: QTextTable):
        rows, cols = table.rows(), table.columns()
        tbl = _el("tbl")
        tblpr = _el("tblPr")
        sid = None
        try:
            sid = self.d.styles["Table Grid"].style_id
        except KeyError:
            pass
        if sid:
            tblpr.append(_el("tblStyle", val=sid))
        tf = table.format()
        tblpr.append(_el("tblW", w="0", type="auto"))
        if tf.alignment() & Qt.AlignHCenter:
            tblpr.append(_el("jc", val="center"))
        borders = _el("tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(_el(side, val="single", sz="4", space="0", color="000000" if tf.border() > 0 else "auto"))
        tblpr.append(borders)
        if tf.layoutDirection() == Qt.RightToLeft:
            tblpr.insert(0, _el("bidiVisual"))
        tblpr.append(_el("tblLook", val="04A0", firstRow="1", lastRow="0", firstColumn="1", lastColumn="0",
                         noHBand="0", noVBand="1"))
        tbl.append(tblpr)
        grid = _el("tblGrid")
        constraints = tf.columnWidthConstraints()
        text_w = self.page.page_px()[0] - self.page.margins_px()[1] - self.page.margins_px()[3]
        widths = []
        for c in range(cols):
            wpx = text_w / max(1, cols)
            if c < len(constraints):
                ln = constraints[c]
                if ln.type() == QTextLength.FixedLength:
                    wpx = ln.rawValue()
                elif ln.type() == QTextLength.PercentageLength:
                    wpx = text_w * ln.rawValue() / 100.0
            widths.append(int(round(wpx * TWIP_PER_PX)))
            grid.append(_el("gridCol", w=widths[-1]))
        tbl.append(grid)
        for r in range(rows):
            tr = _el("tr")
            c = 0
            while c < cols:
                cell = table.cellAt(r, c)
                span = max(1, cell.columnSpan())
                tc = _el("tc")
                tcpr = _el("tcPr")
                tcpr.append(_el("tcW", w=sum(widths[c:c + span]), type="dxa"))
                if span > 1:
                    tcpr.append(_el("gridSpan", val=span))
                if cell.rowSpan() > 1 or cell.row() != r:
                    tcpr.append(_el("vMerge", val="restart") if cell.row() == r else _el("vMerge"))
                cf = cell.format()
                if cf.background().style() != Qt.NoBrush:
                    tcpr.append(_el("shd", val="clear", color="auto", fill=color_hex(cf.background()) or "auto"))
                tc.append(tcpr)
                if cell.row() == r and cell.column() == c:
                    elements = []
                    self._export_cell(cell, elements.append)
                    for el in elements:
                        tc.append(el)
                if tc.find(qn("w:p")) is None or tc[-1].tag != qn("w:p"):
                    tc.append(_el("p"))
                tr.append(tc)
                c += span
            tbl.append(tr)
        return tbl

    def _export_cell(self, cell, append):
        it = cell.begin()
        while not it.atEnd():
            child = it.currentFrame()
            block = it.currentBlock()
            if child is not None and isinstance(child, QTextTable):
                append(self._table(child))
            elif block.isValid():
                append(self._paragraph(block))
            it += 1

    # ------------------------------------------------------------------ section / header / footer
    def _page_setup(self):
        sec = self.d.sections[0]
        wmm, hmm = self.page.page_mm()
        sec.orientation = WD_ORIENT.LANDSCAPE if self.page.orientation == "landscape" else WD_ORIENT.PORTRAIT
        sec.page_width = Mm(wmm)
        sec.page_height = Mm(hmm)
        t, r, b, l = self.page.margins_mm
        sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = Mm(t), Mm(r), Mm(b), Mm(l)
        sec.header_distance = Mm(self.page.header_mm)
        sec.footer_distance = Mm(self.page.footer_mm)

    def _hf_paragraph(self, container, text: str, align: str):
        container.is_linked_to_previous = False
        p = container.paragraphs[0]
        p.alignment = {"left": 0, "center": 1, "right": 2}.get(align, 1)
        pel = p._p
        for chunk in re.split(r"(\{PAGE\}|\{NUMPAGES\}|\{DATE\}|\{TITLE\})", text):
            if not chunk:
                continue
            if chunk in _HF_FIELDS:
                instr, cached = _HF_FIELDS[chunk]
                fs = _el("fldSimple", instr=instr)
                r = _el("r")
                t = _el("t")
                t.text = self.title if chunk == "{TITLE}" else (
                    datetime.now().strftime("%Y-%m-%d") if chunk == "{DATE}" else cached)
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)
                fs.append(r)
                pel.append(fs)
            else:
                for r in self._text_runs(chunk, QTextCharFormat()):
                    pel.append(r)

    def _headers_footers(self):
        sec = self.d.sections[0]
        if self.hf.header:
            self._hf_paragraph(sec.header, self.hf.header, self.hf.header_align)
        if self.hf.footer:
            self._hf_paragraph(sec.footer, self.hf.footer, self.hf.footer_align)
        if self.watermark:
            sec.header.is_linked_to_previous = False
            sec.header.paragraphs[0]._p.append(_watermark_run(self.watermark))

    # ------------------------------------------------------------------ comments
    def _add_comments(self):
        for cid, runs in self._comment_runs.items():
            c = self.comments.get(cid)
            if c is None or not runs:
                continue
            para = None
            try:
                first = Run(runs[0], para)
                last = Run(runs[-1], para)
                text = c.text if hasattr(c, "text") else str(c.get("text", ""))
                author = c.author if hasattr(c, "author") else str(c.get("author", ""))
                initials = getattr(c, "initials", "") or "".join(x[0] for x in author.split()[:2]).upper()
                self.d.add_comment([first, last], text=text, author=author, initials=initials)
            except Exception as exc:
                log.warning("comment export failed: %s", exc)

    # ------------------------------------------------------------------ custom xml (citations)
    def _custom_data(self):
        data = {"sources": self.sources, "citation_style": self.citation_style, "watermark": self.watermark,
                "header_footer": {"header_align": self.hf.header_align, "footer_align": self.hf.footer_align},
                "generator": "MyOffice"}
        root = etree.Element("{%s}data" % CUSTOM_NS, nsmap={"mo": CUSTOM_NS})
        root.text = etree.CDATA(json.dumps(data, ensure_ascii=False))
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        package = self.d.part.package
        used = {str(p.partname) for p in package.iter_parts()}
        n = 1
        while f"/customXml/item{n}.xml" in used or f"/customXml/itemProps{n}.xml" in used:
            n += 1
        item = Part(PackURI(f"/customXml/item{n}.xml"), "application/xml", blob, package)
        props_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="no"?>'
                     f'<ds:datastoreItem ds:itemID="{{{str(uuid.uuid4()).upper()}}}" '
                     'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml">'
                     '<ds:schemaRefs/></ds:datastoreItem>').encode("utf-8")
        props = Part(PackURI(f"/customXml/itemProps{n}.xml"),
                     "application/vnd.openxmlformats-officedocument.customXmlProperties+xml", props_xml, package)
        item.relate_to(props, CUSTOMXMLPROPS_RT)
        self.d.part.relate_to(item, CUSTOMXML_RT)


# ============================================================================ convenience API
def load_docx(path: str, qdoc: QTextDocument, register_image) -> DocxImporter:
    imp = DocxImporter(qdoc, register_image)
    try:
        imp.load(path)
    except DocumentError:
        raise
    except Exception as exc:
        log.exception("DOCX import failed")
        raise DocumentError(f"Could not read the document: {exc}") from exc
    return imp


def save_docx(path: str, qdoc: QTextDocument, page: PageSetup, hf: HeaderFooter, comments: dict, image_bytes,
              title: str = "", author: str = "", sources=None, citation_style: str = "APA",
              watermark: str = "", protection: dict | None = None) -> None:
    DocxExporter(qdoc, page, hf, comments, image_bytes, title, author, sources, citation_style,
                 watermark, protection).save(path)

