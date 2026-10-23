"""Mail merge: «Field» placeholders + CSV / XLSX / JSON data source \u2192 DOCX or PDF output."""
from __future__ import annotations

import csv
import io
import json
import os
import re

from PySide6.QtCore import QRegularExpression, QSizeF, QUrl
from PySide6.QtGui import QImage, QTextCursor, QTextDocument, QTextDocumentFragment, QTextFormat
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QButtonGroup, QDialog, QDialogButtonBox, QFormLayout,
                               QGroupBox, QHBoxLayout, QLabel, QLineEdit, QListWidget, QProgressDialog, QPushButton,
                               QRadioButton, QSpinBox, QTableWidget, QTableWidgetItem, QTextEdit, QVBoxLayout)

from app.services.i18n import _t

FIELD_RE = re.compile(r"\u00ab([^\u00ab\u00bb\n]+)\u00bb")


# ============================================================================ data sources
def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256", "cp1252"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def read_data_source(path: str, max_rows: int = 200000) -> tuple[list[str], list[dict]]:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None) or ()
            headers = [str(h).strip() if h is not None else f"Column{i + 1}" for i, h in enumerate(header)]
            data = []
            for r in rows:
                if r is None or all(v is None or str(v).strip() == "" for v in r):
                    continue
                data.append({h: _fmt_value(r[i] if i < len(r) else None) for i, h in enumerate(headers)})
                if len(data) >= max_rows:
                    break
        finally:
            wb.close()
        return headers, data
    if ext == ".json":
        with open(path, encoding="utf-8-sig") as fh:
            obj = json.load(fh)
        if isinstance(obj, dict):
            obj = obj.get("records") or obj.get("data") or [obj]
        recs = [r for r in obj if isinstance(r, dict)]
        headers = []
        for r in recs:
            for k in r:
                if k not in headers:
                    headers.append(str(k))
        return headers, [{h: _fmt_value(r.get(h)) for h in headers} for r in recs[:max_rows]]
    with open(path, "rb") as fh:
        text = _decode(fh.read())
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header = next(reader, [])
    headers = [h.strip() or f"Column{i + 1}" for i, h in enumerate(header)]
    data = []
    for row in reader:
        if not any(c.strip() for c in row):
            continue
        data.append({h: (row[i].strip() if i < len(row) else "") for i, h in enumerate(headers)})
        if len(data) >= max_rows:
            break
    return headers, data


def _fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "strftime"):
        return v.strftime("%d/%m/%Y") if not getattr(v, "hour", 0) else v.strftime("%d/%m/%Y %H:%M")
    return str(v)


# ============================================================================ merging
def merge_fields_in(doc: QTextDocument) -> list[str]:
    seen = []
    for name in FIELD_RE.findall(doc.toPlainText()):
        if name not in seen:
            seen.append(name)
    return seen


def _prepare(doc: QTextDocument, src: QTextDocument) -> None:
    doc.setUseDesignMetrics(True)
    doc.setDocumentMargin(0)
    doc.setDefaultFont(src.defaultFont())
    doc.setPageSize(src.pageSize() if src.pageSize().height() > 0 else QSizeF(794, 1123))
    doc.rootFrame().setFrameFormat(src.rootFrame().frameFormat())


def merged_clone(src: QTextDocument, record: dict) -> QTextDocument:
    clone = src.clone()
    _prepare(clone, src)
    rx = QRegularExpression("\u00ab([^\u00ab\u00bb\\n]+)\u00bb")
    cur = QTextCursor(clone)
    cur.beginEditBlock()
    pos = QTextCursor(clone)
    while True:
        found = clone.find(rx, pos)
        if found.isNull():
            break
        name = found.selectedText()[1:-1]
        value = record.get(name, record.get(name.strip(), ""))
        fmt = found.charFormat()
        found.insertText(str(value), fmt)
        pos = found
    cur.endEditBlock()
    return clone


def build_combined(src: QTextDocument, records: list[dict], images: dict, progress=None) -> QTextDocument:
    combined = QTextDocument()
    _prepare(combined, src)
    for name, (data, _ext) in images.items():
        combined.addResource(QTextDocument.ImageResource, QUrl(name), QImage.fromData(data))
    cur = QTextCursor(combined)
    for i, rec in enumerate(records):
        if progress and not progress(i, len(records)):
            break
        clone = merged_clone(src, rec)
        start = cur.position()
        if i > 0:
            cur.insertBlock()
            start = cur.position()
        cur.insertFragment(QTextDocumentFragment(clone))
        if i > 0:
            b = combined.findBlock(start)
            c2 = QTextCursor(b)
            bf = b.blockFormat()
            bf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
            c2.setBlockFormat(bf)
    return combined


def safe_filename(pattern: str, record: dict, index: int) -> str:
    name = re.sub(r"\{([^{}]+)\}", lambda m: str(record.get(m.group(1), "")), pattern) or f"record_{index}"
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name).strip(" .") or f"record_{index}"
    return name[:150]


# ============================================================================ dialog
class MailMergeDialog(QDialog):
    def __init__(self, writer_doc):
        super().__init__(writer_doc)
        self.wd = writer_doc
        self.setWindowTitle(_t("Mail Merge"))
        self.resize(860, 620)
        self.headers: list[str] = []
        self.records: list[dict] = []
        lay = QVBoxLayout(self)
        # data source
        g1 = QGroupBox(_t("1. Data source (CSV, Excel .xlsx or JSON)"))
        l1 = QVBoxLayout(g1)
        row = QHBoxLayout()
        self.path = QLineEdit()
        self.path.setReadOnly(True)
        b = QPushButton(_t("Browse\u2026"))
        b.clicked.connect(self.browse)
        row.addWidget(self.path, 1)
        row.addWidget(b)
        l1.addLayout(row)
        self.table = QTableWidget(0, 0)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setMaximumHeight(170)
        l1.addWidget(self.table)
        lay.addWidget(g1)
        # fields
        mid = QHBoxLayout()
        g2 = QGroupBox(_t("2. Insert merge fields into the document"))
        l2 = QVBoxLayout(g2)
        self.fields = QListWidget()
        self.fields.itemDoubleClicked.connect(lambda it: self.insert_field(it.text()))
        l2.addWidget(self.fields)
        ins = QPushButton(_t("Insert Field at Cursor"))
        ins.clicked.connect(lambda: self.fields.currentItem() and self.insert_field(self.fields.currentItem().text()))
        l2.addWidget(ins)
        self.used = QLabel("")
        self.used.setWordWrap(True)
        self.used.setObjectName("Muted")
        l2.addWidget(self.used)
        mid.addWidget(g2, 1)
        g3 = QGroupBox(_t("3. Preview"))
        l3 = QVBoxLayout(g3)
        pr = QHBoxLayout()
        pr.addWidget(QLabel(_t("Record:")))
        self.rec = QSpinBox()
        self.rec.setMinimum(1)
        self.rec.valueChanged.connect(self.preview)
        pr.addWidget(self.rec)
        pr.addStretch(1)
        l3.addLayout(pr)
        self.preview_box = QTextEdit()
        self.preview_box.setReadOnly(True)
        l3.addWidget(self.preview_box)
        mid.addWidget(g3, 2)
        lay.addLayout(mid, 1)
        # output
        g4 = QGroupBox(_t("4. Output"))
        f4 = QFormLayout(g4)
        self.o_single = QRadioButton(_t("One merged Word document (.docx)"))
        self.o_each = QRadioButton(_t("One .docx file per record"))
        self.o_pdf = QRadioButton(_t("One merged PDF"))
        self.o_single.setChecked(True)
        grp = QButtonGroup(self)
        for rb in (self.o_single, self.o_each, self.o_pdf):
            grp.addButton(rb)
            f4.addRow(rb)
        self.pattern = QLineEdit("letter_{%s}" % "Name")
        f4.addRow(_t("File name pattern ({Field}):"), self.pattern)
        rng = QHBoxLayout()
        self.r_from = QSpinBox()
        self.r_to = QSpinBox()
        for s in (self.r_from, self.r_to):
            s.setMinimum(1)
        rng.addWidget(self.r_from)
        rng.addWidget(QLabel(_t("to")))
        rng.addWidget(self.r_to)
        rng.addStretch(1)
        f4.addRow(_t("Records:"), rng)
        lay.addWidget(g4)
        bb = QDialogButtonBox()
        run = bb.addButton(_t("Finish && Merge"), QDialogButtonBox.AcceptRole)
        bb.addButton(QDialogButtonBox.Close)
        run.clicked.connect(self.run)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)
        self._update_used()

    # ------------------------------------------------------------------ data
    def browse(self):
        from app.shell.dialogs import get_open_file, show_error

        path = get_open_file(self, _t("Select Data Source"),
                             "Data (*.csv *.txt *.xlsx *.xlsm *.json);;All files (*)")
        if not path:
            return
        try:
            self.load_source(path)
        except Exception as exc:
            show_error(self, _t("Mail Merge"), _t("Could not read the data source: {0}").format(exc))

    def load_source(self, path: str):
        self.headers, self.records = read_data_source(path)
        self.path.setText(path)
        self.table.setColumnCount(len(self.headers))
        self.table.setHorizontalHeaderLabels(self.headers)
        show = self.records[:50]
        self.table.setRowCount(len(show))
        for r, rec in enumerate(show):
            for c, h in enumerate(self.headers):
                self.table.setItem(r, c, QTableWidgetItem(rec.get(h, "")))
        self.fields.clear()
        self.fields.addItems(self.headers)
        n = max(1, len(self.records))
        for s in (self.rec, self.r_from, self.r_to):
            s.setMaximum(n)
        self.r_to.setValue(n)
        if self.headers:
            self.pattern.setText("{%s}" % self.headers[0])
        self.preview()

    def insert_field(self, name: str):
        cur = self.wd.editor.textCursor()
        cur.insertText("\u00ab%s\u00bb" % name)
        self.wd.editor.setTextCursor(cur)
        self._update_used()
        self.preview()

    def _update_used(self):
        used = merge_fields_in(self.wd.editor.document())
        self.used.setText(_t("Fields in document: {0}").format(", ".join(used) if used else _t("none")))

    def preview(self, *_):
        if not self.records:
            self.preview_box.setPlainText(_t("Choose a data source to preview merged records."))
            return
        rec = self.records[self.rec.value() - 1]
        clone = merged_clone(self.wd.editor.document(), rec)
        for name, (data, _e) in self.wd.editor.image_store.items():
            clone.addResource(QTextDocument.ImageResource, QUrl(name), QImage.fromData(data))
        clone.setPageSize(QSizeF(-1, -1))
        self.preview_box.setDocument(clone)
        self._preview_doc = clone

    # ------------------------------------------------------------------ run
    def run(self):
        from app.shell.dialogs import get_directory, get_save_file, show_error, show_info
        from app.writer.docx_io import save_docx
        from app.writer.render import export_pdf

        if not self.records:
            show_error(self, _t("Mail Merge"), _t("Please choose a data source first."))
            return
        if not merge_fields_in(self.wd.editor.document()):
            show_error(self, _t("Mail Merge"), _t("The document contains no merge fields. Insert at least one field."))
            return
        lo, hi = self.r_from.value(), self.r_to.value()
        if hi < lo:
            lo, hi = hi, lo
        recs = self.records[lo - 1:hi]
        ed = self.wd.editor
        prog = QProgressDialog(_t("Merging records\u2026"), _t("Cancel"), 0, len(recs), self)
        prog.setWindowTitle(_t("Mail Merge"))
        prog.setMinimumDuration(0)

        def tick(i, n):
            prog.setValue(i)
            QApplication.processEvents()
            return not prog.wasCanceled()

        try:
            if self.o_each.isChecked():
                folder = get_directory(self, _t("Output folder"))
                if not folder:
                    return
                written = 0
                for i, rec in enumerate(recs, start=lo):
                    if not tick(i - lo, len(recs)):
                        break
                    clone = merged_clone(ed.document(), rec)
                    hf = type(self.wd.editor.hf)(**self.wd.editor.hf.__dict__)
                    hf.header = FIELD_RE.sub(lambda m: rec.get(m.group(1), ""), hf.header)
                    hf.footer = FIELD_RE.sub(lambda m: rec.get(m.group(1), ""), hf.footer)
                    name = safe_filename(self.pattern.text(), rec, i)
                    path = os.path.join(folder, name + ".docx")
                    k = 2
                    while os.path.exists(path):
                        path = os.path.join(folder, f"{name}_{k}.docx")
                        k += 1
                    save_docx(path, clone, ed.page, hf, {}, ed.image_bytes, title=name)
                    written += 1
                prog.setValue(len(recs))
                show_info(self, _t("Mail Merge"), _t("{0} document(s) written to:\n{1}").format(written, folder))
                return
            pdf = self.o_pdf.isChecked()
            path = get_save_file(self, _t("Save merged document"), "PDF (*.pdf)" if pdf else "Word (*.docx)",
                                 "merged." + ("pdf" if pdf else "docx"), "pdf" if pdf else "docx")
            if not path:
                return
            combined = build_combined(ed.document(), recs, ed.image_store, tick)
            prog.setValue(len(recs))
            if pdf:
                export_pdf(combined, ed.page, ed.hf, path, title=os.path.basename(path))
            else:
                save_docx(path, combined, ed.page, ed.hf, {}, ed.image_bytes, title=os.path.basename(path))
            show_info(self, _t("Mail Merge"), _t("Merged {0} record(s) into:\n{1}").format(len(recs), path))
        except Exception as exc:
            show_error(self, _t("Mail Merge"), _t("Mail merge failed: {0}").format(exc))
        finally:
            prog.close()
