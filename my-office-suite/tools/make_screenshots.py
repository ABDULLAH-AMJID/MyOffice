"""Capture screenshots of every module and all 32 tools into docs/screenshots (used by TESTING.md).

Run:  QT_QPA_PLATFORM=offscreen python tools/make_screenshots.py
"""
from __future__ import annotations

import csv
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QT_LOGGING_RULES", "qt.text.font.db=false")
os.environ["MYOFFICE_NO_FIRST_RUN"] = "1"
WORK = Path(tempfile.mkdtemp(prefix="myoffice_shots_"))
os.environ["MYOFFICE_HOME"] = str(WORK / "home")
OUT = ROOT / "docs" / "screenshots"

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtGui import QTextCursor  # noqa: E402
from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

LOG: list[str] = []


def pump(app, seconds=0.3):
    end = time.time() + seconds
    while time.time() < end:
        app.processEvents()
        time.sleep(0.02)


def shot(widget, name, app, wait=0.3):
    pump(app, wait)
    path = OUT / f"{name}.png"
    widget.grab().save(str(path))
    LOG.append(f"{name}.png")
    print("saved", path.name)


def dialog_shot(dlg, name, app, size=None, wait=0.4):
    if size:
        dlg.resize(*size)
    dlg.show()
    shot(dlg, name, app, wait)
    dlg.close()


def make_samples():
    import docx
    import pymupdf

    s = WORK / "samples"
    s.mkdir(parents=True, exist_ok=True)
    with open(s / "customers.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Name", "City", "Amount"])
        for row in (("Ali Raza", "Sahiwal", "12,500"), ("Sara Khan", "Lahore", "8,900"), ("Usman Tariq", "Karachi", "15,200")):
            w.writerow(row)
    for name, n in (("report_a.pdf", 3), ("report_b.pdf", 2)):
        d = pymupdf.open()
        for i in range(n):
            p = d.new_page()
            p.insert_text((72, 90), f"{name.replace('.pdf', '').replace('_', ' ').title()} \u2014 page {i + 1}", fontsize=22)
            p.insert_text((72, 130), "Quarterly results: revenue grew 25% while costs fell 4%.", fontsize=12)
            p.draw_rect(pymupdf.Rect(72, 170, 520, 330), color=(0.17, 0.42, 0.9), fill=(0.9, 0.94, 1))
            p.insert_text((90, 200), "Highlights", fontsize=14)
        d.save(str(s / name))
    for v, text in (("v1", ["Project plan", "The project starts in October.", "Budget is Rs 2 million."]),
                    ("v2", ["Project plan", "The project starts in November with two teams.", "Budget is Rs 2.5 million.", "Risks are reviewed weekly."])):
        d = docx.Document()
        for t in text:
            d.add_paragraph(t)
        d.save(str(s / f"plan_{v}.docx"))
    return s


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    app = QApplication.instance() or QApplication(sys.argv[:1])
    QMessageBox.exec = lambda self: 0
    from app.services import fonts
    from app.services.config import cfg

    fonts.load_bundled_fonts()
    fonts.install_substitutions()
    cfg().set("author", "Ayesha Khan")
    cfg().set("firewall.consent", True, save=True)
    from app.services.i18n import install_translator
    from app.shell.theme import stylesheet

    install_translator(app, "en_US")
    app.setStyle("Fusion")
    app.setStyleSheet(stylesheet("light"))
    samples = make_samples()
    for p in (samples / "report_a.pdf", samples / "customers.csv"):
        cfg().add_recent(str(p))
    from app.shell.main_window import MainWindow

    win = MainWindow()
    win.resize(1440, 900)
    win.show()
    win.start_services()
    pump(app, 0.5)

    # ---------------------------------------------------------------- home / gallery
    shot(win, "01_home", app)
    win.home.show_page("templates")
    shot(win, "15_templates_gallery", app)
    win.home.show_page("tools")
    shot(win, "00_all_tools", app)
    win.home.show_page("home")

    # ---------------------------------------------------------------- 1 Writer (+ 16-19 add-ons)
    from app.writer.model import Comment, HeaderFooter, apply_style
    from app.writer.revisions import add_comment_anchor

    wd = win.new_writer()
    ed = wd.editor
    c = ed.textCursor()
    apply_style(c, "Title", 11)
    c.insertText("Annual Report 2026")
    c.insertBlock()
    apply_style(c, "Heading 1", 11)
    c.insertText("1. Summary")
    c.insertBlock()
    apply_style(c, "Normal", 11)
    c.insertText("MyOffice Writer supports styles, tables, images, headers and footers, tracked changes, comments, "
                 "citations and right-to-left Urdu text with Nastaleeq shaping.")
    ed.setTextCursor(c)
    c = ed.textCursor()
    c.insertBlock()
    ed.setTextCursor(c)
    ed.item._prepare_typing("\u0627\u0631\u062f\u0648")
    c = ed.textCursor()
    c.insertText("\u0645\u0627\u0626\u06cc \u0622\u0641\u0633 \u0645\u06cc\u06ba \u0627\u0631\u062f\u0648 \u0646\u0633\u062a\u0639\u0644\u06cc\u0642 \u0641\u0648\u0646\u0679 \u06a9\u06d2 \u0633\u0627\u062a\u06be \u062f\u0633\u062a\u0627\u0648\u06cc\u0632\u0627\u062a \u0644\u06a9\u06be\u06cc\u06ba\u06d4")
    ed.setTextCursor(c)
    wd.set_direction(Qt.RightToLeft)
    c = ed.textCursor()
    c.movePosition(QTextCursor.End)
    c.insertBlock()
    ed.setTextCursor(c)
    wd.set_direction(Qt.LeftToRight)
    wd.insert_table(3, 3)
    ed.hf = HeaderFooter(header="Annual Report \u2014 {TITLE}", footer="Page {PAGE} of {NUMPAGES}", header_align="right")
    ed.watermark = "DRAFT"
    ed.comments["c1"] = Comment("c1", "Ayesha Khan", "Please confirm these figures.", initials="AK")
    add_comment_anchor(ed.document(), 0, 6, "c1")
    wd.toggle_track(True)
    c = ed.textCursor()
    c.movePosition(QTextCursor.End)
    ed.setTextCursor(c)
    wd.replace_selection("Inserted with Track Changes on.")
    wd.a_comments_pane.setChecked(True)
    wd.toggle_comments_pane(True)
    wd.a_nav.setChecked(True)
    wd.toggle_nav(True)
    ed.refresh_decorations()
    shot(win, "02_writer", app, 0.6)
    wd.ribbon.select_page("Review")
    shot(win, "18_track_changes_comments", app)
    from app.writer import citations

    wd.sources = [{"id": "khan2024", "type": "book", "author": [{"family": "Khan", "given": "Ayesha"}],
                   "title": "Computing in Urdu", "issued": {"date-parts": [[2024]]}, "publisher": "Punjab Press"},
                  {"id": "ali2025", "type": "article-journal", "author": [{"family": "Ali", "given": "Usman"}],
                   "title": "Open office formats", "container-title": "Journal of Software", "issued": {"date-parts": [[2025]]},
                   "volume": "12", "page": "45-60"}]
    c = ed.textCursor()
    c.movePosition(QTextCursor.End)
    citations.insert_citation(c, ["khan2024"], "", wd.sources, "APA")
    ed.setTextCursor(c)
    wd.insert_bibliography()
    wd.ribbon.select_page("References")
    shot(win, "17_citation", app)
    dialog_shot(citations.SourceManagerDialog(wd, wd.sources), "17_citation_sources", app)
    from app.writer.mailmerge import MailMergeDialog

    mm = MailMergeDialog(wd)
    try:
        mm.load_source(str(samples / "customers.csv"))
    except Exception:
        pass
    dialog_shot(mm, "16_mail_merge", app, (900, 600))
    from app.writer.compare import compare_into, file_paragraphs
    from app.writer.widget import WriterDocument

    res = WriterDocument()
    compare_into(file_paragraphs(str(samples / "plan_v1.docx")), file_paragraphs(str(samples / "plan_v2.docx")),
                 res.editor.document(), "Compare")
    res.title_override = "Comparison"
    win.add_document(res)
    res.ribbon.select_page("Review")
    shot(win, "19_compare_documents", app, 0.5)

    # ---------------------------------------------------------------- 2 Sheet (+ 20-25)
    from app.sheet.model import ChartSpec, CondFormat

    sd = win.new_sheet()
    data = [("Region", "Q1", "Q2", "Q3", "Q4", "Total"), ("North", 120, 135, 150, 160), ("South", 90, 110, 95, 120),
            ("East", 150, 140, 170, 180), ("West", 80, 95, 105, 98), ("\u06a9\u0631\u0627\u0686\u06cc", 60, 75, 88, 92)]
    for r, row in enumerate(data, start=1):
        for col, v in enumerate(row, start=1):
            sd.commit_input(r, col, str(v))
    for r in range(2, 7):
        sd.commit_input(r, 6, f"=SUM(B{r}:E{r})")
    sd.commit_input(7, 1, "Total")
    for col in "BCDEF":
        sd.commit_input(7, " ABCDEF".index(col), f"=SUM({col}2:{col}6)")
    sd.view.select_range(1, 1, 1, 6)
    sd.set_style_attr(bold=True, fill="#1f4e79", color="#ffffff", halign="center")
    sd.view.select_range(7, 1, 7, 6)
    sd.cell_style("Total")
    sd.sheet.cond_formats.append(CondFormat((2, 6, 6, 6), "databar", fill="#63be7b"))
    sd.sheet.cond_formats.append(CondFormat((2, 2, 6, 5), "cell", ">", "140", fill="#c6efce", color="#006100"))
    sd.sheet.charts.append(ChartSpec("column", f"{sd.sheet.name}!A1:E6", "Quarterly Sales", anchor=(2, 8), size=(470, 290), id="c1"))
    sd._rebuild_floats()
    sd.model.cells_changed()
    sd.view.select_range(2, 2, 6, 5)
    shot(win, "03_spreadsheet", app, 0.6)
    from app.sheet.charts import ChartDialog

    dialog_shot(ChartDialog(sd, sd.book, ChartSpec("line", f"{sd.sheet.name}!A1:E6", "Trend")), "23_charts", app)
    from app.sheet.formula.parser import range_name
    from app.sheet.pivot import PivotDialog
    from app.sheet import dialogs as SD

    rows = [("Region", "Product", "Sales"), ("North", "Writer", 120), ("North", "Sheets", 80), ("South", "Writer", 95),
            ("South", "Slides", 60), ("East", "Sheets", 140), ("East", "Writer", 70), ("West", "Slides", 55)]
    sh2 = sd.add_sheet("Data")
    for r, row in enumerate(rows, start=1):
        for col, v in enumerate(row, start=1):
            sd.commit_input(r, col, str(v))
    pv = PivotDialog(sd, sh2, range_name(1, 1, len(rows), 3, "Data"), SD.parse_a1_range)
    pv._take(pv.fields, "Product", pv.cols)
    pv._preview()
    dialog_shot(pv, "20_pivot_table", app, (900, 560))
    from app.sheet.analysis import run_tool

    title, out_rows = run_tool("descriptive", sh2, (1, 3, len(rows), 3), True)
    sd.write_table(out_rows, title)
    shot(win, "21_data_analysis", app, 0.5)
    from app.macros.manager import MacroDialog

    md = MacroDialog(win)
    for i in range(md.list.count()):
        if "Monthly" in md.list.item(i).text():
            md.list.setCurrentRow(i)
    dialog_shot(md, "22_macros", app, (960, 620))
    from app.forms.designer import FormsDialog

    fd = FormsDialog(win, sd)
    fd.sample()
    dialog_shot(fd, "24_forms_survey", app, (1080, 640))
    from app.collab.binding import CollabSession
    from app.collab.dialog import CollabDialog
    from app.collab.server import CollabServer

    srv = CollabServer("127.0.0.1", 18999, persist=False)
    srv.start()
    sess = CollabSession(sd, "ws://127.0.0.1:18999", "budget-2026", "Ayesha")
    sd.collab_session = sess
    sess.start()
    from app.sheet.widget import SheetDocument

    peer = SheetDocument()
    peer_sess = CollabSession(peer, "ws://127.0.0.1:18999", "budget-2026", "Usman")
    pump(app, 1.0)
    peer_sess.start()
    pump(app, 2.0)
    cd = CollabDialog(sd)
    cd._peers(["Usman"])
    dialog_shot(cd, "25_realtime_collaboration", app, (560, 470), 0.6)

    # ---------------------------------------------------------------- 3 Slides
    sl = win.open_path(str(ROOT / "resources" / "templates" / "slide" / "project_status.dpt"), as_copy=True,
                       display="Project Status")
    sl.thumbs.setCurrentRow(1)
    pump(app, 0.3)
    if sl.canvas.items_:
        sl.canvas.items_[-1].setSelected(True)
    sl.ribbon.select_page("Animations")
    sl.a_anim_pane.setChecked(True)
    sl.anim_pane.setVisible(True)
    shot(win, "04_presentation", app, 0.6)
    from app.slide.show import SlideShow

    show = SlideShow(sl, 0)
    show.resize(1280, 720)
    show.show()
    shot(show, "04_presentation_slideshow", app, 0.6)
    show.close()

    # ---------------------------------------------------------------- 4 PDF viewer/editor + 14 edit
    pdfdoc = win.open_path(str(samples / "report_a.pdf"))
    pump(app, 0.5)
    words = pdfdoc.words(0)[:4]
    pdfdoc.annotate_words("highlight", 0, words)
    import pymupdf

    pdfdoc.rect_action("rect", 0, pymupdf.Rect(70, 360, 330, 430), pymupdf.Point(70, 360), pymupdf.Point(330, 430))
    from app.pdf.viewer import insert_rich_text

    pdfdoc._checkpoint()
    insert_rich_text(pdfdoc.doc[0], pymupdf.Rect(340, 450, 560, 520), "\u0645\u0646\u0638\u0648\u0631 \u0634\u062f\u06c1 \u2014 Approved", 14, "#188038", "urdu")
    pdfdoc._edited([0])
    pdfdoc.search.setText("revenue")
    pdfdoc.do_search()
    shot(win, "05_pdf_viewer", app, 0.6)
    pdfdoc.enter_edit_mode()
    shot(win, "14_edit_pdf", app, 0.4)

    # ---------------------------------------------------------------- 5-13 PDF toolbox
    tb = win.open_toolbox()
    names = {"convert": "06_pdf_convert", "merge": "07_merge_pdf", "split": "08_split_pdf", "compress": "09_compress_pdf",
             "ocr": "10_ocr", "sign": "11_sign_pdf", "protect": "12_protect_unlock", "watermark": "13_watermark",
             "extract": "13b_extract"}
    inputs = {"merge": [samples / "report_a.pdf", samples / "report_b.pdf"]}
    from app.pdf import ops

    for key, fname in names.items():
        tb.select(key)
        form = tb.forms[key]
        form.files.list.clear()
        form.files.add([str(p) for p in inputs.get(key, [samples / "report_a.pdf"])])
        if key == "merge":
            outs = form.work(form.files.files(), lambda *_: None)
            tb.finished(form.title, outs)
        elif key == "compress":
            res = ops.compress_pdf(str(samples / "report_a.pdf"), str(WORK / "report_a_compressed.pdf"))
            tb.log(f"report_a.pdf: {res.before:,} \u2192 {res.after:,} bytes ({res.saved_pct:.1f}% smaller, {res.method})")
        shot(win, fname, app, 0.3)
    from app.pdf.toolbox import SignDialog

    dialog_shot(SignDialog(win, 0, None, 842), "11_sign_pdf_dialog", app)

    # ---------------------------------------------------------------- 26-32 system tools
    from app.services import cloud_sync
    from app.tools.addin_dialog import AddinManagerDialog
    from app.tools.backup_dialog import BackupDialog
    from app.tools.cloud_sync_dialog import CloudSyncDialog
    from app.tools.font_manager import FontManagerDialog
    from app.tools.repair_dialog import RepairDialog
    from app.tools.update_dialog import UpdateDialog

    remote = WORK / "cloud_remote"
    cfg().set("sync", {"enabled": False, "backend": "folder", "local_dir": str(WORK / "MyOffice Cloud"),
                       "folder_target": str(remote), "interval_minutes": 10, "delete_remote": False,
                       "s3": {"endpoint": "", "bucket": "", "prefix": "MyOffice/", "region": "", "access_key": "", "secret_key": ""}})
    (WORK / "MyOffice Cloud").mkdir(exist_ok=True)
    shutil.copy(samples / "report_a.pdf", WORK / "MyOffice Cloud" / "report_a.pdf")
    csd = CloudSyncDialog(win)
    rep = cloud_sync.SyncEngine(str(WORK / "MyOffice Cloud"), cloud_sync.FolderBackend(str(remote))).sync_once()
    csd.log.appendPlainText(f"[{rep.started}] {rep.summary()}")
    for f in rep.uploaded:
        csd.log.appendPlainText(f"   uploaded: {f}")
    dialog_shot(csd, "26_cloud_sync", app, (640, 600))
    from app.services import backup

    tmpdoc = WORK / "budget.xlsx"
    sd.save_to(str(tmpdoc))
    backup.make_backup(str(tmpdoc), force=True)
    dialog_shot(BackupDialog(win), "27_backup_recovery", app, (860, 520))
    dialog_shot(AddinManagerDialog(win), "28_addin_manager", app, (800, 480))
    fm = FontManagerDialog(win)
    for i in range(fm.list.count()):
        if fm.list.item(i).text() == "Noto Nastaliq Urdu":
            fm.list.setCurrentRow(i)
    fm.sample.setText("\u06cc\u06c1 \u0627\u0631\u062f\u0648 \u0645\u062a\u0646 \u06c1\u06d2 \u2014 Nastaleeq")
    dialog_shot(fm, "29_font_manager", app, (920, 600))
    import docx

    d = docx.Document()
    d.add_paragraph("Important paragraph")
    d.save(str(WORK / "damaged.docx"))
    data = (WORK / "damaged.docx").read_bytes()
    (WORK / "damaged.docx").write_bytes(data[:-300])
    rd = RepairDialog(win, str(WORK / "damaged.docx"))
    from app.services.repair import repair_file

    r = repair_file(str(WORK / "damaged.docx"))
    rd.result = r
    rd.report.setPlainText(r.summary())
    rd.open_btn.setEnabled(r.ok)
    dialog_shot(rd, "30_file_repair", app, (640, 420))
    import hashlib
    import json

    (WORK / "MyOffice_Setup_1.1.0.exe").write_bytes(b"installer")
    (WORK / "version.json").write_text(json.dumps({"version": "1.1.0", "url": str(WORK / "MyOffice_Setup_1.1.0.exe"),
                                                    "sha256": hashlib.sha256(b"installer").hexdigest(),
                                                    "notes": "\u2022 Faster PDF rendering\n\u2022 New Urdu templates\n\u2022 Bug fixes"}))
    ud = UpdateDialog(win)
    ud.url.setText(str(WORK / "version.json"))
    from app.services.updater import check_for_update

    info = check_for_update(str(WORK / "version.json"), force=True)
    ud.info = info
    ud.status.setText(info.message().split("\n")[0])
    ud.notes.setPlainText(info.notes)
    ud.install_btn.setEnabled(True)
    dialog_shot(ud, "31_update_service", app, (560, 420))
    from app.services.capture import CaptureOptions

    dialog_shot(CaptureOptions(win), "32_screen_capture", app)
    from app.services.capture import grab_window

    grab_window(win).save(str(OUT / "32_screen_capture_result.png"))
    LOG.append("32_screen_capture_result.png")

    # ---------------------------------------------------------------- quality-bar dialogs
    bad = WORK / "corrupt.docx"
    bad.write_bytes(b"PK\x03\x04" + os.urandom(2000))
    from app.shell.document import DocumentError
    from app.writer.widget import WriterDocument as _WD

    try:
        _WD().load(str(bad))
    except DocumentError as exc:
        box, _rb = win.open_failed_box(str(bad), exc)
        box.show()
        shot(box, "92_corrupt_docx_error", app, 0.3)
        box.close()
    box, _f, _o = sd.large_file_box(str(WORK / "sales_2026_full.xlsx"), 104_857_600)
    box.show()
    shot(box, "93_large_xlsx_dialog", app, 0.3)
    box.close()

    # ---------------------------------------------------------------- Urdu UI
    install_translator(app, "ur_PK")
    from app.main import apply_ui_font

    apply_ui_font(app, "ur_PK")
    win2 = MainWindow()
    win2.resize(1440, 900)
    win2.show()
    shot(win2, "90_urdu_ui_home", app, 0.5)
    w2 = win2.new_writer()
    shot(win2, "91_urdu_ui_writer", app, 0.5)

    # cleanup
    for s in (sess, peer_sess):
        s.stop()
    srv.stop()
    for wnd in (win, win2):
        for doc in wnd.documents():
            doc.set_modified(False)
    peer.set_modified(False)
    w2.set_modified(False)
    (OUT / "index.txt").write_text("\n".join(LOG) + "\n", encoding="utf-8")
    print(f"{len(LOG)} screenshots in {OUT}")


if __name__ == "__main__":
    main()
