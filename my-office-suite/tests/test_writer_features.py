"""Writer 1.2: spelling, footnotes, equations, document protection, password encryption, version history,
merge/split documents and the table toolbox."""
import shutil
import subprocess
import time
import zipfile

import pytest

from app.services import ooxml_crypto, spell, versions
from app.writer import equations, notes


def _wait_spell(qapp, sp):
    for _ in range(60):
        if sp.ready():
            return
        sp.load_async()
        time.sleep(0.05)
        qapp.processEvents()


@pytest.fixture
def writer(qapp):
    from app.writer.widget import WriterDocument

    w = WriterDocument()
    yield w
    w.set_modified(False)
    w.deleteLater()


def test_spell_checker_languages(qapp):
    us, gb, es = spell.speller("en_US"), spell.speller("en_GB"), spell.speller("es_ES")
    assert not us.check("recieve") and us.suggest("recieve")[0] == "receive"
    assert us.check("color") and not us.check("colour") and us.suggest("colour") == ["color"]
    assert gb.check("colour") and gb.check("organisation") and not gb.check("color")
    assert es.check("canción") and es.suggest("cancion")[0] == "canción"
    assert us.check("NATO") and us.check("iPhone") and us.check("اردو") and us.check("हिंदी")
    us.add("MyOfficeWord")
    assert us.check("myofficeword")


def test_spelling_marks_and_dialog(writer, qapp):
    from PySide6.QtGui import QTextCharFormat

    from app.writer import features

    writer.editor.textCursor().insertText("Thiss sentense is fine otherwise.")
    _wait_spell(qapp, writer.editor.highlighter.speller)
    writer.editor.highlighter.rehighlight()
    b = writer.editor.document().begin()
    marked = {b.text()[r.start:r.start + r.length] for r in b.layout().formats()
              if r.format.underlineStyle() == QTextCharFormat.SpellCheckUnderline}
    assert marked == {"Thiss", "sentense"}
    d = features.SpellDialog(writer)
    d.next()
    assert d.word.text() == "Thiss"
    d.list.setCurrentRow(0)
    chosen = d.list.item(0).text()
    d.change()
    assert writer.editor.document().toPlainText().startswith(chosen + " ")
    assert d.word.text() == "sentense"  # the dialog moved on to the next misspelling


def test_footnotes_docx_pdf_and_libreoffice(writer, tmp_path):
    c = writer.editor.textCursor()
    c.insertText("Solar power grows")
    notes.insert_footnote(c, "IEA 2025.")
    c.insertText(" while costs fall")
    notes.insert_footnote(c, "Lazard LCOE.")
    c.insertText(".")
    writer.editor.setTextCursor(c)
    c.setPosition(5)
    notes.insert_footnote(c, "First note.")
    doc = writer.editor.document()
    assert doc.toPlainText() == "Solar1 power grows2 while costs fall3."
    assert [t for _p, _l, t, _i in notes.iter_footnotes(doc)] == ["First note.", "IEA 2025.", "Lazard LCOE."]
    path = tmp_path / "notes.docx"
    writer.save_to(str(path))
    z = zipfile.ZipFile(path)
    assert "word/footnotes.xml" in z.namelist() and z.read("word/document.xml").count(b"footnoteReference") == 3
    from app.writer.widget import WriterDocument

    w2 = WriterDocument()
    w2.load(str(path))
    assert [t for _p, _l, t, _i in notes.iter_footnotes(w2.editor.document())] == \
        ["First note.", "IEA 2025.", "Lazard LCOE."]
    pdf = tmp_path / "notes.pdf"
    w2.export_pdf_to(str(pdf))
    import pymupdf

    text = " ".join(pymupdf.open(pdf)[0].get_text().split())
    assert "1 First note." in text and "3 Lazard LCOE." in text
    w2.set_modified(False)
    soffice = shutil.which("soffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path / "lo"), str(path)],
                       capture_output=True, timeout=120)
        lo = tmp_path / "lo" / "notes.pdf"
        if lo.exists():
            assert "Lazard LCOE." in pymupdf.open(lo)[0].get_text()


def test_equations_round_trip(writer, tmp_path):
    from app.writer import features
    from app.writer.model import P_EQUATION

    for _name, latex in equations.EXAMPLES:
        assert equations.render_png(latex)[:4] == b"\x89PNG"
        back = equations.omml_to_latex(equations.latex_to_omml(latex))
        assert equations.render_png(back)[:4] == b"\x89PNG"
    with pytest.raises(equations.EquationError):
        equations.render_png(r"\frac{a}{")
    c = writer.editor.textCursor()
    c.insertText("Energy: ")
    c.insertImage(features._equation_format(writer, r"E = m c^{2}"))
    path = tmp_path / "eq.docx"
    writer.save_to(str(path))
    assert b"<m:oMath" in zipfile.ZipFile(path).read("word/document.xml")
    from app.writer.widget import WriterDocument

    w2 = WriterDocument()
    w2.load(str(path))
    found = []
    b = w2.editor.document().begin()
    it = b.begin()
    while not it.atEnd():
        f = it.fragment()
        it += 1
        if f.charFormat().isImageFormat() and f.charFormat().hasProperty(P_EQUATION):
            found.append(f.charFormat().property(P_EQUATION))
    assert found and "c" in found[0]
    w2.set_modified(False)


def test_document_protection(writer, tmp_path):
    assert notes.legacy_key("Example") == 0x64CEED7E  # ECMA-376 Part 4 example
    attrs = notes.protection_attributes("s3cret")
    assert notes.verify_protection("s3cret", attrs) and not notes.verify_protection("nope", attrs)
    writer.editor.textCursor().insertText("Locked text")
    writer.protection = attrs
    path = tmp_path / "prot.docx"
    writer.save_to(str(path))
    settings = zipfile.ZipFile(path).read("word/settings.xml").decode()
    assert 'w:edit="readOnly"' in settings and 'w:enforcement="1"' in settings
    from app.writer.widget import WriterDocument

    w2 = WriterDocument()
    w2.load(str(path))
    assert w2.protection and w2.a_protect.isChecked()
    assert notes.verify_protection("s3cret", w2.protection)
    w2.set_modified(False)


def test_password_encryption_open_and_save(qapp, tmp_path, monkeypatch):
    import docx
    from PySide6.QtWidgets import QInputDialog

    from app.shell.main_window import MainWindow

    plain = tmp_path / "plain.docx"
    d = docx.Document()
    d.add_paragraph("Top secret plan")
    d.save(plain)
    enc = tmp_path / "secret.docx"
    ooxml_crypto.encrypt_file(plain, enc, "pa55")
    assert ooxml_crypto.is_encrypted(enc) and not ooxml_crypto.is_encrypted(plain)
    with pytest.raises(ooxml_crypto.WrongPassword):
        ooxml_crypto.decrypt_to_temp(enc, "wrong")
    answers = iter([("wrong", True), ("pa55", True)])
    monkeypatch.setattr(QInputDialog, "getText", staticmethod(lambda *a, **k: next(answers)))
    win = MainWindow()
    doc = win.open_path(str(enc))
    assert doc is not None and doc.password == "pa55" and doc.path == enc
    assert "Top secret plan" in doc.editor.document().toPlainText()
    doc.editor.textCursor().insertText("Updated ")
    assert doc.save()
    assert ooxml_crypto.is_encrypted(enc)  # stays protected after saving
    tmp = ooxml_crypto.decrypt_to_temp(enc, "pa55")
    assert "Updated" in "".join(p.text for p in docx.Document(tmp).paragraphs)
    assert versions.list_versions(enc)  # every save is recorded
    doc.set_modified(False)
    win.close()


def test_version_history(tmp_path):
    p = tmp_path / "v.txt"
    p.write_text("one")
    v1 = versions.record(p)
    assert versions.record(p) == v1  # unchanged: no duplicate
    p.write_text("two")
    versions.record(p)
    items = versions.list_versions(p)
    assert len(items) == 2
    versions.restore(items[1], p)
    assert p.read_text() == "one"
    assert len(versions.list_versions(p)) >= 2


def test_merge_split_and_table_tools(writer, tmp_path, monkeypatch):
    import docx

    from app.writer import features

    for i in (1, 2):
        d = docx.Document()
        d.add_heading(f"Chapter {i}", 1)
        d.add_paragraph(f"Text of chapter {i}.")
        d.save(tmp_path / f"c{i}.docx")
    import app.shell.dialogs as dlg

    monkeypatch.setattr(dlg, "get_open_files", lambda *a, **k: [str(tmp_path / "c1.docx"), str(tmp_path / "c2.docx")])
    writer.editor.textCursor().insertText("Introduction")
    features.merge_documents(writer)
    text = writer.editor.document().toPlainText()
    assert "Chapter 1" in text and "Text of chapter 2." in text
    pts = features.split_points(writer.editor.document(), "heading", 0, writer.editor.page.page_px()[1])
    assert len(pts) == 3  # intro + two chapters
    # table toolbox
    c = writer.editor.textCursor()
    c.movePosition(c.MoveOperation.End)
    c.insertBlock()
    start = c.position()
    c.insertText("Name\tScore\nZara\t7\nAli\t9\nOmar\t3")
    c.setPosition(start)
    c.movePosition(c.MoveOperation.End, c.MoveMode.KeepAnchor)
    writer.editor.setTextCursor(c)
    assert features.text_to_table(writer)
    table = writer.editor.textCursor().currentTable()
    assert table is not None and table.rows() == 4
    cell = table.cellAt(1, 0).firstCursorPosition()
    writer.editor.setTextCursor(cell)
    assert features.sort_table(writer, 1, True, True)
    assert [features._cell_text(table, r, 0) for r in range(1, 4)] == ["Ali", "Zara", "Omar"]
    assert features.apply_table_style(writer, "Blue Banded")
    assert table.cellAt(0, 0).format().background().color().name() == "#1a73e8"
