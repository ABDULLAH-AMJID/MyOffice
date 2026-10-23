"""MyOffice 1.1: helper programs bundled in the setup, built-in Office->PDF, on-demand Legacy Formats Pack."""
import faulthandler
import hashlib
import http.server
import os
import re
import sys
import threading
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
ISS = (ROOT / "installer" / "setup.iss").read_text(encoding="utf-8")
sys.path.insert(0, str(ROOT / "tools"))


# ----------------------------------------------------------------------------- bundled helper programs
def _fake_install(tmp_path, monkeypatch):
    for rel in ("MyOffice.exe", "tesseract/tesseract.exe", "gs/gswin64c.exe", "tessdata/urd.traineddata"):
        f = tmp_path / rel
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_bytes(b"x")
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(tmp_path / "MyOffice.exe"))
    return tmp_path


def test_bundled_programs_are_used_first(tmp_path, monkeypatch):
    from app.pdf import ops
    from app.services import paths
    from app.services.config import cfg

    base = _fake_install(tmp_path, monkeypatch)
    cfg().set("ocr.tesseract_cmd", "")
    assert paths.install_dir() == base.resolve()
    assert ops.tesseract_cmd() == str(base.resolve() / "tesseract" / "tesseract.exe")
    assert ops.find_ghostscript() == str(base.resolve() / "gs" / "gswin64c.exe")
    assert ops.private_tessdata() == base.resolve() / "tessdata"
    custom = tmp_path / "my-tesseract.exe"  # an explicit setting still wins
    custom.write_bytes(b"x")
    cfg().set("ocr.tesseract_cmd", str(custom))
    try:
        assert ops.tesseract_cmd() == str(custom)
    finally:
        cfg().set("ocr.tesseract_cmd", "")


def test_no_bundled_programs_in_source_tree(monkeypatch):
    from app.services import paths

    monkeypatch.setattr(sys, "frozen", False, raising=False)
    assert paths.install_dir() is None
    assert paths.bundled_program("gs", "gswin64c.exe") is None


# ----------------------------------------------------------------------------- built-in Office -> PDF
def _samples(folder: Path, urdu: bool = False) -> dict:
    import docx
    import openpyxl
    from pptx import Presentation

    d = docx.Document()
    d.add_heading("Quarterly report", 1)
    d.add_paragraph("\u06cc\u06c1 \u0627\u0631\u062f\u0648 \u0645\u062a\u0646 \u06c1\u06d2" if urdu else "Plain English text.")
    d.save(folder / "a.docx")
    wb = openpyxl.Workbook()
    for i in range(1, 6):
        wb.active.cell(i, 1, i * 10)
    wb.active["B1"] = "=SUM(A1:A5)"
    wb.save(folder / "b.xlsx")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Launch plan"
    prs.save(folder / "c.pptx")
    (folder / "d.csv").write_text("name,amount\nalpha,17\n", encoding="utf-8")
    (folder / "e.html").write_text("<h1>Web page</h1><p>Hello</p>", encoding="utf-8")
    return {"a.docx": "Quarterly report", "b.xlsx": "150", "c.pptx": "Launch plan", "d.csv": "alpha",
            "e.html": "Web page"}


def _text(pdf: str) -> str:
    import pymupdf

    with pymupdf.open(pdf) as doc:
        return " ".join(p.get_text() for p in doc)


def test_office_to_pdf_without_libreoffice(qapp, tmp_path, monkeypatch):
    from app.pdf import ops
    from app.services import office_convert

    monkeypatch.setattr(office_convert, "find_soffice", lambda: None)
    for name, expect in _samples(tmp_path).items():
        out = ops.to_pdf(str(tmp_path / name), str(tmp_path / (name + ".pdf")))
        assert expect in _text(out), name


def test_builtin_export_from_worker_thread(qapp, tmp_path, monkeypatch):
    """The PDF Toolbox converts in a worker thread; rendering must be moved to the GUI thread."""
    from app.pdf import ops
    from app.services import office_convert
    from app.shell.tasks import run_task

    monkeypatch.setattr(office_convert, "find_soffice", lambda: None)
    _samples(tmp_path)
    files = [str(tmp_path / n) for n in ("a.docx", "b.xlsx", "c.pptx")]
    faulthandler.dump_traceback_later(120, exit=True)  # a deadlock must fail the run, not hang it
    try:
        outs = run_task(None, "convert", lambda fs, progress=None: [ops.to_pdf(f, f + ".w.pdf") for f in fs], files)
    finally:
        faulthandler.cancel_dump_traceback_later()
    assert [Path(o).name for o in outs] == ["a.docx.w.pdf", "b.xlsx.w.pdf", "c.pptx.w.pdf"]
    assert "Launch plan" in _text(outs[2])


def test_legacy_formats_still_need_libreoffice(qapp, tmp_path, monkeypatch):
    from app.pdf import ops
    from app.services import builtin_pdf, office_convert

    old = tmp_path / "old.doc"
    old.write_bytes(b"\xd0\xcf\x11\xe0 not really a doc")
    assert not builtin_pdf.can_export(str(old))
    monkeypatch.setattr(office_convert, "find_soffice", lambda: None)
    office_convert.set_missing_handler(None)
    with pytest.raises(ops.PdfOpError, match="Legacy Formats Pack"):
        ops.to_pdf(str(old), str(tmp_path / "old.pdf"))


def test_urdu_documents_prefer_libreoffice_when_installed(qapp, tmp_path, monkeypatch):
    """Qt's PDF engine cannot keep Urdu/Hindi text searchable; LibreOffice can, so it is used if present."""
    from app.pdf import ops
    from app.services import builtin_pdf, office_convert

    (tmp_path / "en").mkdir()
    (tmp_path / "ur").mkdir()
    _samples(tmp_path / "en")
    _samples(tmp_path / "ur", urdu=True)
    assert not builtin_pdf.has_complex_script(str(tmp_path / "en" / "a.docx"))
    assert builtin_pdf.has_complex_script(str(tmp_path / "ur" / "a.docx"))
    used = []

    def fake_convert(src, target, outdir=None, timeout=240):
        used.append(Path(src).parent.name)
        out = Path(outdir) / (Path(src).stem + ".pdf")
        import pymupdf

        doc = pymupdf.open()
        doc.new_page().insert_text((72, 72), "via LibreOffice")
        doc.save(out)
        return str(out)

    monkeypatch.setattr(office_convert, "find_soffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(office_convert, "convert", fake_convert)
    ops.to_pdf(str(tmp_path / "en" / "a.docx"), str(tmp_path / "en.pdf"))
    ops.to_pdf(str(tmp_path / "ur" / "a.docx"), str(tmp_path / "ur.pdf"))
    assert used == ["ur"]
    assert "Quarterly report" in _text(str(tmp_path / "en.pdf"))


def test_missing_libreoffice_handler(tmp_path, monkeypatch):
    from app.services import office_convert

    src = tmp_path / "letter.odt"
    src.write_bytes(b"x")
    calls = []
    monkeypatch.setattr(office_convert, "find_soffice", lambda: None)
    office_convert.set_missing_handler(lambda path: calls.append(path) or False)
    try:
        with pytest.raises(office_convert.LibreOfficeMissing, match="Legacy Formats Pack"):
            office_convert.convert(str(src), "docx", str(tmp_path / "out"))
        assert calls == [str(src)]
        if os.name == "posix":  # the handler "installs" LibreOffice -> the conversion continues
            fake = tmp_path / "soffice"
            fake.write_text('#!/bin/sh\nwhile [ "$1" != "--outdir" ]; do shift; done\n'
                            'out="$2"; src="$3"; b=$(basename "$src"); touch "$out/${b%.*}.docx"\n')
            fake.chmod(0o755)
            state = {"installed": False}
            monkeypatch.setattr(office_convert, "find_soffice", lambda: str(fake) if state["installed"] else None)
            office_convert.set_missing_handler(lambda path: state.update(installed=True) or True)
            res = office_convert.convert(str(src), "docx", str(tmp_path / "out"))
            assert Path(res).name == "letter.docx" and Path(res).is_file()
    finally:
        office_convert.set_missing_handler(None)


def test_cli_convert_uses_builtin_engine(qapp, tmp_path, monkeypatch):
    from app.services import cli, office_convert

    monkeypatch.setattr(office_convert, "find_soffice", lambda: None)
    _samples(tmp_path)
    assert cli._convert(str(tmp_path / "c.pptx"), "pdf", str(tmp_path / "out")) == 0
    assert "Launch plan" in _text(str(tmp_path / "out" / "c.pdf"))
    assert cli.parse(["--self-test", "-o", "r.txt"]).self_test


# ----------------------------------------------------------------------------- Legacy Formats Pack download
PAYLOAD = os.urandom(300_000)
PAYLOAD_SHA = hashlib.sha256(PAYLOAD).hexdigest()


class _Handler(http.server.BaseHTTPRequestHandler):
    ranges: list = []
    ignore_range = False

    def log_message(self, *a):
        pass

    def do_GET(self):
        if self.path.startswith("/missing"):
            self.send_error(404)
            return
        data = PAYLOAD if not self.path.startswith("/bad") else PAYLOAD[::-1]
        rng = self.headers.get("Range")
        _Handler.ranges.append(rng)
        start = 0
        if rng and not _Handler.ignore_range:
            start = int(re.match(r"bytes=(\d+)-", rng).group(1))
            self.send_response(206)
            self.send_header("Content-Range", f"bytes {start}-{len(data) - 1}/{len(data)}")
        else:
            self.send_response(200)
        self.send_header("Content-Length", str(len(data) - start))
        self.end_headers()
        self.wfile.write(data[start:])


@pytest.fixture
def server():
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
    t = threading.Thread(target=httpd.serve_forever, daemon=True)
    t.start()
    _Handler.ranges, _Handler.ignore_range = [], False
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()


def _dl(base, tmp_path, urls=("/pack.msi",), **kw):
    from app.services import legacy_pack

    return legacy_pack.download(urls=[base + u for u in urls], sha256=PAYLOAD_SHA, size=len(PAYLOAD),
                                dest_dir=tmp_path, name="pack.msi", **kw)


def test_pack_download_verifies_and_reports_progress(server, tmp_path):
    seen = []
    f = _dl(server, tmp_path, progress=lambda pct, msg: seen.append(pct))
    assert f.read_bytes() == PAYLOAD and seen and seen[-1] == 100
    assert not (tmp_path / "pack.msi.part").exists()
    _Handler.ranges.clear()
    assert _dl(server, tmp_path) == f and _Handler.ranges == []  # already there: no second download


def test_pack_download_resumes(server, tmp_path):
    (tmp_path / "pack.msi.part").write_bytes(PAYLOAD[:120_000])
    assert _dl(server, tmp_path).read_bytes() == PAYLOAD
    assert _Handler.ranges == ["bytes=120000-"]


def test_pack_download_restarts_when_server_ignores_range(server, tmp_path):
    _Handler.ignore_range = True
    (tmp_path / "pack.msi.part").write_bytes(PAYLOAD[:50_000])
    assert _dl(server, tmp_path).read_bytes() == PAYLOAD


def test_pack_download_uses_mirror_and_rejects_bad_files(server, tmp_path):
    from app.services import legacy_pack

    assert _dl(server, tmp_path, urls=("/missing.msi", "/pack.msi")).read_bytes() == PAYLOAD
    (tmp_path / "pack.msi").unlink()
    with pytest.raises(legacy_pack.PackError, match="checksum"):
        _dl(server, tmp_path, urls=("/bad.msi",))
    assert not (tmp_path / "pack.msi.part").exists() and not (tmp_path / "pack.msi").exists()


def test_pack_download_cancel_keeps_partial_file(server, tmp_path, monkeypatch):
    from app.services import legacy_pack

    monkeypatch.setattr(legacy_pack, "CHUNK", 16_384)
    n = {"calls": 0}

    def cancelled():
        n["calls"] += 1
        return n["calls"] > 3

    with pytest.raises(legacy_pack.DownloadCancelled):
        _dl(server, tmp_path, is_cancelled=cancelled)
    part = tmp_path / "pack.msi.part"
    assert 0 < part.stat().st_size < len(PAYLOAD)
    assert _dl(server, tmp_path).read_bytes() == PAYLOAD  # continues where it stopped


def test_pack_msiexec_command_line(tmp_path):
    from app.services import legacy_pack

    args = legacy_pack.msiexec_arguments(tmp_path / "LO.msi", tmp_path / "LO.log")
    assert args.startswith(f'/i "{tmp_path / "LO.msi"}" /qn /norestart ALLUSERS=1')
    assert "REGISTER_NO_MSO_TYPES=1" in args and args.endswith(f'/L*v "{tmp_path / "LO.log"}"')


def test_pack_dialog_and_offer_flow(qapp, monkeypatch):
    from app.services import office_convert
    from app.shell import legacy_pack_dialog as lpd

    dlg = lpd.LegacyPackDialog(None, "report.doc")
    assert dlg.windowTitle() == "Legacy Formats Pack" and dlg.install_btn.isDefault()
    assert dlg.file_btn.text().startswith("Use Installer File") and dlg.local_msi == ""
    monkeypatch.setattr(office_convert, "available", lambda: True)
    assert lpd.ensure_libreoffice(None, "x.doc") is True
    monkeypatch.setattr(office_convert, "available", lambda: False)
    if sys.platform != "win32":
        assert lpd.ensure_libreoffice(None, "x.doc") is False  # Linux: points to the package manager
    else:
        monkeypatch.setattr(lpd.LegacyPackDialog, "exec", lambda self: 0)  # "Not Now"
        assert lpd.ensure_libreoffice(None, "x.doc") is False


# ----------------------------------------------------------------------------- setup script + staging
def _defines() -> dict:
    return dict(re.findall(r'^#define\s+(\w+)\s+"([^"]*)"', ISS, re.M))


def _ispp(expr: str, defines: dict) -> str:
    return "".join(p.strip()[1:-1] if p.strip().startswith('"') else defines[p.strip()] for p in expr.split("+"))


def test_setup_and_app_pin_the_same_libreoffice():
    from app.services import legacy_pack

    d = _defines()
    assert d["LOVersion"] == legacy_pack.LO_VERSION and d["LOBuild"] == legacy_pack.LO_BUILD
    assert d["LOSha256"] == legacy_pack.LO_SHA256 and d["LOProperties"] == legacy_pack.MSI_PROPERTIES
    assert int(d["LOSizeMB"]) == round(legacy_pack.LO_SIZE / 1e6)
    exprs = dict(re.findall(r"^#define\s+(LOUrl2?)\s+(.+)$", ISS, re.M))
    assert (_ispp(exprs["LOUrl"], d), _ispp(exprs["LOUrl2"], d)) == legacy_pack.LO_URLS


def test_setup_bundles_the_helper_programs():
    for line in ('Source: "{#DepsDir}\\tesseract\\*"; DestDir: "{app}\\tesseract"',
                 'Source: "{#DepsDir}\\tessdata\\*"; DestDir: "{app}\\tessdata"',
                 'Source: "{#DepsDir}\\gs\\*"; DestDir: "{app}\\gs"',
                 'Source: "{#DepsDir}\\licenses\\*"; DestDir: "{app}\\licenses"'):
        assert line in ISS, line
    assert "#error The bundled helper programs are missing" in ISS
    task = next(ln for ln in ISS.splitlines() if ln.startswith('Name: "legacypack"'))
    assert "Flags: unchecked" in task and "Check: (not LibreOfficeInstalled) and" in task
    assert "UserMarkerExists('libreoffice')" in ISS  # uninstaller also offers what the app installed
    from app.services import legacy_pack

    assert "'" + legacy_pack.USER_MARKER_KEY + "'" in ISS and "'" + legacy_pack.SETUP_MARKER_KEY + "'" in ISS
    import inspect

    from app.pdf import ops

    assert 'bundled_program("tesseract", "tesseract.exe")' in inspect.getsource(ops.tesseract_cmd)
    assert 'bundled_program("gs", "gswin64c.exe")' in inspect.getsource(ops.find_ghostscript)


def test_staging_licence_gate(tmp_path):
    import fetch_windows_deps as fwd

    for t in {v[2] for v in fwd.LICENSES.values() if v[2]}:
        assert (ROOT / "installer" / "licenses" / t).is_file(), t
    for d in ("tesseract", "gs", "licenses"):
        (tmp_path / d).mkdir()
    (tmp_path / "tesseract" / "tesseract.exe").write_bytes(b"x")
    (tmp_path / "gs" / "gsdll64.dll").write_bytes(b"x")
    fwd.write_licenses(tmp_path)
    notice = (tmp_path / "licenses" / "THIRD_PARTY_PROGRAMS.txt").read_text(encoding="utf-8")
    assert "AGPL-3.0" in notice and "Apache-2.0" in notice
    (tmp_path / "tesseract" / "libsurprise-1.dll").write_bytes(b"x")  # unknown dependency -> stop
    with pytest.raises(fwd.BuildError, match="libsurprise"):
        fwd.write_licenses(tmp_path)


def test_selftest_report(tmp_path, monkeypatch):
    from app.services import selftest

    monkeypatch.setattr(selftest, "CHECKS", (("good", lambda w: "fine", True), ("info", lambda w: 1 / 0, False)))
    assert selftest.run(str(tmp_path / "r.txt")) == 0
    text = (tmp_path / "r.txt").read_text(encoding="utf-8")
    assert "[ OK ] good: fine" in text and "[INFO] info: ZeroDivisionError" in text
    monkeypatch.setattr(selftest, "CHECKS", (("bad", lambda w: 1 / 0, True),))
    assert selftest.run(str(tmp_path / "r2.txt")) == 1
    assert "RESULT: FAILED" in (tmp_path / "r2.txt").read_text(encoding="utf-8")


def test_selftest_builtin_pdf_check(qapp, tmp_path):
    from app.services import selftest

    assert ".docx" in selftest._check_builtin_pdf(tmp_path)


def test_run_task_updates_progress_in_gui_thread(qapp, monkeypatch):
    """Regression (found by the Wine test): with PySide6 6.12 a plain function connected to the worker's
    signals ran in the worker thread, so the progress dialog was touched from there and Qt crashed."""
    import time

    from PySide6.QtCore import QThread
    from PySide6.QtWidgets import QProgressDialog

    from app.services.builtin_pdf import call_in_gui_thread
    from app.shell.tasks import run_task

    seen = []
    orig = QProgressDialog.setValue

    def record(self, value):
        seen.append((self.windowTitle(), QThread.currentThread() == qapp.thread()))
        return orig(self, value)

    monkeypatch.setattr(QProgressDialog, "setValue", record)

    def slow(progress=None):
        for pct in (10, 60, 90):
            progress(pct, f"{pct}%")
            time.sleep(0.2)  # longer than the dialog's 400 ms minimum duration: it is really shown
        return "ok"

    assert run_task(None, "plain", slow) == "ok"
    # nested: a worker asks the GUI thread for a second task (download inside "Converting document...")
    assert run_task(None, "outer", lambda progress=None: call_in_gui_thread(lambda: run_task(None, "inner", slow))) == "ok"
    assert [t for t, _ in seen] == ["plain"] * 3 + ["inner"] * 3
    assert all(in_gui for _, in_gui in seen)


def test_uninstall_dialog_mentions_bundled_programs_and_libreoffice(qapp, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QLabel

    from app.services import legacy_pack
    from app.services.uninstall import UninstallPlan
    from app.shell.uninstall_dialog import UninstallDialog

    plan = UninstallPlan("installer", tmp_path / "MyOffice", tmp_path / "data", tmp_path / "MyOffice" / "unins000.exe")
    for by_myoffice in (False, True):
        monkeypatch.setattr(legacy_pack, "installed_by_myoffice", lambda: by_myoffice)
        d = UninstallDialog(plan)
        text = " ".join(lab.text() for lab in d.findChildren(QLabel))
        assert "OCR engine and Ghostscript" in text
        assert ("LibreOffice (Legacy Formats Pack) was installed by MyOffice" in text) == by_myoffice
        d.deleteLater()


def test_pack_cleanup_after_install(tmp_path, monkeypatch):
    from app.services import legacy_pack

    own, mine = tmp_path / "downloads", tmp_path / "usb"
    own.mkdir(), mine.mkdir()
    monkeypatch.setattr(legacy_pack, "download_dir", lambda: own)
    for p in (own / "LO.msi", own / "LO.msi.part", own / "LO.install.log", mine / "LO2.msi", own / "LO2.install.log"):
        p.write_bytes(b"x")
    legacy_pack.remove_download(own / "LO.msi")
    legacy_pack.remove_download(mine / "LO2.msi")  # the user's own file is never deleted
    assert list(own.iterdir()) == [] and [p.name for p in mine.iterdir()] == ["LO2.msi"]


def test_pack_accepts_only_the_pinned_installer_file(tmp_path, monkeypatch):
    from app.services import legacy_pack

    good = tmp_path / "LibreOffice.msi"
    good.write_bytes(PAYLOAD)
    monkeypatch.setattr(legacy_pack, "LO_SIZE", len(PAYLOAD))
    monkeypatch.setattr(legacy_pack, "LO_SHA256", PAYLOAD_SHA)
    seen = []
    assert legacy_pack.verify_installer(good, lambda pct, msg: seen.append(pct)) == good and seen[-1] == 100
    bad = tmp_path / "other.msi"
    bad.write_bytes(PAYLOAD[::-1])
    with pytest.raises(legacy_pack.PackError, match="not the LibreOffice"):
        legacy_pack.verify_installer(bad)
    with pytest.raises(legacy_pack.PackError, match="not found"):
        legacy_pack.verify_installer(tmp_path / "missing.msi")


def test_setup_installs_a_libreoffice_msi_lying_next_to_it():
    task = next(ln for ln in ISS.splitlines() if ln.startswith('Name: "legacypacklocal"'))
    assert "Flags: unchecked" not in task and "LocalLibreOfficeAvailable" in task  # ticked when the file is there
    assert "ExpandConstant('{src}\\LibreOffice_{#LOVersion}_Win_x86-64.msi')" in ISS
    assert "GetSHA256OfFile(Msi)" in ISS and "InstallLibreOffice(LocalLibreOfficeMsi, True)" in ISS
    download = next(ln for ln in ISS.splitlines() if ln.startswith('Name: "legacypack";'))
    assert "(not LocalLibreOfficeAvailable)" in download


def test_setup_removes_libreoffice_leftovers_only_when_it_is_gone():
    assert "procedure RemoveLeftoverDir(const Dir, ProgramFile: String);" in ISS
    assert "not FileExists(AddBackslash(Dir) + ProgramFile)" in ISS
    assert "RemoveLeftoverDir(ExpandConstant('{commonpf64}\\LibreOffice'), 'program\\soffice.exe');" in ISS


def test_running_task_dialog_stays_hidden_while_another_flow_is_open(qapp, monkeypatch):
    """Opening a .doc runs "Converting document..."; the Legacy Formats Pack offer that it opens must not be covered
    by that progress dialog (found under Wine).  QProgressDialog shows itself from an internal 4-second timer
    (forceShow), so that path is driven here directly, besides run_task's own 400 ms timer."""
    import time

    from PySide6.QtCore import QCoreApplication, QMetaObject

    from app.services.builtin_pdf import call_in_gui_thread
    from app.shell import tasks

    wait = (tasks.SHOW_AFTER_MS + 300) / 1000
    seen = {}

    def pump(seconds):
        end = time.monotonic() + seconds
        while time.monotonic() < end:
            QCoreApplication.processEvents()
            time.sleep(0.01)

    def flow():  # GUI thread, while the outer task runs and its dialog has not appeared yet
        outer = tasks._ACTIVE[-1]
        seen["outer_before"] = outer.isVisible()
        with tasks.other_tasks_paused():
            QMetaObject.invokeMethod(outer, "forceShow")  # what QProgressDialog's own timer does after 4 s
            outer.show()
            pump(wait)  # run_task's show timer fires in here
            seen["outer_during"] = outer.isVisible()
            tasks.run_task(None, "inner", lambda: time.sleep(wait) or "ok")
            seen["inner_ok"] = True
        seen["outer_after"] = outer.isVisible()
        return "done"

    shown = []
    orig = tasks._show_when_due
    monkeypatch.setattr(tasks, "_show_when_due", lambda dlg: (orig(dlg), shown.append((dlg.windowTitle(),
                                                                                         dlg.isVisible()))))
    assert tasks.run_task(None, "outer", lambda progress=None: call_in_gui_thread(flow)) == "done"
    assert seen == {"outer_before": False, "outer_during": False, "inner_ok": True, "outer_after": True}
    assert ("inner", True) in shown and ("outer", False) in shown  # the flow's own dialog does appear
    assert tasks._ACTIVE == []


def test_pack_checks_disk_space_first(tmp_path, monkeypatch):
    """A full disk gives a clear message before anything is downloaded or installed (not Windows Installer 1603)."""
    import collections
    import shutil

    from app.services import legacy_pack

    usage = collections.namedtuple("usage", "total used free")
    monkeypatch.setenv("ProgramW6432", str(tmp_path))
    monkeypatch.setattr(legacy_pack, "download_dir", lambda: tmp_path)
    monkeypatch.setattr(shutil, "disk_usage", lambda p: usage(10**12, 0, legacy_pack.LO_DISK_NEEDED + 10**8))
    assert legacy_pack.disk_space_shortfall() is None  # enough for the installation from a local file
    assert legacy_pack.pending_download_bytes() == legacy_pack.LO_SIZE
    needed, free, _drive = legacy_pack.disk_space_shortfall(legacy_pack.pending_download_bytes())
    assert needed == legacy_pack.LO_DISK_NEEDED + legacy_pack.LO_SIZE and free < needed  # not with the download
    (tmp_path / (legacy_pack.MSI_NAME + ".part")).write_bytes(b"x" * 1000)  # an interrupted download counts
    assert legacy_pack.pending_download_bytes() == legacy_pack.LO_SIZE - 1000

    def unreadable(p):
        raise OSError("no such drive")
    monkeypatch.setattr(shutil, "disk_usage", unreadable)
    assert legacy_pack.disk_space_shortfall() is None  # cannot be measured: Windows Installer decides
    assert "GetSpaceOnDisk64" in ISS and "if not EnoughSpaceForLibreOffice then" in ISS
