"""Quality bar: 50-page DOCX opens in < 2 s; 10k formulas recalculate in < 1 s."""
import time


def make_50_pages(path):
    import docx

    d = docx.Document()
    para = ("MyOffice performance test paragraph with enough words to wrap over several lines of the page "
            "so that pagination has real work to do. ") * 3
    for ch in range(1, 58):
        d.add_heading(f"Chapter {ch}", 1)
        for _ in range(6):
            d.add_paragraph(para)
        t = d.add_table(rows=3, cols=3)
        for r in range(3):
            for c in range(3):
                t.cell(r, c).text = f"R{r}C{c}"
    d.save(path)


def test_open_50_page_docx_under_2s(qapp, tmp_path):
    from app.writer.widget import WriterDocument

    path = tmp_path / "big.docx"
    make_50_pages(str(path))
    doc = WriterDocument()
    t = time.perf_counter()
    doc.load(str(path))
    pages = doc.editor.page_count()
    elapsed = time.perf_counter() - t
    print(f"\n50-page docx: {pages} pages opened in {elapsed:.3f} s")
    assert pages >= 50
    assert elapsed < 2.0
    doc.set_modified(False)


def test_10k_formulas_recalc_under_1s():
    from app.sheet.model import Book, Cell

    b = Book()
    sh = b.add_sheet()
    for r in range(1, 10001):
        sh.cells[(r, 1)] = Cell(float(r))
        c = Cell()
        sh.cells[(r, 2)] = c
        b.engine.set_formula(sh.id, r, 2, f"=IF(A{r}>5000,A{r}*2,A{r}+1)" if r % 2 else f"=B{r - 1}+A{r}", c)
    sh.invalidate_bounds()
    t = time.perf_counter()
    b.recalc_all()
    elapsed = time.perf_counter() - t
    print(f"\n10,000 formulas recalculated in {elapsed:.3f} s")
    assert b.engine.last_recalc_count == 10000
    assert elapsed < 1.0
    t = time.perf_counter()
    b.set_input(sh, 1, 1, "42")
    assert time.perf_counter() - t < 1.0
