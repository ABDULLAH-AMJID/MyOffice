"""System services: cloud sync, backup/recovery, repair, updater, templates, add-ins, fonts, CLI, i18n."""
import hashlib
import io
import json
import os
import time
import zipfile

import pytest


def test_cloud_sync_three_way(tmp_path):
    from app.services.cloud_sync import FolderBackend, SyncEngine

    remote = tmp_path / "remote"
    a = SyncEngine(str(tmp_path / "A"), FolderBackend(str(remote)), str(tmp_path / "sa.json"))
    b = SyncEngine(str(tmp_path / "B"), FolderBackend(str(remote)), str(tmp_path / "sb.json"))
    (tmp_path / "A" / "doc.txt").write_text("v1")
    assert a.sync_once().uploaded == ["doc.txt"]
    assert b.sync_once().downloaded == ["doc.txt"]
    time.sleep(1.1)
    (tmp_path / "B" / "doc.txt").write_text("v2 from B")
    b.sync_once()
    assert a.sync_once().downloaded == ["doc.txt"] and (tmp_path / "A" / "doc.txt").read_text() == "v2 from B"
    time.sleep(1.1)
    (tmp_path / "A" / "doc.txt").write_text("edit A")
    (tmp_path / "B" / "doc.txt").write_text("edit B")
    b.sync_once()
    rep = a.sync_once()
    assert rep.conflicts == ["doc.txt"]
    assert any("conflict" in n for n in os.listdir(tmp_path / "A"))


def test_s3_backend_with_stub(tmp_path):
    """S3Backend against an in-memory client implementing the boto3 calls it uses."""
    from datetime import datetime, timezone

    from app.services.cloud_sync import S3Backend, SyncEngine

    class Stub:
        def __init__(self):
            self.objs = {}

        def get_paginator(self, _name):
            stub = self

            class P:
                def paginate(self, Bucket, Prefix):
                    yield {"Contents": [{"Key": k, "Size": len(v), "LastModified": datetime.now(timezone.utc),
                                         "ETag": hashlib.md5(v).hexdigest()} for k, v in stub.objs.items() if k.startswith(Prefix)]}
            return P()

        def upload_file(self, local, bucket, key, ExtraArgs=None):
            self.objs[key] = open(local, "rb").read()

        def head_object(self, Bucket, Key):
            v = self.objs[Key]
            return {"ContentLength": len(v), "LastModified": datetime.now(timezone.utc), "ETag": hashlib.md5(v).hexdigest()}

        def download_file(self, bucket, key, local):
            open(local, "wb").write(self.objs[key])

        def delete_object(self, Bucket, Key):
            self.objs.pop(Key, None)

    stub = Stub()
    be = S3Backend("bucket", "MyOffice/", client=stub)
    eng = SyncEngine(str(tmp_path / "L"), be, str(tmp_path / "s.json"))
    (tmp_path / "L" / "x.docx").write_bytes(b"hello")
    assert eng.sync_once().uploaded == ["x.docx"]
    assert stub.objs["MyOffice/x.docx"] == b"hello"


def test_backup_and_recovery(tmp_path):
    from app.services import backup

    f = tmp_path / "report.docx"
    f.write_bytes(b"one")
    assert backup.make_backup(str(f), force=True) is not None
    entries = backup.list_backups("report.docx")
    assert entries and entries[0].original_name == "report.docx"
    rm = backup.RecoveryManager(tmp_path / "rec")
    snap = rm.snapshot("doc1", "Report", "writer", str(f), ".txt", lambda p: open(p, "w").write("snapshot"))
    assert snap.read_text() == "snapshot"
    rm.discard("doc1")
    assert not snap.exists()


def test_repair_damaged_files(tmp_path):
    import docx

    from app.services.repair import repair_file

    d = docx.Document()
    d.add_paragraph("Must survive")
    ok = tmp_path / "ok.docx"
    d.save(str(ok))
    data = ok.read_bytes()
    (tmp_path / "trunc.docx").write_bytes(data[:-200])
    rep = repair_file(str(tmp_path / "trunc.docx"))
    assert rep.ok and "Must survive" in "\n".join(p.text for p in docx.Document(rep.output).paragraphs)
    zin = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zout:
        for n in zin.namelist():
            c = zin.read(n)
            zout.writestr(n, c.replace(b"</w:body>", b"") if n == "word/document.xml" else c)
    (tmp_path / "bad.docx").write_bytes(buf.getvalue())
    rep = repair_file(str(tmp_path / "bad.docx"))
    assert rep.ok and rep.level == "repaired"


def test_updater_local_manifest(tmp_path):
    from app.services import updater

    pkg = tmp_path / "setup.exe"
    pkg.write_bytes(b"installer")
    man = tmp_path / "version.json"
    man.write_text(json.dumps({"version": "9.0.0", "url": str(pkg), "sha256": hashlib.sha256(b"installer").hexdigest()}))
    info = updater.check_for_update(str(man), force=True)
    assert info.available and info.latest == "9.0.0"
    assert os.path.exists(updater.download_update(info, str(tmp_path / "dl")))
    man.write_text(json.dumps({"version": "9.0.0", "url": str(pkg), "sha256": "0" * 64}))
    with pytest.raises(updater.UpdateError):
        updater.download_update(updater.check_for_update(str(man), force=True), str(tmp_path / "dl2"))


def test_templates_catalog(qapp, tmp_path):
    from app.services import templates

    found = templates.list_templates()
    kinds = {t.kind for t in found}
    assert {"writer", "sheet", "slide"} <= kinds and len(found) >= 12
    t = next(t for t in found if t.kind == "writer")
    tmp = templates.instantiate(t)
    assert tmp.endswith(".docx") and zipfile.is_zipfile(tmp)
    src = tmp_path / "mine.docx"
    import docx

    docx.Document().save(str(src))
    saved = templates.save_user_template("writer", "My Letter", str(src), b"\x89PNG fake")
    assert saved.path.endswith(".wpt") and any(x.name == "My Letter" for x in templates.list_templates("writer"))
    templates.delete_user_template(saved)


def test_addin_manager(qapp, tmp_path):
    from app.services.addinmgr import AddinManager

    mgr = AddinManager(None)
    mgr.load_all()
    ids = {a.id for a in mgr.addins.values() if a.loaded}
    assert {"myoffice.word_stats", "myoffice.urdu_digits"} <= ids
    from app.sheet.formula.functions import FUNCTIONS

    assert "URDUNUM" in FUNCTIONS
    # install a new add-in from a ZIP
    z = tmp_path / "hello.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("hello/manifest.json", json.dumps({"id": "test.hello", "name": "Hello", "version": "0.1"}))
        zf.writestr("hello/plugin.py", "def register(api):\n    api.register_sheet_function('HELLO', lambda: 'hi', 0, 0)\n")
    info = mgr.install_zip(str(z))
    assert info.loaded and "HELLO" in FUNCTIONS
    mgr.set_enabled("test.hello", False)
    assert "HELLO" not in FUNCTIONS
    mgr.uninstall("test.hello")
    mgr.unload_all()


def test_font_manager(qapp):
    from app.services import fonts

    fams = fonts.load_bundled_fonts()
    assert "Noto Nastaliq Urdu" in fams
    info = fonts.font_info(str(fonts.bundled_fonts()[0]))
    assert info["faces"][0]["glyphs"] > 100 and info["faces"][0]["scripts"]


def test_cli_and_language(tmp_path, monkeypatch):
    from app.services.cli import parse, run_headless_actions
    from app.services.config import cfg

    args = parse(["-setlng", "ur_PK"])
    assert run_headless_actions(args) == 0 and cfg().get("lang") == "ur_PK"
    args = parse(["--setservers", "update=https://example.com/version.json"])
    assert run_headless_actions(args) == 0 and cfg().get("update.url") == "https://example.com/version.json"
    cfg().set("lang", "en_US", save=True)


def test_translations_exist(qapp):
    from PySide6.QtCore import QTranslator

    from app.services.i18n import qm_path

    for lang in ("en_US", "ur_PK", "hi_IN"):
        assert qm_path(lang).exists()
    tr = QTranslator()
    assert tr.load(str(qm_path("ur_PK")))
    assert tr.translate("MyOffice", "Home") == "ہوم"


def test_capture_and_firewall_api(qapp):
    from app.services import capture, firewall

    img = capture.grab_screen().toImage()
    assert not img.isNull()
    assert firewall.consent() in (None, True, False)
