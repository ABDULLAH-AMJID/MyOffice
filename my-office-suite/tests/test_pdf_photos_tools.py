"""PDF 1.2 (crop, replace, organize, callouts, measurement, scan to PDF), MyOffice Photos (image operations and the
editor) and the new tools (batch rename, Resume Master, recorder)."""
import os

import numpy as np
import pymupdf
import pytest
from PIL import Image, ImageDraw

from app.pdf import pagetools, scan
from app.photos import ops


def _pdf(path, n=4):
    doc = pymupdf.open()
    for i in range(n):
        doc.new_page().insert_text((150, 300), f"Page {i + 1} content", fontsize=20)
    doc.save(path)
    return path


def test_crop_replace_organize(tmp_path):
    d = pymupdf.open(_pdf(str(tmp_path / "a.pdf")))
    assert pagetools.crop_pages(d, [0], None, auto=True) == 1
    box = d[0].cropbox
    assert box.width < 300 and box.height < 60
    pagetools.crop_pages(d, [1], (10, 10, 10, 10))
    assert round(d[1].cropbox.x0) == 28
    src = pymupdf.open()
    src.new_page().insert_text((72, 72), "REPLACEMENT", fontsize=30)
    pagetools.replace_pages(d, [2], src, [0])
    assert [p.get_text().split("\n")[0] for p in d] == ["Page 1 content", "Page 2 content", "REPLACEMENT",
                                                         "Page 4 content"]
    new = pagetools.apply_plan(d, [{"src": 3, "rot": 0}, {"src": 0, "rot": 90}, {"src": -1, "rot": 0}])
    assert new.page_count == 3 and [p.rotation for p in new] == [0, 90, 0]
    assert new[0].get_text().startswith("Page 4") and new[2].get_text() == ""


def test_callout_and_measurement(qapp, tmp_path, monkeypatch):
    import app.pdf.viewer as V

    v = V.PdfDocument()
    v.load(_pdf(str(tmp_path / "m.pdf"), 1))
    monkeypatch.setattr(V.QInputDialog, "getMultiLineText", staticmethod(lambda *a, **k: ("Check this", True)))
    v.add_callout(0, pymupdf.Point(200, 300), pymupdf.Point(350, 200))
    cm = 72 / 2.54
    v.add_distance_measure(0, pymupdf.Point(100, 400), pymupdf.Point(100 + 5 * cm, 400))
    v.add_area_measure(0, [pymupdf.Point(100, 500), pymupdf.Point(100 + 4 * cm, 500),
                           pymupdf.Point(100 + 4 * cm, 500 + 3 * cm), pymupdf.Point(100, 500 + 3 * cm)])
    infos = [(a.type[1], a.info.get("content", "")) for a in v.doc[0].annots()]
    assert ("FreeText", "Check this") in infos
    assert ("Line", "Distance: 5.00 cm") in infos
    assert any(t == "Polygon" and "12.00 cm" in c and "14.00 cm" in c for t, c in infos)
    v._scale = (cm, "m", 100.0)  # 1 cm on paper = 100 m
    assert v.distance_label(pymupdf.Point(0, 0), pymupdf.Point(2.5 * cm, 0)) == "250.00 m"
    v.set_modified(False)


def _photo_of_page(path):
    paper = Image.new("RGB", (1240, 1754), "white")
    dr = ImageDraw.Draw(paper)
    font = ops.find_font(46)
    for k, line in enumerate(["Invoice number 2026-118", "Total amount 45,500 PKR"]):
        dr.text((120, 200 + k * 120), line, fill=(20, 20, 20), font=font)
    photo = Image.new("RGB", (1500, 2000), (60, 50, 45))
    photo.paste(paper.rotate(3, expand=True, fillcolor=(60, 50, 45)), (90, 60))
    photo.save(path, quality=88)
    return path


def test_scan_to_pdf_cleanup_and_ocr(tmp_path):
    img = Image.open(_photo_of_page(str(tmp_path / "photo.jpg")))
    clean = scan.enhance(img, "document")
    assert clean.mode == "L" and abs(clean.width - 1240) < 60 and abs(clean.height - 1754) < 80  # cropped
    bw = scan.enhance(img, "bw")
    assert set(np.unique(np.asarray(bw))) <= {0, 255}
    out = scan.save_scan([clean], str(tmp_path / "scan.pdf"), ocr=True, lang="eng")
    assert pymupdf.open(out).page_count == 1
    from app.pdf.ops import tesseract_cmd

    if tesseract_cmd():
        text = " ".join(pymupdf.open(out)[0].get_text().split())
        assert "Invoice" in text and "45,500" in text


def test_photo_operations(tmp_path):
    w, h = 600, 400
    arr = np.zeros((h, w, 3), np.uint8)
    arr[:240] = (120, 170, 235)
    arr[240:] = (60, 145, 50)
    img = Image.fromarray(arr)
    ImageDraw.Draw(img).rectangle((290, 100, 310, 330), fill=(40, 30, 20))
    mask = Image.new("L", img.size, 0)
    ImageDraw.Draw(mask).rectangle((284, 94, 316, 336), fill=255)
    out = np.asarray(ops.inpaint(img, mask)).astype(int)
    assert np.abs(out[150, 300] - [120, 170, 235]).max() < 25   # sky
    assert np.abs(out[300, 300] - [60, 145, 50]).max() < 25     # grass
    # the horizon continues through the filled area (a soft edge of 1-2 px is fine)
    assert np.abs(out[236, 300] - [120, 170, 235]).max() < 30 and np.abs(out[244, 300] - [60, 145, 50]).max() < 30
    # background removal keeps a subject that touches the bottom edge
    portrait = Image.new("RGB", (400, 500), (40, 110, 200))
    d = ImageDraw.Draw(portrait)
    d.ellipse((130, 80, 270, 240), fill=(224, 172, 140))
    d.rectangle((90, 230, 310, 500), fill=(30, 30, 30))
    a = np.asarray(ops.remove_background(portrait).getchannel("A"))
    assert a[5, 5] == 0 and a[495, 20] == 0 and a[160, 200] == 255 and a[400, 200] == 255
    photo = ops.id_photo(portrait, (35, 45), "#ffffff")
    assert photo.size == (413, 531) and photo.getpixel((3, 3)) == (255, 255, 255)
    sheet, n = ops.print_sheet(photo)
    assert sheet.size == (1800, 1200) and n == 8
    page = Image.new("RGB", (400, 200), "white")
    pd = ImageDraw.Draw(page)
    pd.text((20, 20), "Printed text", fill=(0, 0, 0))
    pd.line((20, 100, 300, 120), fill=(20, 60, 220), width=4)
    cleaned = np.asarray(ops.remove_pen_marks(page))
    assert tuple(cleaned[110, 160]) == (255, 255, 255)
    for name in ops.FILTERS:
        assert ops.apply_filter(img, name).size == img.size
    assert ops.upscale(img, 2).size == (1200, 800)
    assert ops.adjust(img, ops.Adjust(brightness=30, vignette=50, temperature=-40)).size == img.size
    big = Image.fromarray((np.random.default_rng(3).random((600, 800, 3)) * 30 +
                           np.linspace(0, 200, 800)[None, :, None]).astype(np.uint8))
    assert ops.compress(big, str(tmp_path / "s.jpg"), "small") < ops.compress(big, str(tmp_path / "h.jpg"), "high")
    lossless = tmp_path / "l.webp"
    ops.compress(big, str(lossless), "lossless")
    assert np.array_equal(np.asarray(Image.open(lossless).convert("RGB")), np.asarray(big))


def test_photos_editor(qapp, tmp_path):
    from app.photos.widget import PhotosDocument

    p = tmp_path / "pic.png"
    Image.new("RGB", (300, 200), "#3366cc").save(p)
    doc = PhotosDocument()
    doc.load(str(p))
    assert doc.image.size == (300, 200) and not doc.is_modified()
    doc.apply(ops.rotate, 90)
    assert doc.image.size == (200, 300) and doc.is_modified()
    doc.undo()
    assert doc.image.size == (300, 200)
    doc.redo()
    assert doc.image.size == (200, 300)
    doc.apply(ops.crop, (10, 10, 110, 60))
    assert doc.image.size == (100, 50)
    out = tmp_path / "out.jpg"
    doc.save_to(str(out))
    assert Image.open(out).size == (100, 50)
    doc._filter_clicked(type("It", (), {"data": lambda self, role: "Sepia"})())
    assert doc.image.size == (100, 50)
    doc.set_modified(False)


def test_batch_rename(tmp_path):
    from app.tools import batch_rename as BR

    files = []
    for n in ("b.txt", "a.txt", "c.txt"):
        f = tmp_path / n
        f.write_text(n)
        files.append(str(f))
    names = BR.new_names(files, "doc_{n}", start=1, digits=2)
    assert names == ["doc_01.txt", "doc_02.txt", "doc_03.txt"]
    with pytest.raises(ValueError):
        BR.plan(files, ["same.txt"] * 3)
    done = BR.execute(BR.plan(files, names))
    assert sorted(os.listdir(tmp_path)) == ["doc_01.txt", "doc_02.txt", "doc_03.txt"]
    assert (tmp_path / "doc_01.txt").read_text() == "b.txt"
    BR.execute([(dst, src) for src, dst in done])  # undo
    assert sorted(os.listdir(tmp_path)) == ["a.txt", "b.txt", "c.txt"]
    swapped = BR.execute(BR.plan([str(tmp_path / "a.txt"), str(tmp_path / "b.txt")], ["b.txt", "a.txt"]))
    assert len(swapped) == 2 and (tmp_path / "a.txt").read_text() == "b.txt"


def test_resume_master(tmp_path):
    import docx

    from app.tools import resume

    for style in resume.STYLES:
        path = tmp_path / f"cv_{style.replace(' ', '_')}.docx"
        resume.build_docx(resume.SAMPLE, style, str(path))
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs) + "\n".join(
            c.text for t in d.tables for row in t.rows for c in row.cells)
        assert "Ayesha Khan" in text and "Punjab Textiles" in text and "IFRS" in text


def test_video_recorder_writes_mp4(qapp, tmp_path):
    from PySide6.QtCore import QEventLoop, QSize, QTimer
    from PySide6.QtGui import QColor, QImage

    from app.services import recorder

    ok, _why = recorder.multimedia_available()
    if not ok:
        pytest.skip("Qt Multimedia not available")
    out = tmp_path / "rec.mp4"
    rec = recorder.FrameRecorder(str(out), QSize(320, 180), fps=10)
    errors = []
    rec.error.connect(errors.append)
    rec.start()
    loop = QEventLoop()
    count = [0]

    def frame():
        img = QImage(320, 180, QImage.Format_ARGB32)
        img.fill(QColor.fromHsv(count[0] * 20 % 360, 200, 220))
        rec.push(img)
        count[0] += 1
        if count[0] == 20:
            timer.stop()
            rec.stop()
            QTimer.singleShot(1200, loop.quit)
    timer = QTimer(interval=50, timeout=frame)
    timer.start()
    loop.exec()
    assert not errors and out.exists() and out.stat().st_size > 5000
    data = out.read_bytes()
    assert b"ftyp" in data and b"moov" in data


def test_ocr_language_choices_only_offer_installed_languages(monkeypatch):
    """Found in the installed 1.2 build: Photos > Picture to Text offered "Arabic + English", but the setup ships
    only Urdu, Hindi and English OCR data, so that choice always ended in a Tesseract error."""
    from app.pdf import ops

    items = [("urd+eng", "Urdu + English"), ("eng", "English"), ("hin+eng", "Hindi + English"),
             ("ara+eng", "Arabic + English")]
    monkeypatch.setattr(ops, "ocr_languages", lambda: ["eng", "hin", "osd", "urd"])   # what MyOffice Setup installs
    assert [code for code, _label in ops.usable_ocr_choices(items)] == ["urd+eng", "eng", "hin+eng"]
    monkeypatch.setattr(ops, "ocr_languages", lambda: ["ara", "eng"])
    assert [code for code, _label in ops.usable_ocr_choices(items)] == ["eng", "ara+eng"]
    monkeypatch.setattr(ops, "ocr_languages", lambda: [])      # no Tesseract: the OCR call explains what is missing
    assert ops.usable_ocr_choices(items) == items
