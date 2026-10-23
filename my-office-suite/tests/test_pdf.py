"""PDF suite: merge, split, compress, protect/unlock, watermark, extract, convert, sign, OCR, repair."""
import os
import shutil

import pymupdf
import pytest

from app.pdf import ops


def test_merge_pdfs(sample_pdfs, tmp_path):
    out = tmp_path / "merged.pdf"
    ops.merge_pdfs(sample_pdfs, str(out))
    d = pymupdf.open(str(out))
    assert d.page_count == 5
    assert [t[1] for t in d.get_toc()] == ["a", "b"]
    assert "b.pdf page 1" in d[3].get_text()


def test_merge_rejects_encrypted_without_password(sample_pdfs, tmp_path):
    locked = tmp_path / "locked.pdf"
    ops.protect_pdf(sample_pdfs[0], str(locked), "pw")
    with pytest.raises(ops.PdfOpError):
        ops.merge_pdfs([str(locked), sample_pdfs[1]], str(tmp_path / "x.pdf"))


def test_split_modes(sample_pdfs, tmp_path):
    merged = tmp_path / "m.pdf"
    ops.merge_pdfs(sample_pdfs, str(merged))
    assert len(ops.split_pdf(str(merged), str(tmp_path / "r"), "ranges", "1-2, 4, 5-")) == 3
    assert len(ops.split_pdf(str(merged), str(tmp_path / "e"), "every", every=2)) == 3
    assert len(ops.split_pdf(str(merged), str(tmp_path / "b"), "bookmarks")) == 2
    with pytest.raises(ops.PdfOpError):
        ops.parse_ranges("9-12", 5)


def test_compress(sample_pdfs, tmp_path):
    src = tmp_path / "img.pdf"
    d = pymupdf.open()
    p = d.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 1200, 900), 0)
    pix.clear_with(180)
    p.insert_image(p.rect, pixmap=pix)
    d.save(str(src))
    res = ops.compress_pdf(str(src), str(tmp_path / "small.pdf"), "ebook")
    assert res.after <= res.before and os.path.exists(res.output)


def test_protect_unlock(sample_pdfs, tmp_path):
    locked = tmp_path / "locked.pdf"
    ops.protect_pdf(sample_pdfs[0], str(locked), "secret")
    assert ops.pdf_info(str(locked))["needs_password"]
    with pytest.raises(ops.PdfOpError):
        ops.unlock_pdf(str(locked), str(tmp_path / "x.pdf"), "wrong")
    ops.unlock_pdf(str(locked), str(tmp_path / "open.pdf"), "secret")
    assert not ops.pdf_info(str(tmp_path / "open.pdf"))["needs_password"]


def test_watermark_and_extract(sample_pdfs, tmp_path):
    out = tmp_path / "wm.pdf"
    ops.watermark_pdf(sample_pdfs[0], str(out), text="CONFIDENTIAL", opacity=0.3)
    assert "CONFIDENTIAL" in pymupdf.open(str(out))[0].get_text()
    txt = ops.pdf_to_text(sample_pdfs[0], str(tmp_path / "a.txt"))
    assert "quick brown fox" in open(txt, encoding="utf-8").read()
    assert pymupdf.open(ops.extract_pages(sample_pdfs[0], str(tmp_path / "p.pdf"), "2-3")).page_count == 2


def test_convert_from_pdf(sample_pdfs, tmp_path):
    import docx
    import openpyxl

    d = ops.pdf_to_docx(sample_pdfs[0], str(tmp_path / "a.docx"))
    assert "quick brown fox" in "\n".join(p.text for p in docx.Document(d).paragraphs)
    x = ops.pdf_to_xlsx(sample_pdfs[0], str(tmp_path / "a.xlsx"))
    assert openpyxl.load_workbook(x).worksheets
    assert len(ops.pdf_to_images(sample_pdfs[0], str(tmp_path / "img"), dpi=50)) == 3


def test_images_to_pdf(tmp_path):
    from PIL import Image

    img = tmp_path / "x.png"
    Image.new("RGB", (200, 120), "red").save(img)
    out = ops.images_to_pdf([str(img), str(img)], str(tmp_path / "imgs.pdf"))
    assert pymupdf.open(out).page_count == 2


def test_sign_and_verify(sample_pdfs, tmp_path):
    p12 = tmp_path / "me.p12"
    ops.create_self_signed_p12(str(p12), "Test Signer", "t@example.com", "pw")
    signed = tmp_path / "signed.pdf"
    ops.sign_pdf(sample_pdfs[0], str(signed), str(p12), "pw", page=0, box=(350, 50, 560, 120), reason="Approve")
    sigs = ops.verify_signatures(str(signed))
    assert len(sigs) == 1 and sigs[0].intact and sigs[0].valid
    # tampering breaks the signature
    data = bytearray(signed.read_bytes())
    idx = data.find(b"page 1")
    data[idx:idx + 6] = b"PAGE 9"
    tampered = tmp_path / "tampered.pdf"
    tampered.write_bytes(bytes(data))
    assert not all(s.intact for s in ops.verify_signatures(str(tampered)))


@pytest.mark.skipif(ops.tesseract_cmd() is None, reason="Tesseract not installed")
def test_ocr_urdu_english(qapp, tmp_path):
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QColor, QFont, QImage, QPainter

    img = QImage(1400, 400, QImage.Format_RGB32)
    img.fill(QColor("white"))
    p = QPainter(img)
    f = QFont("DejaVu Sans")
    f.setPixelSize(56)
    p.setFont(f)
    p.setPen(Qt.black)
    p.drawText(40, 120, "Invoice number 4521")
    f2 = QFont("DejaVu Sans")  # Naskh-style Urdu glyphs, shaped by HarfBuzz
    f2.setPixelSize(64)
    p.setFont(f2)
    p.drawText(40, 300, "پاکستان زندہ باد")
    p.end()
    path = tmp_path / "scan.png"
    img.save(str(path))
    pdf = ops.images_to_pdf([str(path)], str(tmp_path / "scan.pdf"))
    res = ops.ocr_pdf(pdf, str(tmp_path / "searchable.pdf"), "urd+eng", dpi=200)
    assert res["ocr_pages"] == 1
    page = pymupdf.open(res["output"])[0]
    text = ops.fix_rtl_text(page.get_text())
    assert "Invoice" in text and "4521" in text
    assert "پاکستان زندہ باد" in text  # Urdu recognised and restored to reading order
    assert any(page.search_for(v) for v in ops.search_variants("زندہ باد"))  # searchable


def test_rtl_visual_logical_conversion():
    assert ops.visual_to_logical("داب ہدنز ناتسکاپ") == "پاکستان زندہ باد"
    assert ops.visual_to_logical(ops.visual_to_logical("قیمت 4521 روپے")) == "قیمت 4521 روپے"
    assert ops.fix_rtl_text("Invoice 12\nداب ہدنز") == "Invoice 12\nزندہ باد"


@pytest.mark.skipif(shutil.which("soffice") is None and shutil.which("libreoffice") is None,
                    reason="LibreOffice not installed")
def test_legacy_office_to_pdf(tmp_path):
    """Legacy formats (here RTF) go through LibreOffice; .docx/.xlsx/.pptx use MyOffice's own engines
    (tests/test_offline_setup.py)."""
    import unicodedata

    src = tmp_path / "letter.rtf"
    src.write_text(r"{\rtf1\ansi Converted by LibreOffice headless\par}", encoding="ascii")
    out = ops.to_pdf(str(src), str(tmp_path / "letter.pdf"))
    text = unicodedata.normalize("NFKC", pymupdf.open(out)[0].get_text())  # "ff" may come out as a ligature
    assert "LibreOffice headless" in text


def test_replace_text(sample_pdfs, tmp_path):
    n = ops.replace_text(sample_pdfs[0], str(tmp_path / "r.pdf"), "lazy dog", "happy cat")
    assert n == 3 and "happy cat" in pymupdf.open(str(tmp_path / "r.pdf"))[0].get_text()
