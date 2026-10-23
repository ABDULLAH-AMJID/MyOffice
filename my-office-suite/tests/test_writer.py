"""Writer: DOCX save/load round trip, formatting, Urdu, tables, tracked changes, comments, PDF export."""
import zipfile

import pytest
from PySide6.QtCore import Qt
from PySide6.QtGui import QFont, QTextCharFormat, QTextCursor


@pytest.fixture
def writer(qapp):
    from app.writer.widget import WriterDocument

    d = WriterDocument()
    yield d
    d.set_modified(False)
    d.deleteLater()


def build_sample(doc):
    from app.writer.model import Comment, HeaderFooter, apply_style
    from app.writer.revisions import add_comment_anchor

    ed = doc.editor
    c = ed.textCursor()
    apply_style(c, "Title", 11)
    c.insertText("Quarterly Report")
    c.insertBlock()
    apply_style(c, "Heading 1", 11)
    c.insertText("Summary")
    c.insertBlock()
    apply_style(c, "Normal", 11)
    bold = QTextCharFormat()
    bold.setFontWeight(QFont.Bold)
    c.insertText("Bold words", bold)
    c.insertText(" and normal text.", QTextCharFormat())
    c.insertBlock()
    ed.setTextCursor(c)
    ed.item._prepare_typing("اردو")
    c = ed.textCursor()
    c.insertText("یہ اردو متن ہے")
    ed.setTextCursor(c)
    doc.set_direction(Qt.RightToLeft)
    c = ed.textCursor()
    c.movePosition(QTextCursor.End)
    c.insertBlock()
    ed.setTextCursor(c)
    doc.set_direction(Qt.LeftToRight)
    doc.insert_table(2, 3)
    ed.hf = HeaderFooter(header="Header {TITLE}", footer="Page {PAGE} of {NUMPAGES}")
    ed.watermark = "DRAFT"
    ed.comments["c1"] = Comment("c1", "Tester", "Please review")
    add_comment_anchor(ed.document(), 0, 9, "c1")


def test_docx_round_trip(writer, tmp_path):
    from app.writer.widget import WriterDocument

    build_sample(writer)
    path = tmp_path / "report.docx"
    writer.save_to(str(path))
    assert zipfile.is_zipfile(path)
    other = WriterDocument()
    other.load(str(path))
    text = other.editor.document().toPlainText()
    assert "Quarterly Report" in text and "Summary" in text and "یہ اردو متن ہے" in text
    assert other.editor.hf.footer == "Page {PAGE} of {NUMPAGES}"
    assert other.editor.hf.header == "Header {TITLE}"
    assert other.editor.watermark == "DRAFT"
    assert len(other.editor.comments) == 1
    # heading level and RTL direction survive
    doc = other.editor.document()
    levels = [doc.findBlockByNumber(i).blockFormat().headingLevel() for i in range(doc.blockCount())]
    assert 1 in levels
    rtl = [b for b in (doc.findBlockByNumber(i) for i in range(doc.blockCount())) if "اردو" in b.text()][0]
    assert rtl.textDirection() == Qt.RightToLeft
    # bold run preserved
    blk = [doc.findBlockByNumber(i) for i in range(doc.blockCount()) if "Bold words" in doc.findBlockByNumber(i).text()][0]
    it = blk.begin()
    assert it.fragment().charFormat().fontWeight() >= QFont.DemiBold
    # table preserved
    from PySide6.QtGui import QTextTable

    assert any(isinstance(f, QTextTable) for f in doc.rootFrame().childFrames())
    other.set_modified(False)


def test_tracked_changes_round_trip(writer, tmp_path):
    from app.writer.model import P_DEL, P_INS
    from app.writer.revisions import fragments
    from app.writer.widget import WriterDocument

    ed = writer.editor
    c = ed.textCursor()
    c.insertText("Original sentence.")
    ed.setTextCursor(c)
    writer.toggle_track(True)
    c = ed.textCursor()
    c.movePosition(QTextCursor.End)
    ed.setTextCursor(c)
    writer.replace_selection(" Added text")
    path = tmp_path / "tracked.docx"
    writer.save_to(str(path))
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "<w:ins " in xml
    other = WriterDocument()
    other.load(str(path))
    kinds = {("ins" if f.hasProperty(P_INS) else "del" if f.hasProperty(P_DEL) else "plain")
             for _p, _l, f in fragments(other.editor.document())}
    assert "ins" in kinds
    assert other.editor.revisions.accept_all() >= 1
    assert "Added text" in other.editor.document().toPlainText()
    other.set_modified(False)


def test_corrupt_docx_raises_document_error(qapp, tmp_path):
    from app.shell.document import DocumentError
    from app.writer.widget import WriterDocument

    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"PK\x03\x04 this is not a real docx")
    with pytest.raises(DocumentError):
        WriterDocument().load(str(bad))


def test_pdf_export_and_formats(writer, tmp_path):
    import pymupdf

    from app.writer.widget import WriterDocument

    build_sample(writer)
    pdf = tmp_path / "out.pdf"
    writer.export_pdf_to(str(pdf))
    d = pymupdf.open(str(pdf))
    assert d.page_count >= 1
    text = d[0].get_text()
    assert "Quarterly Report" in text and "Page 1 of" in text
    for ext in ("odt", "html", "md", "txt"):
        p = tmp_path / f"out.{ext}"
        writer.save_to(str(p))
        assert p.stat().st_size > 0
    again = WriterDocument()
    again.load(str(tmp_path / "out.md"))
    assert "Quarterly Report" in again.editor.document().toPlainText()


def test_compare_documents(qapp):
    from PySide6.QtGui import QTextDocument

    from app.writer.compare import compare_into, count_revisions

    out = QTextDocument()
    stats = compare_into(["Hello world", "Second line", "Gone"], ["Hello brave world", "Second line", "New paragraph"], out)
    ins, dele = count_revisions(out)
    assert stats.inserted_words >= 1 and stats.deleted_words >= 1
    assert ins > 0 and dele > 0


def test_citations(writer):
    from app.writer import citations

    sources = [{"id": "khan2024", "type": "book", "author": [{"family": "Khan", "given": "Ayesha"}],
                "title": "Urdu Computing", "issued": {"date-parts": [[2024]]}, "publisher": "Lahore Press"}]
    c = writer.editor.textCursor()
    citations.insert_citation(c, ["khan2024"], "", sources, "APA")
    assert "(Khan, 2024)" in writer.editor.document().toPlainText()
    entry = "".join(seg[0] if isinstance(seg, (tuple, list)) else str(seg) for seg in citations.bib_entry("MLA", sources[0]))
    assert "Khan" in entry and "Urdu Computing" in entry


def test_mail_merge(qapp, tmp_path):
    import csv

    from app.writer import mailmerge

    p = tmp_path / "people.csv"
    with open(p, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Name", "City"])
        w.writerow(["Ali", "Sahiwal"])
        w.writerow(["Sara", "Lahore"])
    headers, rows = mailmerge.read_data_source(str(p))
    assert headers == ["Name", "City"] and len(rows) == 2
    from PySide6.QtGui import QTextDocument

    tpl = QTextDocument()
    tpl.setPlainText("Dear «Name» from «City»")
    out = mailmerge.merged_clone(tpl, rows[1])
    assert out.toPlainText() == "Dear Sara from Lahore"


def test_numbering_level_without_marker_is_not_a_list(qapp, tmp_path):
    """LibreOffice's DOCX export links Heading 1 to an outline level with numFmt "none" and an empty lvlText:
    the heading must not show a bullet, while a real numbered paragraph keeps its list."""
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from PySide6.QtGui import QTextDocument

    from app.writer.docx_io import load_docx

    d = docx.Document()
    numbering = d.part.numbering_part.element
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="none"/><w:suff w:val="nothing"/>'
        '<w:lvlText w:val=""/><w:lvlJc w:val="left"/></w:lvl></w:abstractNum>'))
    numbering.insert(1, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="91">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>'
        '<w:lvlJc w:val="left"/></w:lvl></w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="91"><w:abstractNumId w:val="91"/></w:num>'))
    style = d.styles["Heading 1"].element
    style.get_or_add_pPr().insert(0, parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="90"/></w:numPr>'))
    d.add_heading("Outline heading", 1)
    item = d.add_paragraph("Numbered item")
    item._p.get_or_add_pPr().insert(0, parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="91"/></w:numPr>'))
    path = tmp_path / "lo_headings.docx"
    d.save(path)

    qdoc = QTextDocument()
    load_docx(str(path), qdoc, lambda *a, **k: None)
    blocks = {}
    b = qdoc.begin()
    while b.isValid():
        blocks[b.text()] = b
        b = b.next()
    assert blocks["Outline heading"].textList() is None
    assert blocks["Numbered item"].textList() is not None
